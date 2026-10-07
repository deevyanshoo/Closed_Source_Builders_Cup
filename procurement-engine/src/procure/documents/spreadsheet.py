"""XLSX cells as located segments, for example `Offers!C7`.

Every non-empty cell on every sheet becomes one segment, in sheet then row order. Cell text
is the stored value, not Excel's display formatting; the number format is kept beside it.
Numbers use the shortest exact decimal spelling of the stored value ("2615", "3.42").

Formulas are never evaluated here. A formula with a cached result reports that result; one
without (common for generated files) reports the formula text, flagged
`formula_without_cached_value`. Hidden sheets, rows and columns and near-white font are
flagged rather than dropped.
"""
from __future__ import annotations

import datetime as dt
import colorsys
import io
import posixpath
import re
import zipfile
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.xml.constants import PKG_REL_NS, REL_NS, SHEET_MAIN_NS, XLSM, XLSX, XLTM, XLTX

from .models import DocumentReadError, make_segment

MAX_SHEET_CELLS = 200_000
MAX_XLSX_MEMBERS = 1_024
MAX_XLSX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_XLSX_EXPANSION_RATIO = 200
_PLAIN_SHEET = re.compile(r"^[A-Za-z_][A-Za-z0-9_.]*$")
_CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_WORKSHEET_REL_TYPE = REL_NS + "/worksheet"
_WORKBOOK_TYPES = {XLSX, XLSM, XLTX, XLTM}


def _package_path(target: str, parent: str = "") -> str:
    """Resolve an internal OOXML target without accepting path aliases or traversal."""
    if not target or "\\" in target or "\x00" in target or "?" in target or "#" in target:
        raise DocumentReadError("unreadable", "XLSX has an invalid package path")
    absolute = target.startswith("/")
    parts = target[1:].split("/") if absolute else target.split("/")
    if any(part in ("", ".", "..") or ":" in part for part in parts):
        raise DocumentReadError("unreadable", "XLSX has an ambiguous or traversing package path")
    path = "/".join(parts) if absolute else posixpath.join(parent, *parts)
    if not path or path.startswith("/"):
        raise DocumentReadError("unreadable", "XLSX has an invalid package path")
    return path


def _worksheet_members(archive: zipfile.ZipFile, members: dict[str, zipfile.ZipInfo]):
    """Follow the same workbook sheet relationships that openpyxl will load."""
    manifest = members.get("[Content_Types].xml")
    if manifest is None:
        raise DocumentReadError("unreadable", "XLSX has no content types")
    with archive.open(manifest) as source:
        content_types = ElementTree.parse(source).getroot()
    if content_types.tag != f"{{{_CONTENT_TYPES_NS}}}Types":
        raise DocumentReadError("unreadable", "XLSX has invalid content types")
    workbook_paths = []
    for override in content_types.findall(f"{{{_CONTENT_TYPES_NS}}}Override"):
        if override.get("ContentType") in _WORKBOOK_TYPES:
            workbook_paths.append(_package_path(override.get("PartName", "")))
    if len(workbook_paths) != 1:
        raise DocumentReadError("unreadable", "XLSX workbook part is missing or ambiguous")
    workbook_path = workbook_paths[0]
    if workbook_path not in members:
        raise DocumentReadError("unreadable", "XLSX workbook part is missing")
    rels_path = posixpath.join(posixpath.dirname(workbook_path), "_rels",
                               posixpath.basename(workbook_path) + ".rels")
    if rels_path not in members:
        raise DocumentReadError("unreadable", "XLSX workbook relationships are missing")
    with archive.open(members[workbook_path]) as source:
        workbook = ElementTree.parse(source).getroot()
    with archive.open(members[rels_path]) as source:
        relationships = ElementTree.parse(source).getroot()
    if workbook.tag != f"{{{SHEET_MAIN_NS}}}workbook" or relationships.tag != f"{{{PKG_REL_NS}}}Relationships":
        raise DocumentReadError("unreadable", "XLSX workbook relationships are invalid")
    by_id = {}
    for rel in relationships:
        if rel.tag != f"{{{PKG_REL_NS}}}Relationship":
            raise DocumentReadError("unreadable", "XLSX workbook relationship is malformed")
        rel_id = rel.get("Id")
        if not rel_id or rel_id in by_id or not rel.get("Type") or not rel.get("Target"):
            raise DocumentReadError("unreadable", "XLSX workbook relationship is missing or duplicated")
        by_id[rel_id] = rel
    sheets = workbook.find(f"{{{SHEET_MAIN_NS}}}sheets")
    if sheets is None:
        raise DocumentReadError("unreadable", "XLSX workbook has no sheet list")
    sheet_members = []
    seen_targets = set()
    for sheet in sheets:
        if sheet.tag != f"{{{SHEET_MAIN_NS}}}sheet":
            raise DocumentReadError("unreadable", "XLSX workbook sheet is malformed")
        rel = by_id.get(sheet.get(f"{{{REL_NS}}}id"))
        if rel is None or rel.get("Type") != _WORKSHEET_REL_TYPE or rel.get("TargetMode") not in (None, "Internal"):
            raise DocumentReadError("unreadable", "XLSX worksheet relationship is missing or unsupported")
        target = _package_path(rel.get("Target"), posixpath.dirname(workbook_path))
        if target in seen_targets or target not in members or members[target].is_dir():
            raise DocumentReadError("unreadable", "XLSX worksheet target is missing or duplicated")
        seen_targets.add(target)
        sheet_members.append(members[target])
    return sheet_members


def sheet_locator(sheet: str, coordinate: str) -> str:
    if _PLAIN_SHEET.match(sheet):
        return f"{sheet}!{coordinate}"
    return "'" + sheet.replace("'", "''") + f"'!{coordinate}"


def cell_text(value) -> tuple[str, str]:
    """(text, value_type) for a stored cell value."""
    if isinstance(value, str):
        return value, "string"
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), "boolean"
    if isinstance(value, int):
        return str(value), "number"
    if isinstance(value, float):
        return (str(int(value)) if value.is_integer() else repr(value)), "number"
    if isinstance(value, dt.datetime):
        return (value.date().isoformat() if value.time() == dt.time() else value.isoformat()), "date"
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat(), "date" if isinstance(value, dt.date) else "time"
    return str(value), "other"


def _theme_colors(wb) -> list[str | None]:
    try:
        root = ElementTree.fromstring(wb.loaded_theme)
        scheme = root.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}clrScheme")
        colors = {}
        for entry in scheme:
            color = next(iter(entry), None)
            if color is None:
                continue
            kind = color.tag.rsplit("}", 1)[-1]
            colors[entry.tag.rsplit("}", 1)[-1]] = (
                color.get("lastClr") if kind == "sysClr" else color.get("val") if kind == "srgbClr" else None)
        # Spreadsheet theme indices use light/dark pairs in this order, unlike XML's dk1/lt1 order.
        order = ("lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3",
                 "accent4", "accent5", "accent6", "hlink", "folHlink")
        return [colors.get(name) for name in order]
    except (TypeError, ValueError, IndexError, AttributeError, ElementTree.ParseError):
        return []


def _color_rgb(color, wb, theme_colors: list[str | None]) -> str | None:
    if color is None:
        return None
    if color.type == "rgb":
        value = color.rgb
    elif color.type == "indexed":
        index = color.indexed
        palette = wb._colors if isinstance(wb._colors, (list, tuple)) else COLOR_INDEX
        value = palette[index] if isinstance(index, int) and 0 <= index < len(palette) else None
    elif color.type == "theme":
        index = color.theme
        value = theme_colors[index] if isinstance(index, int) and 0 <= index < len(theme_colors) else None
    else:
        return None
    if not isinstance(value, str) or len(value) not in (6, 8):
        return None
    value = value[-6:]
    try:
        channels = [int(value[i:i + 2], 16) for i in (0, 2, 4)]
        tint = float(color.tint or 0)
    except (TypeError, ValueError):
        return None
    if not -1 <= tint <= 1:
        return None
    if tint:
        hue, lightness, saturation = colorsys.rgb_to_hls(*(channel / 255 for channel in channels))
        lightness = lightness * (1 + tint) if tint < 0 else lightness * (1 - tint) + tint
        channels = [round(channel * 255) for channel in colorsys.hls_to_rgb(hue, lightness, saturation)]
    return "".join(f"{channel:02X}" for channel in channels)


def _near_white_font(cell, wb, theme_colors: list[str | None]) -> bool:
    color = cell.font.color if cell.font else None
    if color is None:
        return False
    rgb = _color_rgb(color, wb, theme_colors)
    if rgb is None:
        return True  # Unknown color is unsafe to assume visible on a white sheet.
    return all(int(rgb[i:i + 2], 16) >= 0xF0 for i in (0, 2, 4))


def _preflight(data: bytes) -> None:
    """Check ZIP expansion and stored cell coordinates before openpyxl allocates cells."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_MEMBERS:
                raise DocumentReadError("too_large", f"XLSX has more than {MAX_XLSX_MEMBERS} ZIP members")
            expanded = 0
            by_name = {}
            for info in members:
                name = info.filename[:-1] if info.is_dir() else info.filename
                _package_path(name)
                if name in by_name:
                    raise DocumentReadError("unreadable", "XLSX has duplicate ZIP member names")
                by_name[name] = info
                expanded += info.file_size
                if expanded > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise DocumentReadError("too_large", "XLSX declared uncompressed content is too large")
                if info.file_size > MAX_XLSX_EXPANSION_RATIO * max(info.compress_size, 1):
                    raise DocumentReadError("too_large", "XLSX ZIP member expands too much")
            occupied = rectangular_work = 0
            for info in _worksheet_members(archive, by_name):
                max_row = max_col = 0
                with archive.open(info) as sheet:
                    root = None
                    for event, element in ElementTree.iterparse(sheet, events=("start", "end")):
                        if root is None:
                            root = element
                        if event == "end":
                            if element.tag.rsplit("}", 1)[-1] == "row":
                                root.clear()
                            continue
                        if element.tag.rsplit("}", 1)[-1] != "c":
                            continue
                        coordinate = element.get("r")
                        if not coordinate:
                            raise DocumentReadError("unreadable", "XLSX cell has no coordinate")
                        try:
                            letters, row = coordinate_from_string(coordinate)
                            col = column_index_from_string(letters)
                        except ValueError as exc:
                            raise DocumentReadError("unreadable", "XLSX cell has an invalid coordinate") from exc
                        max_row, max_col = max(max_row, row), max(max_col, col)
                        occupied += 1
                        if occupied > MAX_SHEET_CELLS or rectangular_work + max_row * max_col > MAX_SHEET_CELLS:
                            raise DocumentReadError("too_large", f"XLSX exceeds {MAX_SHEET_CELLS} cell work units")
                rectangular_work += max_row * max_col
    except (zipfile.BadZipFile, RuntimeError, OSError, ElementTree.ParseError) as exc:
        raise DocumentReadError("unreadable", f"XLSX archive could not be checked: {exc}") from exc


def _hidden_columns(ws) -> set[int]:
    hidden = set()
    for dim in ws.column_dimensions.values():
        if dim.hidden and dim.min and dim.max:
            hidden.update(range(dim.min, dim.max + 1))
    return hidden


def read_xlsx(data: bytes, document_id: str):
    _preflight(data)
    try:
        wb = load_workbook(io.BytesIO(data), data_only=False)
        cached = load_workbook(io.BytesIO(data), data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise DocumentReadError("unreadable", f"XLSX could not be opened: {exc}") from exc

    segments, count = [], 0
    theme_colors = _theme_colors(wb)
    for ws in wb.worksheets:
        sheet_flags = () if ws.sheet_state == "visible" else ("hidden_sheet",)
        hidden_cols = _hidden_columns(ws)
        cached_ws = cached[ws.title]
        for _, cell in sorted(ws._cells.items()):
            if cell.value is None:
                continue
            count += 1
            if count > MAX_SHEET_CELLS:
                raise DocumentReadError("too_large", f"Workbook has more than {MAX_SHEET_CELLS} non-empty cells")
            flags = list(sheet_flags)
            formula = None
            if cell.data_type == "f":
                formula = cell.value if isinstance(cell.value, str) else getattr(cell.value, "text", str(cell.value))
                result = cached_ws[cell.coordinate].value
                if result is None:
                    text, value_type = formula, "formula"
                    flags.append("formula_without_cached_value")
                else:
                    text, value_type = cell_text(result)
            else:
                text, value_type = cell_text(cell.value)
            if not text.strip():
                continue
            if ws.row_dimensions[cell.row].hidden:
                flags.append("hidden_row")
            if cell.column in hidden_cols:
                flags.append("hidden_column")
            if _near_white_font(cell, wb, theme_colors):
                flags.append("low_visibility")
            segments.append(make_segment(
                document_id, "sheet_cell", sheet_locator(ws.title, cell.coordinate), text,
                sheet=ws.title, cell=cell.coordinate, row=cell.row, column=cell.column, flags=tuple(flags),
                value_type=value_type, number_format=cell.number_format, formula=formula))
    metadata = {"sheets": wb.sheetnames,
                "hidden_sheets": [ws.title for ws in wb.worksheets if ws.sheet_state != "visible"]}
    return segments, [], metadata
