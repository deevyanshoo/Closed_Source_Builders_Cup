"""Fixture-only host admission and native binary64 workbook snapshots.

There is deliberately no observed registration API. The registry is an in-process
host configuration seam, not a signature or protection against hostile same-UID
Python. Product requests cannot mint bindings by copying fields or APPROVED text.
The parent sees the entire snapshot; this is not physical test-data isolation.
"""
from __future__ import annotations

import hashlib
import io
import json
import math
import os
import re
import stat
import sys
import weakref
import zipfile
from dataclasses import dataclass
from fractions import Fraction
from pathlib import Path, PurePosixPath
from types import MappingProxyType
from xml.etree import ElementTree

MAX_SOURCE = 12 * 1024 * 1024
MAX_MEMBER = 16 * 1024 * 1024
MAX_EXPANSION = 64 * 1024 * 1024
SCHEMA = "observed-admission-fixture.v1"
SERIES = {"copper": (65, "Copper", "($/mt)"), "aluminum": (63, "Aluminum", "($/mt)")}
OFFLINE_ACCEPTANCE_SHA256 = "ea92d939b06e9c48a25bd03becc0039135e98af90efe380aff6adf41bdee4983"
_HELD_OBSERVED_SHA256 = "ea1c350827878ea3bbe30e3cda16a13fd3bd5b409b8458940dc94a36b5a33154"
_REGISTRY = weakref.WeakKeyDictionary()
_RUNS = weakref.WeakKeyDictionary()
_RUN_LIVES = weakref.WeakKeyDictionary()


def digest(raw):
    return hashlib.sha256(raw).hexdigest()


def canonical(value):
    return json.dumps(value, sort_keys=True, separators=(",", ":"), allow_nan=False).encode()


def _pairs(pairs):
    out = {}
    for key, value in pairs:
        if key in out:
            raise ValueError("duplicate JSON key")
        out[key] = value
    return out


def decode(raw):
    def invalid(_):
        raise ValueError("nonfinite JSON")
    value = json.loads(raw, object_pairs_hook=_pairs, parse_constant=invalid)
    canonical(value)
    return value


def _absolute(path):
    path = Path(path)
    if not path.is_absolute() or ".." in path.parts:
        raise ValueError("absolute canonical local path required")
    return path


def _identity(s, directory=False):
    base = (s.st_dev, s.st_ino, s.st_mode, s.st_uid)
    # Ancestor content timestamps change when an owned sibling output is created.
    # Anchor directory identity/mode/owner, not unrelated directory contents.
    return base if directory else (*base, s.st_nlink, s.st_size, s.st_mtime_ns, s.st_ctime_ns)


class Lifetime:
    """Own every descriptor immediately, including failures at the first fstat."""

    def __init__(self):
        self.fds = []
        self.records = []
        self.closed = False

    def open(self, path, *, private=False, limit=MAX_SOURCE):
        path = _absolute(path)
        parent = None
        for index, part in enumerate(path.parts):
            if index == 0:
                fd = os.open("/", os.O_RDONLY | os.O_DIRECTORY)
                at = Path("/")
            else:
                at = at / part
                directory = index < len(path.parts) - 1
                flags = os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK
                if directory:
                    flags |= os.O_DIRECTORY
                fd = os.open(part, flags, dir_fd=parent)
            self.fds.append(fd)  # before fstat, never retry close
            info = os.fstat(fd)
            directory = index < len(path.parts) - 1
            if directory and not stat.S_ISDIR(info.st_mode):
                raise ValueError("ancestor not directory")
            if not directory:
                if not stat.S_ISREG(info.st_mode) or info.st_nlink != 1 or info.st_size > limit:
                    raise ValueError("regular single-link bounded file required")
                if private:
                    pinfo = os.fstat(parent)
                    if info.st_uid != os.getuid() or pinfo.st_uid != os.getuid() or pinfo.st_mode & 0o077:
                        raise ValueError("source/control root must be private")
            self.records.append((at, fd, _identity(info, directory), directory))
            parent = fd
        return fd

    def check(self, *, postclose=False):
        for path, fd, expected, directory in self.records:
            info = os.lstat(path)
            if _identity(info, directory) != expected:
                raise ValueError("path identity changed")
            if not postclose and _identity(os.fstat(fd), directory) != expected:
                raise ValueError("descriptor identity changed")

    def read(self, fd, limit):
        chunks, total = [], 0
        while True:
            chunk = os.read(fd, min(65536, limit + 1 - total))
            if not chunk:
                break
            chunks.append(chunk); total += len(chunk)
            if total > limit:
                raise ValueError("bounded read exceeded")
        self.check()
        return b"".join(chunks)

    def close(self):
        if self.closed:
            raise ValueError("lifetime already closed")
        self.closed = True
        failed = False
        for fd in reversed(self.fds):
            try:
                os.close(fd)
            except OSError:
                failed = True
        try:
            self.check(postclose=True)
        except (OSError, ValueError):
            failed = True
        if failed:
            raise ValueError("close/postclose identity unverified")


@dataclass(frozen=True, eq=False)
class RootAdmissionBinding:
    """Opaque identity; constructing this type does not register authority."""
    fixture_id: str


@dataclass(frozen=True, eq=False)
class RootObservedRunBinding:
    fixture_id: str


def _registered(binding, registry):
    try:
        value = registry.get(binding)
    except TypeError:
        value = None
    if value is None:
        raise ValueError("registered fixture binding required; observed execution HOLD")
    return value


def register_owned_fixture(source_file, admission_raw, *, controls=()):
    """Trusted test-host setup ONLY. No grant for real originals or real fitting.

    Controls are owning-host paths, never product request data. Hashing the owned
    fixture happens in its generating test, not here before permission checking.
    """
    grant = decode(admission_raw)
    required = {"schema_version", "fixture_id", "is_synthetic", "observed_authenticity",
                "observed_execution_authorized", "source_sha256", "source_bytes", "as_of_utc"}
    if (set(grant) != required or grant["schema_version"] != SCHEMA
            or not isinstance(grant["fixture_id"], str) or not 0 < len(grant["fixture_id"]) <= 128):
        raise ValueError("fixture grant shape mismatch")
    if (grant["is_synthetic"] is not True or grant["observed_authenticity"] is not False
            or grant["observed_execution_authorized"] is not False):
        raise ValueError("only synthetic fixture registration; observed execution HOLD")
    if grant["as_of_utc"] != "2026-10-08T00:00:00.000000Z":
        raise ValueError("fixture UTC proxy mismatch")
    sha = grant["source_sha256"]
    if not isinstance(sha, str) or len(sha) != 64 or any(c not in "0123456789abcdef" for c in sha):
        raise ValueError("fixture source hash invalid")
    if sha == _HELD_OBSERVED_SHA256 or Path(source_file).name == "worldbank_pinksheet_monthly.xlsx":
        raise ValueError("named observed original is not a fixture; observed execution HOLD")
    if type(grant["source_bytes"]) is not int or not 0 < grant["source_bytes"] <= MAX_SOURCE:
        raise ValueError("fixture source size invalid")
    life = Lifetime()
    try:
        life.open(source_file, private=True)
        pins = []
        for path in controls:
            fd = life.open(path, private=True, limit=1024 * 1024)
            pins.append((str(_absolute(path)), digest(life.read(fd, 1024 * 1024))))
        code = []
        for path in (Path(__file__), Path(__file__).with_name("observed_runner.py"),
                     Path(__file__).with_name("forecast_scorecard.py"),
                     Path(__file__).parents[1] / "forecast/models.py"):
            fd = life.open(path, limit=1024 * 1024)
            code.append((str(path), digest(life.read(fd, 1024 * 1024))))
        identities = tuple((str(p), identity, directory) for p, _, identity, directory in life.records)
        life.check()
    finally:
        life.close()
    binding = RootAdmissionBinding(grant["fixture_id"])
    from importlib.metadata import version
    runtime = {"python": sys.version.split()[0], "executable": str(Path(sys.executable).resolve()),
               "packages": {name: version(name) for name in ("openpyxl", "numpy", "pandas")}}
    _REGISTRY[binding] = MappingProxyType({"source": str(_absolute(source_file)),
                                         "raw": bytes(admission_raw), "grant": MappingProxyType(grant),
                                         "controls": tuple(pins), "code": tuple(code), "identities": identities,
                                         "runtime_raw": canonical(runtime), "acceptance_sha256": OFFLINE_ACCEPTANCE_SHA256})
    return binding


def register_owned_fixture_run(admission, protocol_path, manifest_path, output_dir):
    record = _registered(admission, _REGISTRY)
    paths = tuple(str(_absolute(p)) for p in (protocol_path, manifest_path))
    if paths != tuple(p for p, _ in record["controls"]):
        raise ValueError("foreign control paths")
    output = _absolute(output_dir)
    if output.parent != Path(record["source"]).parent.parent or output.name in {"inputs", "source"}:
        raise ValueError("output must be a distinct owned fixture-root child")
    if os.path.lexists(output):
        raise ValueError("output already has an entry")
    parent = os.lstat(output.parent)
    if not stat.S_ISDIR(parent.st_mode) or parent.st_uid != os.getuid() or parent.st_mode & 0o077:
        raise ValueError("output parent must be private")
    binding = RootObservedRunBinding(admission.fixture_id)
    _RUNS[binding] = MappingProxyType({"admission": admission, "protocol": paths[0], "manifest": paths[1],
                                     "output": str(output), "parent_identity": _identity(parent, True)})
    return binding


def fixture_run_record(binding):
    return _registered(binding, _RUNS)


def fixture_admission_record(binding):
    return _registered(binding, _REGISTRY)


def _begin_fixture_run(binding):
    _registered(binding, _RUNS)
    if binding in _RUN_LIVES:
        raise ValueError("fixture run already attempted; no marker reset")
    token = object()
    _RUN_LIVES[binding] = {"token": token, "completion_sha256": None, "failed": False,
                          "created_identity": None, "sealed_identity": None}
    return token


class _RunDirectoryAnchor(Lifetime):
    """Discrete run identity check; own directory FDs through verification/close.

    The inherited file reader is unchanged. This anchor opens directories only,
    without following any ancestor, and uses the same close-once/postclose guard.
    """

    def __init__(self, output, identities):
        super().__init__()
        output = _absolute(output)
        parent = None
        try:
            for index, part in enumerate(output.parts):
                at = Path("/") if index == 0 else at / part
                flags = os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW | os.O_NONBLOCK
                fd = os.open("/" if index == 0 else part, flags, dir_fd=parent)
                self.fds.append(fd)  # ownership before even the first fstat
                info = os.fstat(fd)
                if not stat.S_ISDIR(info.st_mode):
                    raise ValueError("run directory identity/type changed")
                identity = _identity(info, True)
                self.records.append((at, fd, identity, True))
                if at == output.parent and identity != identities[0]:
                    raise ValueError("run parent identity changed")
                if at == output and identity != identities[1]:
                    raise ValueError("run output identity changed")
                parent = fd
            self.check()
        except BaseException:
            self.close()
            raise

    def __enter__(self):
        try:
            self.check()
        except BaseException:
            self.close()
            raise
        return self

    def __exit__(self, *_):
        self.close()


def _capture_fixture_output(binding, token, created_identity):
    """Owning host only: capture the mkdir identity with its issued run token."""
    run = _registered(binding, _RUNS)
    state = _registered(binding, _RUN_LIVES)
    if (state["token"] is not token or state["failed"] or state["created_identity"] is not None
            or state["completion_sha256"] is not None):
        raise ValueError("foreign/stale host output token")
    identities = (run["parent_identity"], tuple(created_identity))
    with _RunDirectoryAnchor(run["output"], identities):
        pass
    state["created_identity"] = identities


def _seal_fixture_run(binding, token, completion_raw):
    run = _registered(binding, _RUNS)
    state = _registered(binding, _RUN_LIVES)
    if state["token"] is not token or state["failed"] or state["completion_sha256"] is not None:
        raise ValueError("foreign/stale host closure token")
    identities = state["created_identity"]
    if identities is None or identities[0] != run["parent_identity"]:
        raise ValueError("host-created output identity absent")
    with _RunDirectoryAnchor(run["output"], identities):
        pass
    # Issue both commitments only after the anchor's close/postclose succeeds.
    state.update(completion_sha256=digest(completion_raw), sealed_identity=identities)


def _fail_fixture_run(binding, token):
    state = _registered(binding, _RUN_LIVES)
    if state["token"] is not token:
        raise ValueError("foreign host token")
    state["failed"] = True


def _fixture_completion_sha(binding):
    state = _registered(binding, _RUN_LIVES)
    if state["failed"] or state["completion_sha256"] is None:
        raise ValueError("host-issued closure absent or failed")
    return state["completion_sha256"]


def _fixture_completion_identity(binding):
    run = _registered(binding, _RUNS)
    _fixture_completion_sha(binding)
    state = _registered(binding, _RUN_LIVES)
    identities = state["sealed_identity"]
    if (identities is None or identities != state["created_identity"]
            or identities[0] != run["parent_identity"]):
        raise ValueError("host-sealed output identity absent or stale")
    return identities


def _zip_guard(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        members = archive.infolist()
        if len(members) > 1024 or len({m.filename for m in members}) != len(members):
            raise ValueError("ZIP member count or duplicate")
        total = 0
        for member in members:
            path = PurePosixPath(member.filename)
            if path.is_absolute() or ".." in path.parts or "\\" in member.filename or member.flag_bits & 1:
                raise ValueError("ZIP traversal/encryption")
            if member.file_size > MAX_MEMBER or stat.S_ISLNK(member.external_attr >> 16):
                raise ValueError("ZIP member limit/type")
            actual = 0; chunks = []
            with archive.open(member) as stream:
                while True:
                    chunk = stream.read(65536)
                    if not chunk:
                        break
                    actual += len(chunk); total += len(chunk)
                    if actual > MAX_MEMBER or total > MAX_EXPANSION:
                        raise ValueError("ZIP expanded byte limit")
                    if member.filename.endswith((".xml", ".rels")):
                        chunks.append(chunk)
            if actual != member.file_size:
                raise ValueError("ZIP expansion metadata mismatch")
            if chunks:
                data = b"".join(chunks)
                markup = data.replace(b"\x00", b"").upper()  # include UTF16/32 markup
                if b"<!DOCTYPE" in markup or b"<!ENTITY" in markup:
                    raise ValueError("XML DTD/entity forbidden")
                root = ElementTree.fromstring(data)
                if member.filename.startswith("xl/worksheets/"):
                    xml_rows = [node for node in root.iter() if node.tag.rsplit("}", 1)[-1] == "row"]
                    cells = [node for node in root.iter() if node.tag.rsplit("}", 1)[-1] == "c"]
                    row_ids = [int(node.attrib.get("r", "0")) for node in xml_rows]
                    if (len(xml_rows) > 1006 or len(cells) > 100000 or row_ids != sorted(set(row_ids))
                            or any(index < 1 or index > 1006 for index in row_ids)):
                        raise ValueError("actual XML row/cell expansion bound")
                    seen, max_row, max_col = set(), 0, 0
                    for cell in cells:
                        ref = cell.attrib.get("r", "")
                        matched = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", ref)
                        if matched is None or ref in seen:
                            raise ValueError("actual XML duplicate/invalid cell")
                        seen.add(ref)
                        col = 0
                        for letter in matched[1]:
                            col = col * 26 + ord(letter) - 64
                        max_col, max_row = max(max_col, col), max(max_row, int(matched[2]))
                    if max_row > 1006 or max_row * max_col > 100000:
                        raise ValueError("actual XML cell expansion bound")
                if member.filename.endswith(".rels") and any(
                        node.attrib.get("TargetMode", "").lower() == "external" for node in root):
                    raise ValueError("external workbook relationship")


def _native(cell):
    value = cell.value
    if cell.data_type in {"f", "e"} or type(value) not in {int, float}:
        raise ValueError("required native numeric cell invalid")
    try:
        converted = float(value)
    except (OverflowError, ValueError) as exc:
        raise ValueError("native numeric overflow") from exc
    if not math.isfinite(converted) or converted < sys.float_info.min:
        raise ValueError("native numeric nonpositive/nonfinite/subnormal")
    if type(value) is int and Fraction(value) != Fraction.from_float(converted):
        raise ValueError("lossy integer conversion")
    return converted


def _parse_native(raw):
    from openpyxl import load_workbook
    _zip_guard(raw)
    stream = io.BytesIO(raw)
    book = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
    try:
        sheet = book["Monthly Prices"]
        if not sheet.max_row or not sheet.max_column or sheet.max_row * sheet.max_column > 100000:
            raise ValueError("sheet cell bound")
        if sheet.max_row - 6 > 1000:
            raise ValueError("sheet month-row bound")
        for _, (col, label, unit) in SERIES.items():
            if sheet.cell(5, col).value != label or sheet.cell(6, col).value != unit:
                raise ValueError("native header/unit mismatch")
        series = {key: {} for key in SERIES}
        labels = []
        cells_seen = 0
        for row in sheet.iter_rows(min_row=7, max_col=sheet.max_column):
            cells_seen += len(row)
            if cells_seen > 100000 or len(labels) >= 1000:
                raise ValueError("actual sheet expansion bound")
            label = row[0].value
            if label is None:
                raise ValueError("blank source month")
            expected_index = 1960 * 12 + len(labels)
            expected = f"{expected_index // 12:04d}M{expected_index % 12 + 1:02d}"
            if label != expected:
                raise ValueError("source month labels must be unique consecutive")
            labels.append(label)
            month = label[:4] + "-" + label[5:]
            if "2016-09" <= month <= "2026-08":
                for key, (col, _, _) in SERIES.items():
                    series[key][month] = _native(row[col - 1])
        if len(labels) != 801 or labels[-1] != "2026M09" or any(len(s) != 120 for s in series.values()):
            raise ValueError("exact 801 labels/120 model months required")
        return series
    finally:
        book.close()
        stream.close()


class ObservedNativeRead:
    def __init__(self, lifetime, series, receipt):
        self.lifetime = lifetime
        self.series = MappingProxyType({key: MappingProxyType(values) for key, values in series.items()})
        self.receipt = MappingProxyType(receipt)

    def check(self):
        self.lifetime.check()

    def close(self):
        self.lifetime.close()

    def __enter__(self):
        self.check()
        return self

    def __exit__(self, *_):
        self.close()


def open_observed_native_snapshot(source_file: Path, admission_raw: bytes, *, expected_admission: RootAdmissionBinding):
    record = _registered(expected_admission, _REGISTRY)  # before any body/metadata access
    if str(_absolute(source_file)) != record["source"] or admission_raw != record["raw"]:
        raise ValueError("foreign admission/source binding")
    for path, identity, directory in record["identities"]:
        if _identity(os.lstat(path), directory) != identity:
            raise ValueError("registered identity changed")
    life = Lifetime()
    try:
        for path, sha in record["controls"] + record["code"]:
            fd = life.open(path, private=(path, sha) in record["controls"], limit=1024 * 1024)
            if digest(life.read(fd, 1024 * 1024)) != sha:
                raise ValueError("control/code hash changed")
        fd = life.open(source_file, private=True)
        raw = life.read(fd, MAX_SOURCE)
        grant = record["grant"]
        if len(raw) != grant["source_bytes"] or digest(raw) != grant["source_sha256"]:
            raise ValueError("anchored source size/hash mismatch")
        values = _parse_native(raw)  # same bytes; never reopen the path
        life.check()
        basis = {key: {month: {"hex": value.hex(), "numerator": str(Fraction.from_float(value).numerator),
                              "denominator": str(Fraction.from_float(value).denominator)}
                       for month, value in series.items()} for key, series in values.items()}
        receipt = {**grant, "full_parent_snapshot_exposed": True, "physical_test_actuals_untouched": False,
                   "numeric_policy": "native-cell binary64; no XML decimal certificate",
                   "source_labels": 801, "required_months": 120, "stored_value_basis_sha256": digest(canonical(basis)),
                   "stored_value_basis": basis, "code_sha256": dict(record["code"]),
                   "control_sha256": dict(record["controls"]), "runtime": decode(record["runtime_raw"]),
                   "offline_acceptance_sha256": record["acceptance_sha256"]}
        return ObservedNativeRead(life, values, receipt)
    except BaseException:
        life.close()
        raise

# Native host configuration is private, in-process and deliberately nonpersistent.
# Identity registries are not hostile same-UID security or a signature service.
import threading
from datetime import datetime, timezone
from collections.abc import Mapping

_NATIVE_HOSTS = weakref.WeakKeyDictionary()
_NATIVE_RUNS = weakref.WeakKeyDictionary()
_NATIVE_LIVES = weakref.WeakKeyDictionary()
_NATIVE_ACTIVE = None
_NATIVE_LOCK = threading.RLock()
_NATIVE_PACKAGES = frozenset(('numpy', 'pandas', 'openpyxl', 'scikit-learn', 'statsmodels', 'lightgbm'))
_NATIVE_CODE = frozenset(('procurement-engine/src/procure/evaluation/observed_admission.py',
                         'procurement-engine/src/procure/evaluation/observed_runner.py',
                         'procurement-engine/src/procure/evaluation/forecast_scorecard.py',
                         'procurement-engine/src/procure/forecast/models.py'))
_NATIVE_TESTS = frozenset(('procurement-engine/tests/test_forecast_observed_admission.py',
                          'procurement-engine/tests/test_observed_forecast_runner.py'))
NATIVE_BUDGETS = {'fit_seconds': 60, 'per_series_seconds': 900, 'total_seconds': 1800,
                  'max_attempts': 288, 'validation_max_attempts': 144, 'test_max_attempts': 144,
                  'cleanup_reserve_seconds': 3, 'finalizer_reserve_seconds': 30,
                  'per_model_seconds': {'naive': 120, 'ets': 360, 'gbm': 360},
                  'validation_model_seconds': {'naive': 48, 'ets': 144, 'gbm': 144},
                  'test_model_seconds': {'naive': 72, 'ets': 216, 'gbm': 216}, 'retries': 0}
NATIVE_CALENDAR = {'as_of': '2026-10-08', 'training_start': '2016-09', 'validation_start': '2021-09',
                   'validation_end': '2023-08', 'purge_start': '2023-09', 'purge_end': '2024-02',
                   'test_start': '2024-03', 'test_end': '2026-02', 'latest_eligible_actual': '2026-08',
                   'source_months': 801, 'selected_months': 120, 'validation_origins': 24,
                   'purged_origins': 6, 'test_origins': 24, 'minimum_training_months': 60,
                   'training_cutoff': 'origin_minus_1', 'horizons': [1, 3, 6], 'model_leads': [2, 4, 7],
                   'fit_steps': 7, 'gbm_refit_every': 1, 'base_seed': 42, 'n_paths': 1000,
                   'bootstrap_rng': 'numpy.PCG64', 'quantile_method': 'linear',
                   'calibration_grid': [.8, .9, 1., 1.1, 1.2, 1.35, 1.5],
                   'selection_relative_improvement': .02}


def _native_keys(value, keys):
    if not isinstance(value, Mapping) or set(value) != set(keys):
        raise ValueError('native exact schema keys required')


def _native_sha(value, length=64):
    if not isinstance(value, str) or re.fullmatch('[0-9a-f]{%d}' % length, value) is None:
        raise ValueError('native hash/revision invalid')


def _native_id(value):
    if not isinstance(value, str) or re.fullmatch('[A-Za-z0-9][A-Za-z0-9_-]{0,63}', value) is None:
        raise ValueError('native identifier invalid')


def _native_path(value):
    if not isinstance(value, str) or str(_absolute(value)) != value:
        raise ValueError('native canonical path required')
    return Path(value)


def _native_utc(value):
    if not isinstance(value, str) or re.fullmatch(r'\d{4}-\d\d-\d\dT\d\d:\d\d:\d\d\.\d{6}Z', value) is None:
        raise ValueError('native canonical UTC required')
    return datetime.strptime(value, '%Y-%m-%dT%H:%M:%S.%fZ').replace(tzinfo=timezone.utc)


def _native_equal(actual, expected):
    # JSON bool/int and int/float equality must not weaken the frozen protocol.
    if type(actual) is not type(expected):
        raise ValueError('native frozen field type drift')
    if isinstance(expected, dict):
        _native_keys(actual, expected)
        for key in expected:
            _native_equal(actual[key], expected[key])
    elif isinstance(expected, list):
        if len(actual) != len(expected):
            raise ValueError('native frozen list drift')
        for a, b in zip(actual, expected):
            _native_equal(a, b)
    elif actual != expected:
        raise ValueError('native frozen field drift')


def _native_identity_doc(value, directory=False, link=False):
    keys = ('dev', 'ino', 'mode', 'uid') if directory else ('dev', 'ino', 'mode', 'uid', 'nlink', 'size', 'mtime_ns', 'ctime_ns')
    _native_keys(value, keys)
    if any(type(value[k]) is not int or value[k] < 0 for k in keys):
        raise ValueError('native identity type invalid')
    if directory:
        valid = stat.S_ISDIR(value['mode'])
    elif link:
        valid = stat.S_ISLNK(value['mode'])
    else:
        valid = stat.S_ISREG(value['mode']) and value['nlink'] == 1
    if not valid:
        raise ValueError('native identity file type/link invalid')
    return tuple(value[k] for k in keys)


def _native_ancestors(path, ancestors):
    if not isinstance(ancestors, (list, tuple)) or not 1 <= len(ancestors) <= 32:
        raise ValueError('native ancestor bound')
    expected = list(reversed(path.parents))
    if [x.get('path') for x in ancestors] != [str(p) for p in expected]:
        raise ValueError('native ancestor chain mismatch')
    for item in ancestors:
        _native_keys(item, ('path', 'identity'))
        _native_identity_doc(item['identity'], True)


def _native_pin_schema(pin):
    _native_keys(pin, ('path', 'bytes', 'sha256', 'identity', 'ancestors'))
    path = _native_path(pin['path'])
    if type(pin['bytes']) is not int or not 0 < pin['bytes'] <= 16 * 1024 * 1024:
        raise ValueError('native file size bound')
    _native_sha(pin['sha256'])
    identity = _native_identity_doc(pin['identity'])
    if identity[5] != pin['bytes']:
        raise ValueError('native size identity mismatch')
    _native_ancestors(path, pin['ancestors'])


class _NativeLifetime(Lifetime):
    """Native-only prefix sharing; every FD is owned before first fstat."""
    def __init__(self):
        super().__init__()
        self._dirs = {}

    def open(self, path, *, private=False, limit=MAX_SOURCE):
        if self.closed:
            raise ValueError('native lifetime closed')
        path = _absolute(path)
        parent = None
        for index, part in enumerate(path.parts):
            at = Path('/') if index == 0 else at / part
            directory = index < len(path.parts) - 1
            if directory and at in self._dirs:
                fd = self._dirs[at]
                self.check()
            else:
                if len(self.fds) >= 128:
                    raise ValueError('native descriptor capacity unavailable')
                flags = os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK | (os.O_DIRECTORY if directory else 0)
                fd = os.open('/' if index == 0 else part, flags, dir_fd=parent)
                self.fds.append(fd)
                info = os.fstat(fd)
                if directory and not stat.S_ISDIR(info.st_mode):
                    raise ValueError('native ancestor type changed')
                if not directory and (not stat.S_ISREG(info.st_mode) or info.st_nlink != 1 or info.st_size > limit):
                    raise ValueError('native bounded single-link file required')
                if not directory and private:
                    p = os.fstat(parent)
                    if info.st_uid != os.getuid() or p.st_uid != os.getuid() or p.st_mode & 0o077:
                        raise ValueError('native private input parent required')
                self.records.append((at, fd, _identity(info, directory), directory))
                if directory:
                    self._dirs[at] = fd
            parent = fd
        return fd


def _native_pin_open(life, pin, *, private=False, body=True):
    fd = life.open(pin['path'], private=private, limit=pin['bytes'])
    if _identity(os.fstat(fd)) != _native_identity_doc(pin['identity']):
        raise ValueError('native pinned file identity changed')
    for item in pin['ancestors']:
        if _identity(os.lstat(item['path']), True) != _native_identity_doc(item['identity'], True):
            raise ValueError('native pinned ancestor changed')
    if body:
        raw = life.read(fd, pin['bytes'])
        if len(raw) != pin['bytes'] or digest(raw) != pin['sha256']:
            raise ValueError('native pinned file hash changed')
        return raw
    return fd


def _native_freeze(value):
    if isinstance(value, Mapping):
        return MappingProxyType({k: _native_freeze(v) for k, v in value.items()})
    if isinstance(value, (list, tuple)):
        return tuple(_native_freeze(v) for v in value)
    return value


def _native_plain(value):
    if isinstance(value, Mapping):
        return {k: _native_plain(v) for k, v in value.items()}
    if isinstance(value, tuple):
        return [_native_plain(v) for v in value]
    return value


@dataclass(frozen=True, eq=False)
class RootNativeHostAuthority:
    acceptance_id: str


@dataclass(frozen=True, eq=False)
class RootNativeRunBinding:
    run_id: str


def _issue_root_native_host_authority(*, acceptance_id, accepted_control, integrated_head, integrated_tree, profile):
    """ROOT private bootstrap only, once per parent; never called by products."""
    global _NATIVE_ACTIVE
    _native_id(acceptance_id)
    _native_pin_schema(accepted_control)
    _native_sha(integrated_head, 40); _native_sha(integrated_tree, 40)
    if profile not in ('native_observed', 'native_fixture_test'):
        raise ValueError('native host profile required')
    with _NATIVE_LOCK:
        if _NATIVE_ACTIVE is not None:
            raise ValueError('native issuer already configured; no reset/rotation')
        life = _NativeLifetime()
        try:
            _native_pin_open(life, accepted_control, private=True)
        finally:
            life.close()
        authority = RootNativeHostAuthority(acceptance_id)
        record = {'pin': _native_freeze(accepted_control), 'head': integrated_head, 'tree': integrated_tree,
                  'profile': profile, 'epoch': object(), 'used_controls': set(), 'used_ids': set(), 'used_outputs': set()}
        _NATIVE_HOSTS[authority] = record
        _NATIVE_ACTIVE = authority
        return authority


def _native_host(authority):
    if authority is not _NATIVE_ACTIVE:
        raise ValueError('issued active native host authority required')
    return _registered(authority, _NATIVE_HOSTS)


def _native_control_schema(c, host, authority):
    required = ('schema_version', 'status', 'profile', 'acceptance_id', 'run_id', 'integrated_head', 'integrated_tree',
                'integration_receipt_sha256', 'design_acceptance_sha256', 'numerical_review_sha256', 'runtime_review_sha256',
                'issued_at_utc', 'expires_at_utc', 'analysis_as_of_utc', 'source', 'controls', 'code', 'test_artifact_blobs',
                'runtime', 'calendar', 'outputs', 'budgets', 'permissions')
    _native_keys(c, required)
    observed = host['profile'] == 'native_observed'
    status = 'ROOT_ACCEPTED_NATIVE_OBSERVED_RUN' if observed else 'ROOT_ACCEPTED_NATIVE_FIXTURE_TEST'
    if (c['schema_version'], c['status'], c['profile'], c['acceptance_id'], c['integrated_head'], c['integrated_tree']) != (
            'root-native-run-control.v1', status, host['profile'], authority.acceptance_id, host['head'], host['tree']):
        raise ValueError('native authority/control/profile mismatch')
    _native_id(c['run_id'])
    for key in ('integration_receipt_sha256', 'design_acceptance_sha256', 'numerical_review_sha256', 'runtime_review_sha256'):
        _native_sha(c[key])
    issued, expiry, asof = map(_native_utc, (c['issued_at_utc'], c['expires_at_utc'], c['analysis_as_of_utc']))
    if not issued <= datetime.now(timezone.utc) < expiry or issued >= expiry or asof.date().isoformat() != '2026-10-08':
        raise ValueError('native expiry/as-of refused')
    _native_equal(c['calendar'], NATIVE_CALENDAR); _native_equal(c['budgets'], NATIVE_BUDGETS)
    _native_equal(c['permissions'], {'local_source_body_read': True, 'local_statistical_execution': observed,
                                   'observed_execution': observed, 'synthetic_test_execution': not observed,
                                   'network': False, 'provider': False, 'source_download': False, 'sql': False})
    source = c['source']
    _native_keys(source, ('file', 'source_id', 'source_url', 'retrieved_utc', 'rights_disposition_sha256', 'sheet',
                         'month_range', 'series', 'representation', 'full_parent_snapshot_exposed',
                         'physical_test_actuals_untouched', 'vintage_status', 'availability_proxy'))
    _native_pin_schema(source['file'])
    _native_sha(source['rights_disposition_sha256'])
    if source['file']['bytes'] > MAX_SOURCE or _native_utc(source['retrieved_utc']) > asof:
        raise ValueError('native source size/retrieval invalid')
    if not isinstance(source['source_id'], str) or not isinstance(source['source_url'], str):
        raise ValueError('native source provenance type invalid')
    _native_equal({k: source[k] for k in ('sheet', 'month_range', 'representation', 'full_parent_snapshot_exposed',
                                         'physical_test_actuals_untouched', 'vintage_status', 'availability_proxy')},
                  {'sheet': 'Monthly Prices', 'month_range': '1960M01-2026M09', 'representation': 'openpyxl_native_binary64',
                   'full_parent_snapshot_exposed': True, 'physical_test_actuals_untouched': False,
                   'vintage_status': 'latest_vintage_only_historical_publication_unknown',
                   'availability_proxy': 'month_m_available_at_end_of_m_plus_1; not historical publication proof'})
    _native_equal(source['series'], {
        'copper': {'column': 'BM', 'index': 65, 'label': 'Copper', 'native_unit': '($/mt)', 'unit': 'USD/metric tonne', 'currency': 'USD'},
        'aluminum': {'column': 'BK', 'index': 63, 'label': 'Aluminum', 'native_unit': '($/mt)', 'unit': 'USD/metric tonne', 'currency': 'USD'}})
    if observed:
        if (source['source_id'] != 'worldbank_pinksheet_monthly_oct2026' or source['file']['sha256'] != _HELD_OBSERVED_SHA256
                or source['file']['bytes'] != 586099
                or source['rights_disposition_sha256'] != 'f34d8348dc638969f61b629c77e20a277b108a169848844cc2b944ee7a37c1ee'
                or source['retrieved_utc'] != '2026-10-08T07:44:34.000000Z'
                or source['source_url'] != 'https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Monthly.xlsx'):
            raise ValueError('native observed source/rights pin mismatch')
    elif (source['file']['sha256'] == _HELD_OBSERVED_SHA256 or source['source_id'] != 'owned_fixture'
          or Path(source['file']['path']).name == 'worldbank_pinksheet_monthly.xlsx'):
        raise ValueError('observed original cannot become native test fixture')
    controls = c['controls']
    _native_keys(controls, ('protocol', 'manifest', 'root_acceptance_path'))
    for key in ('protocol', 'manifest'):
        _native_pin_schema(controls[key])
    if controls['root_acceptance_path'] != host['pin']['path']:
        raise ValueError('native external control path mismatch')
    if not isinstance(c['code'], list) or len(c['code']) != 4 or {p.get('repository_path') for p in c['code']} != _NATIVE_CODE:
        raise ValueError('native exact code pins required')
    repo = Path(__file__).resolve().parents[4]
    for item in c['code']:
        _native_keys(item, ('repository_path', 'file', 'git_blob'))
        _native_sha(item['git_blob'], 40); _native_pin_schema(item['file'])
        if item['file']['path'] != str(repo / item['repository_path']):
            raise ValueError('native foreign code path')
    _native_keys(c['test_artifact_blobs'], _NATIVE_TESTS)
    for blob in c['test_artifact_blobs'].values():
        _native_sha(blob, 40)
    runtime = c['runtime']
    _native_keys(runtime, ('executable_alias', 'resolved_executable', 'alias_chain', 'python_version', 'packages', 'model_backends'))
    _native_path(runtime['executable_alias']); _native_pin_schema(runtime['resolved_executable'])
    if runtime['executable_alias'] != sys.executable or runtime['resolved_executable']['path'] != str(Path(sys.executable).resolve()):
        raise ValueError('native interpreter path mismatch')
    if runtime['python_version'] != sys.version.split()[0]:
        raise ValueError('native Python version mismatch')
    _native_equal(runtime['model_backends'], {'naive': 'naive', 'ets': 'ets', 'gbm': 'sklearn'})
    if not isinstance(runtime['alias_chain'], list) or len(runtime['alias_chain']) > 8:
        raise ValueError('native alias bound')
    for item in runtime['alias_chain']:
        _native_keys(item, ('path', 'target', 'identity', 'ancestors'))
        _native_path(item['path']); _native_identity_doc(item['identity'], link=True)
        _native_ancestors(Path(item['path']), item['ancestors'])
        if not isinstance(item['target'], str) or not item['target']:
            raise ValueError('native alias target invalid')
    if not isinstance(runtime['packages'], list) or len(runtime['packages']) != 6 or {p.get('distribution') for p in runtime['packages']} != _NATIVE_PACKAGES:
        raise ValueError('native six package pins required')
    for item in runtime['packages']:
        _native_keys(item, ('distribution', 'version', 'metadata', 'module_origin'))
        if not isinstance(item['version'], str) or not 0 < len(item['version']) <= 128 or any(ord(x) < 32 for x in item['version']):
            raise ValueError('native package version invalid')
        _native_pin_schema(item['metadata']); _native_path(item['module_origin'])
    outputs = c['outputs']
    _native_keys(outputs, ('path', 'parent', 'parent_identity', 'owner_uid', 'create_mode', 'absent_required', 'overwrite'))
    output = _native_path(outputs['path'])
    if str(output.parent) != outputs['parent'] or outputs['owner_uid'] != os.getuid() or type(outputs['owner_uid']) is not int:
        raise ValueError('native output parent/owner mismatch')
    parent = _native_identity_doc(outputs['parent_identity'], True)
    if parent[3] != os.getuid() or parent[2] & 0o077:
        raise ValueError('native output parent must be private')
    _native_equal({k: outputs[k] for k in ('create_mode', 'absent_required', 'overwrite')},
                  {'create_mode': 448, 'absent_required': True, 'overwrite': False})
    inputs = [Path(source['file']['path']), Path(host['pin']['path']), Path(controls['protocol']['path']), Path(controls['manifest']['path'])]
    if any(output == p or output in p.parents or output.parent == p.parent for p in inputs) or len(set(inputs)) != len(inputs):
        raise ValueError('native distinct input/output paths required')


def _native_alias_guard(runtime):
    current = Path(sys.executable)
    observed = []
    while current.is_symlink():
        if len(observed) >= 8:
            raise ValueError('native executable alias cycle')
        observed.append((str(current), os.readlink(current), _identity(os.lstat(current))))
        target = Path(os.readlink(current))
        next_path = target if target.is_absolute() else current.parent / target
        current = next_path.parent.resolve() / next_path.name
    expected = [(x['path'], x['target'], _native_identity_doc(x['identity'], link=True)) for x in runtime['alias_chain']]
    if observed != expected or str(current.resolve()) != runtime['resolved_executable']['path']:
        raise ValueError('native executable alias changed')
    for item in runtime['alias_chain']:
        for ancestor in item['ancestors']:
            if _identity(os.lstat(ancestor['path']), True) != _native_identity_doc(ancestor['identity'], True):
                raise ValueError('native executable ancestor changed')


def _native_pins_lifetime(run):
    if run['control']['profile'] in ('generic-observed-v1','generated-fixture-test-v1'):
        return _generic_pins_lifetime(run)
    c = run['control']; life = _NativeLifetime()
    try:
        _native_alias_guard(c['runtime'])
        _native_pin_open(life, run['root_pin'], private=True)
        p_raw = _native_pin_open(life, c['controls']['protocol'], private=True)
        m_raw = _native_pin_open(life, c['controls']['manifest'], private=True)
        for item in c['code']:
            _native_pin_open(life, item['file'])
        _native_pin_open(life, c['runtime']['resolved_executable'])
        from importlib.machinery import PathFinder
        from email.parser import BytesParser
        for item in c['runtime']['packages']:
            raw = _native_pin_open(life, item['metadata'])
            if BytesParser().parsebytes(raw).get('Version') != item['version']:
                raise ValueError('native package metadata version drift')
            module = {'scikit-learn': 'sklearn'}.get(item['distribution'], item['distribution'])
            # Top-level find_spec resolves an origin without importing the package.
            spec = PathFinder.find_spec(module)
            if spec is None or spec.origin != item['module_origin']:
                raise ValueError('native module origin drift')
        _native_pin_open(life, c['source']['file'], private=True, body=False)
        p, m = decode(p_raw), decode(m_raw)
        from .observed_runner import _check_protocol
        observed = c['profile'] == 'native_observed'
        _check_protocol({**p, 'protocol_sha256': digest(p_raw), 'source_sha256': c['source']['file']['sha256'],
                         'snapshot_sha256': c['source']['file']['sha256']}, '2026-10-08')
        if (p['observed_execution_authorized'] is not observed or m.get('observed_execution_authorized') is not observed
                or m.get('is_synthetic') is not (not observed) or m.get('observed_authenticity') is not observed
                or m.get('source_id') != c['source']['source_id'] or m.get('sha256') != c['source']['file']['sha256']
                or type(m.get('bytes')) is not int or m['bytes'] != c['source']['file']['bytes']):
            raise ValueError('native control body/source/profile mismatch')
        life.check()
        return life
    except BaseException:
        life.close()
        raise


def register_root_accepted_native_run(authority, root_acceptance_raw, *, source_file, protocol_path, manifest_path, output_dir):
    host = _native_host(authority)  # no source metadata or body before identity
    if not isinstance(root_acceptance_raw, bytes) or len(root_acceptance_raw) > 128 * 1024 or digest(root_acceptance_raw) != host['pin']['sha256']:
        raise ValueError('native external control SHA mismatch')
    c = decode(root_acceptance_raw)
    _native_control_schema(c, host, authority)
    if tuple(map(str, (source_file, protocol_path, manifest_path, output_dir))) != (
            c['source']['file']['path'], c['controls']['protocol']['path'], c['controls']['manifest']['path'], c['outputs']['path']):
        raise ValueError('native caller paths mismatch')
    with _NATIVE_LOCK:
        if (host['pin']['sha256'] in host['used_controls'] or c['run_id'] in host['used_ids'] or c['outputs']['path'] in host['used_outputs']):
            raise ValueError('native reservation consumed; no reset')
        host['used_controls'].add(host['pin']['sha256']); host['used_ids'].add(c['run_id']); host['used_outputs'].add(c['outputs']['path'])
    run = _native_freeze({'authority': authority, 'epoch': host['epoch'], 'control': c, 'root_pin': host['pin'],
                          'output': c['outputs']['path'], 'parent_identity': _native_identity_doc(c['outputs']['parent_identity'], True),
                          'raw': root_acceptance_raw})
    life = _native_pins_lifetime(run)
    try:
        if os.path.lexists(run['output']) or _identity(os.lstat(Path(run['output']).parent), True) != run['parent_identity']:
            raise ValueError('native absent output/private parent changed')
    finally:
        life.close()
    binding = RootNativeRunBinding(c['run_id'])
    _NATIVE_RUNS[binding] = run
    _NATIVE_LIVES[binding] = {'state': 'READY', 'token': None, 'created': None, 'sealed': None, 'closure_sha': None}
    return binding


def native_run_record(binding):
    run = _registered(binding, _NATIVE_RUNS)
    host = _native_host(run['authority'])
    if run['epoch'] is not host['epoch']:
        raise ValueError('native foreign host epoch')
    return run


def _begin_native_run(binding):
    run = native_run_record(binding)
    with _NATIVE_LOCK:
        state = _registered(binding, _NATIVE_LIVES)
        if state['state'] != 'READY' or datetime.now(timezone.utc) >= _native_utc(run['control']['expires_at_utc']):
            raise ValueError('native attempted/expired run; no reset')
        state.update(state='RUNNING', token=object())
        return state['token']


def _native_state(binding, token=None):
    native_run_record(binding)
    state = _registered(binding, _NATIVE_LIVES)
    if token is not None and state['token'] is not token:
        raise ValueError('native foreign execution token')
    return state


def _capture_native_output(binding, token, identity):
    run = native_run_record(binding); state = _native_state(binding, token)
    if state['state'] != 'RUNNING' or state['created'] is not None:
        raise ValueError('native output capture already consumed')
    ids = (run['parent_identity'], tuple(identity))
    with _RunDirectoryAnchor(run['output'], ids):
        pass
    state['created'] = ids


def _seal_native_run(binding, token, raw):
    run = native_run_record(binding); state = _native_state(binding, token)
    if state['state'] != 'RUNNING' or state['created'] is None:
        raise ValueError('native completion token unavailable')
    with _RunDirectoryAnchor(run['output'], state['created']):
        if os.path.lexists(Path(run['output']) / 'FAILURE.json'):
            raise ValueError('native failure entry dominates')
    state.update(state='SEALED', sealed=state['created'], closure_sha=digest(raw))


def _fail_native_run(binding, token):
    _native_state(binding, token)['state'] = 'FAILED'


def _native_completion(binding):
    state = _native_state(binding)
    if state['state'] != 'SEALED' or not state['sealed'] or not state['closure_sha']:
        raise ValueError('native issued completion absent/failed')
    return state['sealed'], state['closure_sha']


class _NativeRead(ObservedNativeRead):
    def __enter__(self):
        try:
            self.check()
        except BaseException:
            self.close()
            raise
        return self


def open_root_native_snapshot(binding):
    run = native_run_record(binding)
    if _native_state(binding)['state'] != 'RUNNING':
        raise ValueError('native running host token required before source body')
    life = _native_pins_lifetime(run)
    try:
        c = run['control']; pin = c['source']['file']
        raw = _native_pin_open(life, pin, private=True)  # one source-body read after all pins
        values = _parse_native(raw)
        life.check()
        basis = {sid: {month: {'hex': value.hex(), 'numerator': str(Fraction.from_float(value).numerator),
                              'denominator': str(Fraction.from_float(value).denominator)} for month, value in series.items()}
                 for sid, series in values.items()}
        observed = c['profile'] == 'native_observed'
        receipt = {'source_sha256': pin['sha256'], 'source_bytes': pin['bytes'], 'source_id': c['source']['source_id'],
                   'is_synthetic': not observed, 'observed_authenticity': observed, 'observed_execution_authorized': observed,
                   'full_parent_snapshot_exposed': True, 'physical_test_actuals_untouched': False,
                   'numeric_policy': 'native-cell binary64; no XML decimal certificate',
                   'source_labels': 801, 'required_months': 120, 'stored_value_basis': basis,
                   'stored_value_basis_sha256': digest(canonical(basis)), 'root_control_sha256': run['root_pin']['sha256']}
        return _NativeRead(life, values, receipt)
    except BaseException:
        life.close()
        raise

# Generic profiles are explicit host identities, separate from legacy/narrow runs.
GENERIC_CALENDAR = {k:v for k,v in NATIVE_CALENDAR.items() if k != 'source_months'}
GENERIC_CALENDAR = {**GENERIC_CALENDAR, 'as_of':'2026-10-09'}
_GENERIC_CODE = _NATIVE_CODE | {'procurement-engine/src/procure/evaluation/series_catalog.py'}
_GENERIC_TESTS = _NATIVE_TESTS | {'procurement-engine/tests/test_forecast_series_catalog.py'}
_GENERIC_ORIGINALS = {'wb':('ea1c350827878ea3bbe30e3cda16a13fd3bd5b409b8458940dc94a36b5a33154',586099),
                      'bis':('c3abed13e05392664999b363703ef40e6ac6f2851b2c3e5a0deabdd70f462989',22714802)}


@dataclass(frozen=True, eq=False)
class RootGenericRunBinding(RootNativeRunBinding):
    """Constructing/copying this object never registers authority."""


def _issue_root_generic_host_authority(*, acceptance_id, accepted_control, integrated_head, integrated_tree, profile):
    """Private ROOT bootstrap, one shared active host epoch per process."""
    global _NATIVE_ACTIVE
    from .series_catalog import PROFILES
    _native_id(acceptance_id);_native_pin_schema(accepted_control)
    _native_sha(integrated_head,40);_native_sha(integrated_tree,40)
    if profile not in PROFILES:raise ValueError('generic issued profile required')
    with _NATIVE_LOCK:
        if _NATIVE_ACTIVE is not None:raise ValueError('native issuer already configured; no reset/rotation')
        life=_NativeLifetime()
        try:_native_pin_open(life,accepted_control,private=True)
        finally:life.close()
        authority=RootNativeHostAuthority(acceptance_id)
        _NATIVE_HOSTS[authority]={'pin':_native_freeze(accepted_control),'head':integrated_head,'tree':integrated_tree,
                                'profile':profile,'epoch':object(),'used_controls':set(),'used_ids':set(),'used_outputs':set()}
        _NATIVE_ACTIVE=authority
        return authority


def _generic_source_pin(pin,limit):
    # New source-specific limit; the legacy 12MiB/16MiB pin rules are untouched.
    _native_keys(pin,('path','bytes','sha256','identity','ancestors'))
    path=_native_path(pin['path']);_native_sha(pin['sha256'])
    if type(pin['bytes']) is not int or not 0<pin['bytes']<=limit:raise ValueError('generic source size bound')
    if _native_identity_doc(pin['identity'])[5]!=pin['bytes']:raise ValueError('generic size identity mismatch')
    _native_ancestors(path,pin['ancestors'])


def _generic_runtime_schema(runtime):
    _native_keys(runtime,('executable_alias','resolved_executable','alias_chain','python_version','packages','model_backends'))
    _native_path(runtime['executable_alias']);_native_pin_schema(runtime['resolved_executable'])
    if (runtime['executable_alias']!=sys.executable or runtime['resolved_executable']['path']!=str(Path(sys.executable).resolve())
            or runtime['python_version']!=sys.version.split()[0]):raise ValueError('generic interpreter pin mismatch')
    _native_equal(runtime['model_backends'],{'naive':'naive','ets':'ets','gbm':'sklearn'})
    if not isinstance(runtime['alias_chain'],list) or len(runtime['alias_chain'])>8:raise ValueError('generic alias bound')
    for item in runtime['alias_chain']:
        _native_keys(item,('path','target','identity','ancestors'))
        path=_native_path(item['path']);_native_identity_doc(item['identity'],link=True);_native_ancestors(path,item['ancestors'])
        if not isinstance(item['target'],str) or not item['target']:raise ValueError('generic alias target invalid')
    if not isinstance(runtime['packages'],list) or len(runtime['packages'])!=6 or {p.get('distribution') for p in runtime['packages']}!=_NATIVE_PACKAGES:
        raise ValueError('generic six package pins required')
    for item in runtime['packages']:
        _native_keys(item,('distribution','version','metadata','module_origin'))
        if not isinstance(item['version'],str) or not 0<len(item['version'])<=128 or any(ord(c)<32 for c in item['version']):
            raise ValueError('generic version invalid')
        _native_pin_schema(item['metadata']);_native_path(item['module_origin'])


def _generic_control_schema(c,host,authority):
    from .series_catalog import MAX_BIS_BYTES,MAX_WB_BYTES
    required=('schema_version','status','profile','acceptance_id','run_id','integrated_head','integrated_tree',
              'integration_receipt_sha256','design_acceptance_sha256','numerical_review_sha256','runtime_review_sha256',
              'issued_at_utc','expires_at_utc','analysis_as_of_utc','sources','catalog','selected_ids','batch_index',
              'regime_dispositions','controls','code','test_artifact_blobs','runtime','calendar','outputs','budgets','permissions')
    _native_keys(c,required);observed=host['profile']=='generic-observed-v1'
    if (c['schema_version'],c['status'],c['profile'],c['acceptance_id'],c['integrated_head'],c['integrated_tree'])!=(
        'root-generic-run-control.v1','ROOT_ACCEPTED_GENERIC_OBSERVED_RUN' if observed else 'ROOT_ACCEPTED_GENERIC_FIXTURE_TEST',
        host['profile'],authority.acceptance_id,host['head'],host['tree']):raise ValueError('generic authority/control mismatch')
    _native_id(c['run_id'])
    for k in ('integration_receipt_sha256','design_acceptance_sha256','numerical_review_sha256','runtime_review_sha256'):_native_sha(c[k])
    issued,expiry,asof=map(_native_utc,(c['issued_at_utc'],c['expires_at_utc'],c['analysis_as_of_utc']))
    if not issued<=datetime.now(timezone.utc)<expiry or issued>=expiry or asof.date().isoformat()!='2026-10-09':
        raise ValueError('generic expiry/as-of refused')
    _native_equal(c['calendar'],GENERIC_CALENDAR);_native_equal(c['budgets'],NATIVE_BUDGETS)
    _native_equal(c['permissions'],{'local_source_body_read':True,'local_statistical_execution':observed,
                  'observed_execution':observed,'synthetic_test_execution':not observed,'network':False,'provider':False,
                  'source_download':False,'sql':False})
    _native_keys(c['sources'],('wb','bis'))
    for key,limit in [('wb',MAX_WB_BYTES),('bis',MAX_BIS_BYTES)]:
        source=c['sources'][key]
        _native_keys(source,('file','provider','source_id','source_url','retrieved_utc','rights_disposition_sha256','source_disposition_sha256'))
        _generic_source_pin(source['file'],limit)
        for k in ('rights_disposition_sha256','source_disposition_sha256'):_native_sha(source[k])
        if _native_utc(source['retrieved_utc'])>asof:raise ValueError('generic source retrieval after as-of')
        for k in ('provider','source_id','source_url'):
            if not isinstance(source[k],str) or not 0<len(source[k])<=512:raise ValueError('generic source field bound')
        if source['provider']!=('World Bank' if key=='wb' else 'BIS'):raise ValueError('generic provider mismatch')
        original=_GENERIC_ORIGINALS[key]
        if observed:
            if (source['file']['sha256'],source['file']['bytes'])!=original:raise ValueError('generic observed original pin mismatch')
            if source['rights_disposition_sha256']!='f34d8348dc638969f61b629c77e20a277b108a169848844cc2b944ee7a37c1ee':
                raise ValueError('generic published rights disposition mismatch')
        elif source['file']['sha256'] in {v[0] for v in _GENERIC_ORIGINALS.values()} or Path(source['file']['path']).name in {
                'worldbank_pinksheet_monthly.xlsx','bis-xru-monthly-averages.csv'}:
            raise ValueError('observed original cannot become generic fixture')
    _native_pin_schema(c['catalog'])
    if c['catalog']['bytes']>8*1024*1024:raise ValueError('generic catalog byte bound')
    ids=c['selected_ids']
    if not isinstance(ids,list) or not 1<=len(ids)<=2 or any(not isinstance(s,str) or not 0<len(s)<=512 for s in ids) or ids!=sorted(set(ids)):
        raise ValueError('generic selected ordered initial batch bound')
    if type(c['batch_index']) is not int or c['batch_index']<0 or not isinstance(c['regime_dispositions'],dict):
        raise ValueError('generic finite batch/regime shape invalid')
    controls=c['controls'];_native_keys(controls,('protocol','manifest','root_acceptance_path'))
    for k in ('protocol','manifest'):_native_pin_schema(controls[k])
    if controls['root_acceptance_path']!=host['pin']['path']:raise ValueError('generic root control path mismatch')
    if not isinstance(c['code'],list) or len(c['code'])!=5 or {p.get('repository_path') for p in c['code']}!=_GENERIC_CODE:
        raise ValueError('generic exact five code pins required')
    repo=Path(__file__).resolve().parents[4]
    for item in c['code']:
        _native_keys(item,('repository_path','file','git_blob'));_native_sha(item['git_blob'],40);_native_pin_schema(item['file'])
        if item['file']['path']!=str(repo/item['repository_path']):raise ValueError('generic foreign code path')
    _native_keys(c['test_artifact_blobs'],_GENERIC_TESTS)
    for b in c['test_artifact_blobs'].values():_native_sha(b,40)
    _generic_runtime_schema(c['runtime'])
    out=c['outputs'];_native_keys(out,('path','parent','parent_identity','owner_uid','create_mode','absent_required','overwrite'))
    path=_native_path(out['path']);parent=_native_identity_doc(out['parent_identity'],True)
    if str(path.parent)!=out['parent'] or type(out['owner_uid']) is not int or out['owner_uid']!=os.getuid() or parent[3]!=os.getuid() or parent[2]&0o077:
        raise ValueError('generic private output parent/owner mismatch')
    _native_equal({k:out[k] for k in ('create_mode','absent_required','overwrite')},{'create_mode':448,'absent_required':True,'overwrite':False})
    inputs=[Path(s['file']['path'])for s in c['sources'].values()]+[Path(c['catalog']['path']),Path(host['pin']['path'])]+[
        Path(controls[k]['path'])for k in ('protocol','manifest')]
    if len(set(inputs))!=len(inputs) or any(path==p or path in p.parents or path.parent==p.parent for p in inputs):
        raise ValueError('generic distinct inputs/output required')


def _generic_pins_lifetime(run):
    c=run['control'];life=_NativeLifetime()
    try:
        _native_alias_guard(c['runtime']);_native_pin_open(life,run['root_pin'],private=True)
        pr=_native_pin_open(life,c['controls']['protocol'],private=True)
        mr=_native_pin_open(life,c['controls']['manifest'],private=True)
        catalog=_native_pin_open(life,c['catalog'],private=True)
        decoded=decode(catalog)
        from .series_catalog import canonical_catalog
        if set(decoded)!= {'schema_version','series'} or decoded['schema_version']!='forecast-series-catalog.v1' or canonical_catalog(decoded['series'])!=catalog:
            raise ValueError('generic canonical sorted catalog required')
        for item in c['code']:
            raw=_native_pin_open(life,item['file'])
            if hashlib.sha1(b'blob '+str(len(raw)).encode()+b'\0'+raw).hexdigest()!=item['git_blob']:
                raise ValueError('generic code Git blob drift')
        _native_pin_open(life,c['runtime']['resolved_executable'])
        from importlib.machinery import PathFinder
        from email.parser import BytesParser
        for item in c['runtime']['packages']:
            raw=_native_pin_open(life,item['metadata']);module={'scikit-learn':'sklearn'}.get(item['distribution'],item['distribution'])
            spec=PathFinder.find_spec(module)
            if BytesParser().parsebytes(raw).get('Version')!=item['version'] or spec is None or spec.origin!=item['module_origin']:
                raise ValueError('generic package metadata/origin drift')
        for source in c['sources'].values():_native_pin_open(life,source['file'],private=True,body=False)
        from .observed_runner import PROTOCOL_FIELDS
        observed=c['profile']=='generic-observed-v1'
        p=decode(pr);_native_equal(p,{**PROTOCOL_FIELDS,'as_of':'2026-10-09','profile':c['profile'],
                        'status':'APPROVED_OBSERVED_EXECUTION' if observed else 'INFRASTRUCTURE_ONLY_OBSERVED_EXECUTION_HOLD',
                        'observed_execution_authorized':observed})
        _native_equal(decode(mr),{'schema_version':'generic-source-manifest.v1','profile':c['profile'],
              'sources':{k:{'sha256':s['file']['sha256'],'bytes':s['file']['bytes'],'source_id':s['source_id']}for k,s in c['sources'].items()}})
        life.check();return life
    except BaseException:
        life.close();raise


def register_root_accepted_generic_run(authority,root_acceptance_raw):
    host=_native_host(authority) # issued identity before metadata/body
    if host['profile'] not in ('generic-observed-v1','generated-fixture-test-v1'):raise ValueError('generic issued host required')
    if not isinstance(root_acceptance_raw,bytes) or len(root_acceptance_raw)>128*1024 or digest(root_acceptance_raw)!=host['pin']['sha256']:
        raise ValueError('generic external control SHA mismatch')
    c=decode(root_acceptance_raw);_generic_control_schema(c,host,authority)
    with _NATIVE_LOCK:
        if host['used_controls'] or host['used_ids'] or host['used_outputs']:raise ValueError('generic reservation consumed; fresh host required')
        host['used_controls'].add(host['pin']['sha256']);host['used_ids'].add(c['run_id']);host['used_outputs'].add(c['outputs']['path'])
    run=_native_freeze({'authority':authority,'epoch':host['epoch'],'control':c,'root_pin':host['pin'],'output':c['outputs']['path'],
                       'parent_identity':_native_identity_doc(c['outputs']['parent_identity'],True),'raw':root_acceptance_raw})
    life=_generic_pins_lifetime(run)
    try:
        if os.path.lexists(run['output']) or _identity(os.lstat(Path(run['output']).parent),True)!=run['parent_identity']:
            raise ValueError('generic output absence/parent drift')
    finally:life.close()
    binding=RootGenericRunBinding(c['run_id']);_NATIVE_RUNS[binding]=run
    _NATIVE_LIVES[binding]={'state':'READY','token':None,'created':None,'sealed':None,'closure_sha':None}
    return binding


def open_root_generic_snapshot(binding):
    if not isinstance(binding,RootGenericRunBinding):raise ValueError('generic issued run identity required')
    run=native_run_record(binding)
    if _native_state(binding)['state']!='RUNNING':raise ValueError('generic running host required before body')
    life=_generic_pins_lifetime(run)
    try:
        c=run['control'];sources=_native_plain(c['sources'])
        # No reopen: hash and parse exactly the bytes returned by owned no-follow FDs.
        raw={k:_native_pin_open(life,s['file'],private=True)for k,s in c['sources'].items()}
        from .series_catalog import build_catalog,finite_batches
        snap=build_catalog(raw['wb'],raw['bis'],sources=sources,profile=c['profile'],regime_dispositions=_native_plain(c['regime_dispositions']))
        if digest(snap.canonical_bytes)!=c['catalog']['sha256'] or len(snap.canonical_bytes)!=c['catalog']['bytes']:
            raise ValueError('generic reconstructed catalog drift')
        queue=finite_batches(snap.records)
        if c['batch_index']>=len(queue) or tuple(c['selected_ids'])!=queue[c['batch_index']]:raise ValueError('generic finite queue batch mismatch')
        selected={sid:dict(snap.values[sid]) for sid in c['selected_ids']}
        basis={sid:{m:{'hex':v.hex(),'numerator':str(Fraction.from_float(v).numerator),
                      'denominator':str(Fraction.from_float(v).denominator)}for m,v in series.items()}for sid,series in selected.items()}
        receipt={'catalog':list(snap.records),'catalog_sha256':c['catalog']['sha256'],'batch_queue':[list(b)for b in queue],
                 'selected_ids':list(c['selected_ids']),'batch_index':c['batch_index'],'required_months':120,
                 'catalog_status':{r['series_id']:('selected_current_batch' if r['series_id'] in c['selected_ids'] else
                     'queued_separate_root_control_required' if r['forecast_eligible'] else 'unavailable')for r in snap.records},
                 'stored_value_basis':basis,'stored_value_basis_sha256':digest(canonical(basis)),
                 'root_control_sha256':run['root_pin']['sha256'],'full_parent_snapshot_exposed':True,
                 'physical_test_actuals_untouched':False,'is_synthetic':c['profile']=='generated-fixture-test-v1',
                 'observed_authenticity':c['profile']=='generic-observed-v1','observed_execution_authorized':c['profile']=='generic-observed-v1',
                 'availability_proxy':'retrospective latest-vintage m+1; not PIT'}
        life.check();return _NativeRead(life,selected,receipt)
    except BaseException:
        life.close();raise
