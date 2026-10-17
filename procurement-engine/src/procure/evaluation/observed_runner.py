"""Offline observed-forecast runner infrastructure. Real observed execution is held."""
from __future__ import annotations

import hashlib
import html
import json
import math
import os
import signal
import subprocess
import sys
import threading
import time
from collections import defaultdict
from decimal import Decimal
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd

from .forecast_scorecard import write_run

ARTIFACTS = ("predictions.json", "metrics.csv", "manifest.json", "report.md")
GRID = (0.8, 0.9, 1.0, 1.1, 1.2, 1.35, 1.5)
MONTHS = {"training_start": "2016-09", "validation_start": "2021-09", "validation_end": "2023-08",
          "test_start": "2024-03", "test_end": "2026-02"}
PROTOCOL_FIELDS = {
    "schema_version": "observed-forecast-runner.v1", "as_of": "2026-10-08",
    **MONTHS, "purge_start": "2023-09", "purge_end": "2024-02",
    "latest_eligible_actual": "2026-08",
    "availability_proxy": "month_m_available_at_end_of_m_plus_1; not historical publication proof",
    "horizons": [1, 3, 6], "model_leads_from_last_training_month": [2, 4, 7],
    "training_cutoff": "origin_minus_1", "minimum_consecutive_training_months": 60,
    "validation_effective_origins": 24, "purged_origins": 6, "test_origins": 24,
    "base_seed": 42, "n_paths": 1000, "fit_steps": 7, "gbm_refit_every": 1,
    "bootstrap_rng": "numpy.PCG64", "quantile_method": "linear",
    "calibration_grid": list(GRID), "selection_relative_improvement": 0.02,
    "fit_timeout_seconds": 60, "per_series_deadline_seconds": 900,
    "total_deadline_seconds": 1800,
}


def _sha(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def _unique_pairs(pairs):
    out = {}
    for key, value in pairs:
        if key in out:
            raise ValueError(f"duplicate JSON key: {key}")
        out[key] = value
    return out


def _no_constant(value):
    raise ValueError(f"nonfinite JSON constant: {value}")


def strict_json(path: str | Path, max_bytes: int = 1024 * 1024) -> tuple[dict, bytes]:
    source = Path(path)
    if source.stat().st_size > max_bytes:
        raise ValueError("JSON size limit exceeded")
    raw = source.read_bytes()
    value = json.loads(raw, object_pairs_hook=_unique_pairs, parse_constant=_no_constant)
    if not isinstance(value, dict):
        raise ValueError("JSON root must be an object")
    json.dumps(value, allow_nan=False)
    return value, raw


def _month(value: str) -> int:
    if not isinstance(value, str) or len(value) != 7 or value[4] != "-" or not value[:4].isdigit() or not value[5:].isdigit():
        raise ValueError("invalid YYYY-MM month")
    year, month = int(value[:4]), int(value[5:])
    if year < 1 or month < 1 or month > 12:
        raise ValueError("invalid YYYY-MM month")
    return year * 12 + month - 1


def _name(index: int) -> str:
    return f"{index // 12:04d}-{index % 12 + 1:02d}"


def _range(start: str, end: str) -> list[str]:
    return [_name(index) for index in range(_month(start), _month(end) + 1)]


def _check_protocol(protocol: dict, as_of: str) -> None:
    generic = isinstance(protocol, dict) and protocol.get("profile") in {"generic-observed-v1", "generated-fixture-test-v1"}
    expected_asof = "2026-10-09" if generic else "2026-10-08"
    if not isinstance(protocol, dict) or as_of != protocol.get("as_of") or as_of != expected_asof:
        raise ValueError("as_of/protocol mismatch")
    try:
        date.fromisoformat(as_of)
    except (TypeError, ValueError) as exc:
        raise ValueError("invalid as_of") from exc
    for field, expected in {**PROTOCOL_FIELDS, "as_of": expected_asof}.items():
        actual = protocol.get(field)
        if type(actual) is not type(expected) or actual != expected or (
                isinstance(expected, list) and any(type(a) is not type(e) for a, e in zip(actual, expected))):
            raise ValueError(f"protocol {field} mismatch")
    status = protocol.get("status")
    authorized = protocol.get("observed_execution_authorized")
    if type(authorized) is not bool or (status, authorized) not in {
        ("INFRASTRUCTURE_ONLY_OBSERVED_EXECUTION_HOLD", False),
        ("APPROVED_OBSERVED_EXECUTION", True),
    }:
        raise ValueError("protocol authorization mismatch")
    for field in ("snapshot_sha256", "protocol_sha256", "source_sha256"):
        value = protocol.get(field)
        if not isinstance(value, str) or len(value) != 64 or any(c not in "0123456789abcdef" for c in value):
            raise ValueError(f"protocol invalid {field}")


def _check_series(series: dict, start: str, end: str) -> None:
    if not isinstance(series, dict) or list(sorted(series)) != _range(start, end):
        raise ValueError("series must contain unique consecutive monthly values with no missing month")
    for value in series.values():
        if isinstance(value, bool) or not isinstance(value, (float, int)) or not math.isfinite(value) or value <= 0:
            raise ValueError("series contains nonpositive or nonfinite price")


def plan_common_origins(series_by_id: dict[str, dict], as_of: str, protocol: dict) -> dict:
    """Reject incomplete fixed cohorts; never choose dates from observed model scores."""
    _check_protocol(protocol, as_of)
    if not isinstance(series_by_id, dict) or not series_by_id:
        raise ValueError("no series supplied")
    latest_actual = "2026-08"  # m+1 proxy as of 2026-10-08
    for series in series_by_id.values():
        _check_series(series, MONTHS["training_start"], latest_actual)
    validation = _range(MONTHS["validation_start"], MONTHS["validation_end"])
    test = _range(MONTHS["test_start"], MONTHS["test_end"])
    purged = _range("2023-09", "2024-02")
    if len(validation) != 24 or len(purged) != 6 or len(test) != 24:
        raise ValueError("invalid frozen origin counts")
    if _month(validation[-1]) > _month(test[0]) - 7:
        raise ValueError("validation purge boundary breached")
    if _month(latest_actual) < _month(test[-1]) + 6 or _month(validation[0]) - 1 - _month(MONTHS["training_start"]) + 1 < 60:
        raise ValueError("insufficient_history")
    return {"validation": validation, "purged": purged, "test": test, "latest_actual": latest_actual}


def canonical_seed(protocol: dict, series_id: str, origin: str) -> int:
    _check_protocol(protocol, protocol.get("as_of"))
    material = {"base_seed": 42, "origin": origin, "protocol_sha256": protocol["protocol_sha256"],
                "series_id": series_id, "snapshot_sha256": protocol["snapshot_sha256"]}
    raw = json.dumps(material, sort_keys=True, separators=(",", ":"), allow_nan=False).encode()
    return int.from_bytes(hashlib.sha256(raw).digest()[:8], "big", signed=False)


def forecast_at_origin(series: dict, origin: str, model: str, protocol: dict, series_spec: dict,
                       split: str, fitter=None) -> list[dict]:
    """Create three frozen rows from one past-only fit; fake fitters support offline tests."""
    _check_protocol(protocol, protocol.get("as_of"))
    if model not in {"naive", "ets", "gbm"} or split not in {"validation", "test"}:
        raise ValueError("unsupported model or split")
    allowed = _range(MONTHS["validation_start"], MONTHS["validation_end"]) if split == "validation" else _range(MONTHS["test_start"], MONTHS["test_end"])
    if origin not in allowed:
        raise ValueError("origin outside frozen split")
    cutoff = _name(_month(origin) - 1)
    targets = [_name(_month(origin) + h) for h in (1, 3, 6)]
    required = _range(MONTHS["training_start"], targets[-1])
    if list(sorted(series)) != required and not set(required).issubset(series):
        raise ValueError("missing or nonconsecutive series months")
    training_months = _range(MONTHS["training_start"], cutoff)
    if any(month not in series for month in training_months + targets):
        raise ValueError("missing training or target month")
    training = {month: series[month] for month in training_months}
    _check_series(training, MONTHS["training_start"], cutoff)
    if any(not isinstance(series[t], (int, float)) or isinstance(series[t], bool) or not math.isfinite(series[t]) or series[t] <= 0 for t in targets):
        raise ValueError("invalid actual target")
    from .observed_admission import canonical
    frozen = canonical(forecast_from_training(training, origin, model, protocol, series_spec, split, fitter))
    return attach_frozen_actuals(frozen, {target: series[target] for target in targets},
                                expected_prediction_sha256=_sha(frozen))


def forecast_from_training(training: dict[str, float], origin: str, model: str, protocol: dict,
                           series_spec: dict, split: str, fitter=None) -> list[dict]:
    """Numerical seam only, no execution authority and no target actuals accepted."""
    _check_protocol(protocol, protocol.get("as_of"))
    if model not in {"naive", "ets", "gbm"} or split not in {"validation", "test"}:
        raise ValueError("unsupported model or split")
    allowed = _range(MONTHS[split + "_start"], MONTHS[split + "_end"])
    if origin not in allowed:
        raise ValueError("origin outside frozen split")
    cutoff = _name(_month(origin) - 1)
    targets = [_name(_month(origin) + h) for h in (1, 3, 6)]
    _check_series(training, MONTHS["training_start"], cutoff)
    training = {month: training[month] for month in _range(MONTHS["training_start"], cutoff)}
    if len(training) < 60:
        raise ValueError("insufficient_history")
    p0 = float(training[cutoff])
    prices = list(training.values())
    scale = float(np.mean(np.abs(np.diff(prices))))
    log_series = pd.Series(np.log(prices), index=pd.date_range(MONTHS["training_start"] + "-01", periods=len(prices), freq="MS"))
    if fitter is None:
        fitter = fit_in_owned_child
    name = protocol["snapshot_sha256"] + ":" + series_spec["series_id"]
    result = fitter(model, log_series, 7, name, 1)
    drift = np.asarray(result.drift, dtype=float)
    residuals = np.asarray(result.residuals, dtype=float)
    noise_scale = result.noise_scale
    if drift.shape != (7,) or not np.isfinite(drift).all() or residuals.ndim != 1 or len(residuals) < 12 or not np.isfinite(residuals).all():
        raise ValueError("invalid fitted drift or residuals")
    if isinstance(noise_scale, bool) or not isinstance(noise_scale, (int, float)) or not math.isfinite(noise_scale) or noise_scale <= 0:
        raise ValueError("invalid fitted noise_scale")
    seed = canonical_seed(protocol, series_spec["series_id"], origin)
    draws = np.random.Generator(np.random.PCG64(seed)).random((1000, 7))
    indices = np.floor(draws * len(residuals)).astype(int)
    with np.errstate(over="ignore", invalid="ignore", under="ignore"):
        paths = p0 * np.exp(np.cumsum(drift + float(noise_scale) * residuals[indices], axis=1))
    if not np.isfinite(paths).all() or (paths <= 0).any():
        raise ValueError("nonfinite or nonpositive forecast paths")
    quantiles = np.quantile(paths, [0.1, 0.5, 0.9], axis=0, method="linear")
    backend = getattr(result, "backend", series_spec.get("backend", "fixture" if series_spec.get("is_synthetic") else "unknown"))
    if backend == "unknown":
        raise ValueError("backend identity unavailable")
    if model == "gbm" and backend not in {"sklearn", "fixture"}:
        raise ValueError("unsupported_unpinned_backend")
    common = {key: series_spec[key] for key in ("series_kind", "series_id", "unit", "source_id", "is_synthetic")}
    rows = []
    for horizon, lead, target in zip((1, 3, 6), (2, 4, 7), targets):
        rows.append({**common, "split": split, "origin": origin, "target": target, "horizon": horizon,
                     "model": model, "backend": backend, "training_cutoff": cutoff, "p0": p0,
                     "p10": float(quantiles[0, lead - 1]), "p50": float(quantiles[1, lead - 1]),
                     "p90": float(quantiles[2, lead - 1]), "mase_scale": scale})
    return rows


def attach_frozen_actuals(frozen_predictions: bytes, actuals: dict[str, float], *,
                          expected_prediction_sha256: str) -> list[dict]:
    """Pure attachment; the admitted host separately owns causal release authority."""
    from .observed_admission import decode
    if _sha(frozen_predictions) != expected_prediction_sha256:
        raise ValueError("foreign/stale prediction freeze")
    rows = decode(frozen_predictions)
    if not isinstance(rows, list) or not rows or any(not isinstance(row, dict) or "actual" in row for row in rows):
        raise ValueError("prediction freeze must be actual-free")
    identities = {(row.get("series_id"), row.get("model"), row.get("origin"), row.get("split")) for row in rows}
    if len(rows) != 3 or len(identities) != 1 or {row.get("horizon") for row in rows} != {1, 3, 6}:
        raise ValueError("duplicate or foreign prediction identity")
    targets = {row["target"] for row in rows}
    if len(targets) != 3 or any(row["target"] != _name(_month(row["origin"]) + row["horizon"]) for row in rows):
        raise ValueError("duplicate or foreign target identity")
    if not isinstance(actuals, dict) or set(actuals) != targets or any(
            type(value) not in {int, float} or not math.isfinite(value) or value <= 0 for value in actuals.values()):
        raise ValueError("invalid actual target")
    return [{**row, "actual": float(actuals[row["target"]])} for row in rows]


def _scaled_interval(row: dict, multiplier: float) -> dict:
    p10, p50, p90 = row["p10"], row["p50"], row["p90"]
    if min(p10, p50, p90) <= 0 or not all(math.isfinite(x) for x in (p10, p50, p90)):
        raise ValueError("invalid interval")
    try:
        low = p50 if p10 == p50 else min(
            p50, math.exp(math.log(p50) + multiplier * (math.log(p10) - math.log(p50))))
        high = p50 if p90 == p50 else max(
            p50, math.exp(math.log(p50) + multiplier * (math.log(p90) - math.log(p50))))
    except OverflowError as exc:
        raise ValueError("calibrated interval overflow") from exc
    if not math.isfinite(low) or not math.isfinite(high) or low <= 0 or high <= 0:
        raise ValueError("calibrated interval overflow")
    return {**row, "p10": low, "p90": high}


def calibrate_intervals(validation_rows: list[dict]) -> tuple[float, list[dict]]:
    if len(validation_rows) != 24 or len({row["origin"] for row in validation_rows}) != 24:
        raise ValueError("calibration requires 24 distinct validation origins")
    candidates = []
    for multiplier in GRID:
        try:
            rows = [_scaled_interval(row, multiplier) for row in validation_rows]
        except ValueError:
            continue
        coverage = sum(row["p10"] <= row["actual"] <= row["p90"] for row in rows) / 24
        candidates.append((abs(coverage - 0.8), multiplier, rows))
    if not candidates:
        raise ValueError("calibration_unavailable")
    _, multiplier, rows = min(candidates, key=lambda item: (item[0], item[1]))
    return multiplier, rows


def select_models(validation_rows: list[dict]) -> dict[tuple[str, int], str]:
    grouped = defaultdict(lambda: defaultdict(dict))
    for row in validation_rows:
        key = (row["series_id"], row["horizon"])
        if row["origin"] in grouped[key][row["model"]]:
            raise ValueError("duplicate validation cell")
        grouped[key][row["model"]][row["origin"]] = abs(Decimal(str(row["p50"])) - Decimal(str(row["actual"])))
    selections = {}
    for key, models in grouped.items():
        naive = models.get("naive", {})
        if len(naive) != 24:
            raise ValueError("naive missing validation cell")
        eligible = []
        naive_mae = sum(naive.values()) / 24
        for model in ("ets", "gbm"):
            cells = models.get(model, {})
            if set(cells) != set(naive):
                continue  # whole candidate unavailable; caller ledger records why
            error = sum(cells.values()) / 24
            if naive_mae > 0 and error < Decimal("0.98") * naive_mae:
                eligible.append((error, 0 if model == "ets" else 1, model))
        selections[key] = min(eligible)[2] if eligible else "naive"
    return selections


def _finalize_cohorts(validation_rows: list[dict], test_rows: dict, failed: dict,
                      ledger: list[dict]) -> tuple[list[dict], list[dict]]:
    """Retain only complete model cohorts; a missing naive removes its series."""
    unavailable_series = {series_id for (series_id, model), _ in failed.items() if model == "naive"}
    available = set(test_rows) - set(failed)
    available -= {(series_id, model) for series_id, model in available if series_id in unavailable_series}
    for item in ledger:
        if item["status"] != "pending":
            continue
        key = (item["series_id"], item["model"])
        if key in failed and item["model"] == "naive":
            item.update(status="not_evaluated", reason=failed[key])
        elif item["series_id"] in unavailable_series:
            item.update(status="not_evaluated", reason="naive_baseline_unavailable")
        elif key in failed:
            item.update(status="not_evaluated", reason=failed[key])
        elif key in available:
            item.update(status="scored", reason=None)
        else:
            item.update(status="not_evaluated", reason="missing_cohort")
    rows = [row for row in validation_rows if (row["series_id"], row["model"]) in available]
    rows += [row for key, cohort in test_rows.items() if key in available for row in cohort]
    return rows, ledger


def verify_stage1_run(run: str | Path, expected_input_sha256: str) -> bool:
    root = Path(run)
    state, _ = strict_json(root / "RUN_STATE.json")
    if state.get("status") != "complete" or set(state.get("artifact_sha256", {})) != set(ARTIFACTS):
        raise ValueError("Stage1 run incomplete or missing four artifact hashes")
    for name in ARTIFACTS:
        if _sha((root / name).read_bytes()) != state["artifact_sha256"][name]:
            raise ValueError(f"Stage1 artifact hash mismatch: {name}")
    manifest, _ = strict_json(root / "manifest.json")
    if manifest.get("input_sha256") != expected_input_sha256 or _sha((root / "predictions.json").read_bytes()) != expected_input_sha256:
        raise ValueError("Stage1 input hash mismatch")
    return True


def load_verified_native_series(manifest: dict, series_id: str, source_file: str | Path,
                                *, allow_synthetic_fixture: bool = False,
                                through_month: str | None = None) -> dict[str, float]:
    """Hash and header-gated local workbook loader; tests may opt into marked fixtures."""
    if not isinstance(manifest, dict) or not isinstance(manifest.get("is_synthetic"), bool):
        raise ValueError("invalid source manifest")
    if manifest["is_synthetic"]:
        if not allow_synthetic_fixture:
            raise ValueError("synthetic source is not observed")
    elif manifest.get("eligibility_status") != "APPROVED_OBSERVED":
        raise ValueError("source eligibility HOLD")
    if not isinstance(source_file, (str, Path)) or not Path(source_file).is_absolute():
        raise ValueError("source-file must be absolute local path")
    path = Path(source_file)
    if path.stat().st_size > 12 * 1024 * 1024 or path.stat().st_size != manifest.get("bytes"):
        raise ValueError("source byte size mismatch or limit")
    raw = path.read_bytes()
    if _sha(raw) != manifest.get("sha256"):
        raise ValueError("source hash mismatch")
    spec = manifest.get("series", {}).get(series_id)
    if not isinstance(spec, dict) or not all(isinstance(spec.get(key), str) and spec[key] for key in ("label", "column", "unit")):
        raise ValueError("series manifest missing")
    from openpyxl import load_workbook
    from openpyxl.utils.cell import column_index_from_string
    try:
        column = column_index_from_string(spec["column"])
    except ValueError as exc:
        raise ValueError("invalid source column") from exc
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book[manifest["sheet"]]
        if sheet.cell(5, column).value != spec["label"] or sheet.cell(6, column).value != spec["unit"]:
            raise ValueError("source header/unit mismatch")
        declared = manifest.get("month_range")
        if not isinstance(declared, str) or len(declared) != 15 or declared[7] != "-":
            raise ValueError("invalid declared month range")
        def source_month(label):
            if not isinstance(label, str) or len(label) != 7 or label[4] != "M":
                raise ValueError("invalid monthly source label")
            month = label[:4] + "-" + label[5:]
            _month(month)
            return month
        declared_months = _range(source_month(declared[:7]), source_month(declared[8:]))
        labels = [source_month(row[0]) for row in sheet.iter_rows(min_row=7, min_col=1, max_col=1, values_only=True)
                  if row[0] is not None]
        if labels != declared_months:
            raise ValueError("source month range mismatch")
        series = {}
        for row in sheet.iter_rows(min_row=7, min_col=1, max_col=column, values_only=True):
            label = row[0]
            if label is None and row[column - 1] is None:
                continue
            month = source_month(label)
            if through_month is not None and _month(month) > _month(through_month):
                break
            value = row[column - 1]
            if month in series:
                raise ValueError("duplicate source month")
            if isinstance(value, bool) or not isinstance(value, (float, int)) or not math.isfinite(value) or value <= 0:
                raise ValueError("missing/nonpositive/nonfinite source price")
            series[month] = float(value)
        if not series or list(series) != _range(min(series), max(series)):
            raise ValueError("source months not consecutive")
        return series
    finally:
        book.close()


class OwnedChildCleanupError(OSError):
    """An owned fit group could not be confirmed cleaned up."""


def run_bounded_child(command: list[str], payload: bytes, *, timeout: float = 60.0) -> bytes:
    """Run one owned child with a clean environment; terminate/reap its group."""
    if not command or timeout <= 0 or timeout > 60 or len(payload) > 1024 * 1024:
        raise ValueError("invalid owned-child bounds")
    src_root = Path(__file__).resolve().parents[2]
    environment = {"PATH": os.defpath, "PYTHONPATH": str(src_root), "PYTHONNOUSERSITE": "1"}
    armed = threading.current_thread() is threading.main_thread()
    interrupted = False
    if armed:
        previous = signal.getsignal(signal.SIGTERM)
        def interrupt(_signum, _frame):
            nonlocal interrupted
            interrupted = True
        signal.signal(signal.SIGTERM, interrupt)
    child = None
    try:
        child = subprocess.Popen(command, stdin=subprocess.PIPE, stdout=subprocess.PIPE,
                                 stderr=subprocess.PIPE, env=environment, start_new_session=True)
        deadline = time.monotonic() + timeout
        pending_input = payload
        while True:
            if interrupted:
                raise InterruptedError("owned fit interrupted")
            remaining = deadline - time.monotonic()
            if remaining <= 0:
                raise subprocess.TimeoutExpired(command, timeout)
            try:
                stdout, _stderr = child.communicate(pending_input, timeout=min(0.2, remaining))
                break
            except subprocess.TimeoutExpired:
                pending_input = None
                continue
        if interrupted:
            raise InterruptedError("owned fit interrupted")
    except (subprocess.TimeoutExpired, KeyboardInterrupt, InterruptedError) as exc:
        if child is None:
            raise
        cleanup_unverified = False
        try:
            try:
                os.killpg(child.pid, signal.SIGTERM)
            except OSError:
                # The later KILL/status check resolves a TERM race or denial.
                pass
            try:
                child.wait(timeout=0.3)
            except subprocess.TimeoutExpired:
                pass
            except OSError:
                cleanup_unverified = True
            # A descendant may retain stdout/stderr after the leader exits.
            try:
                os.killpg(child.pid, signal.SIGKILL)
            except ProcessLookupError:
                pass
            except OSError:
                # EPERM can race with the last member exiting. Only ESRCH
                # establishes that the owned group is gone.
                for _ in range(5):
                    try:
                        os.killpg(child.pid, 0)
                    except ProcessLookupError:
                        break
                    except OSError:
                        pass
                    time.sleep(0.05)
                else:
                    cleanup_unverified = True
        finally:
            try:
                direct_alive = child.poll() is None
            except OSError:
                direct_alive = True
                cleanup_unverified = True
            if direct_alive:
                try:
                    child.kill()
                except ProcessLookupError:
                    pass
                except OSError:
                    cleanup_unverified = True
            try:
                child.wait(timeout=2)
            except (subprocess.TimeoutExpired, OSError):
                cleanup_unverified = True
            for pipe in (child.stdin, child.stdout, child.stderr):
                if pipe is not None:
                    try:
                        pipe.close()
                    except OSError:
                        cleanup_unverified = True
        if cleanup_unverified:
            raise OwnedChildCleanupError("owned_child_cleanup_unverified") from None
        if isinstance(exc, (KeyboardInterrupt, InterruptedError)):
            raise
        raise TimeoutError("owned child fit timed out and was reaped")
    finally:
        if armed:
            signal.signal(signal.SIGTERM, previous)
    if child.returncode != 0 or len(stdout) > 1024 * 1024:
        raise RuntimeError("owned child failed; stderr suppressed")
    return stdout


def fit_in_owned_child(model: str, log_series: pd.Series, horizon: int, name: str, refit_every: int,
                       *, remaining_seconds: float = 60.0, result_contract: str = "legacy", expected_versions=None):
    if result_contract == "native-v1":
        return _fit_in_native_child(model, log_series, horizon, name, refit_every, remaining_seconds, expected_versions)
    if result_contract != "legacy" or expected_versions is not None:
        raise NativeChildProtocolError("unknown native contract selector")
    from types import SimpleNamespace
    if horizon != 7 or refit_every != 1 or model not in {"naive", "ets", "gbm"}:
        raise ValueError("fit protocol mismatch")
    payload = json.dumps({"model": model, "name": name, "months": [v.strftime("%Y-%m") for v in log_series.index],
                          "values": [float(v) for v in log_series.values]}, allow_nan=False).encode()
    response = run_bounded_child([sys.executable, "-m", "procure.evaluation.observed_runner", "--fit-child"], payload,
                                 timeout=min(60.0, remaining_seconds))
    result = json.loads(response, object_pairs_hook=_unique_pairs, parse_constant=_no_constant)
    if result.get("status") != "ok":
        raise RuntimeError(result.get("reason", "fit_unavailable"))
    return SimpleNamespace(drift=np.asarray(result["drift"]), residuals=np.asarray(result["residuals"]),
                           noise_scale=result["noise_scale"], backend=result["backend"], versions=result["versions"])


def _fit_child_main() -> int:
    from importlib.metadata import version
    from procure.forecast import models
    try:
        payload = json.loads(sys.stdin.buffer.read(1024 * 1024 + 1), object_pairs_hook=_unique_pairs, parse_constant=_no_constant)
        if payload["model"] == "gbm" and models.GBM_BACKEND != "sklearn":
            response = {"status": "unavailable", "reason": "unsupported_unpinned_backend"}
        else:
            index = pd.to_datetime([month + "-01" for month in payload["months"]])
            values = pd.Series(payload["values"], index=index)
            models.clear_cache()
            fit = models.fit_series(payload["model"], values, 7, name=payload["name"], refit_every=1)
            response = {"status": "ok", "drift": fit.drift.tolist(), "residuals": np.asarray(fit.residuals).tolist(),
                        "noise_scale": float(fit.noise_scale),
                        "backend": models.GBM_BACKEND if payload["model"] == "gbm" else payload["model"],
                        "versions": {package: version(package) for package in ("numpy", "pandas", "scikit-learn", "statsmodels")}}
        sys.stdout.write(json.dumps(response, allow_nan=False))
        return 0
    except Exception:
        sys.stdout.write('{"status":"unavailable","reason":"fit_failed"}')
        return 0


def _svg_chart(rows: list[dict], label: str) -> str:
    """Dependency-free chart from frozen prediction rows, with safe XML text."""
    ordered = sorted(rows, key=lambda row: row["origin"])
    values = [float(row[key]) for row in ordered for key in ("actual", "p10", "p50", "p90")]
    if not values or not all(math.isfinite(value) for value in values):
        raise ValueError("chart values unavailable")
    low, high = min(values), max(values)
    if low == high:
        high = low + 1
    def point(index, value):
        x = 60 + index * 720 / max(len(ordered) - 1, 1)
        y = 340 - (float(value) - low) * 280 / (high - low)
        return f"{x:.2f},{y:.2f}"
    safe_label = html.escape(label, quote=True)
    safe_label = "".join(f"\\u{ord(char):04x}" if ord(char) < 32 or 0x7f <= ord(char) <= 0x9f or char in "\u2028\u2029" else char for char in safe_label)
    lines = ['<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 840 400" role="img">',
             f'<title>{safe_label}</title>',
             '<rect width="840" height="400" fill="white"/>']
    if any(row.get("is_synthetic") for row in ordered):
        lines.append('<text x="420" y="210" text-anchor="middle" opacity="0.35">SYNTHETIC TEST FIXTURE</text>')
    lines.append('<text x="40" y="25">Latest-vintage retrospective; month m available at m+1 proxy</text>')
    for key, color in (("p10", "#86a6c4"), ("p90", "#86a6c4"), ("actual", "#17324d"), ("p50", "#b45a37")):
        points = " ".join(point(i, row[key]) for i, row in enumerate(ordered))
        lines.append(f'<polyline fill="none" stroke="{color}" stroke-width="2" points="{points}"/>')
    lines.append('</svg>')
    return "\n".join(lines) + "\n"


def publish_runner(output_dir: str | Path, stage1_input: dict, ledger: list, selection: dict,
                   calibration: dict, provenance: dict) -> dict:
    """Publish only after nested Stage1 and every runner artifact hash verifies."""
    for field in ("source_manifest_sha256", "protocol_sha256", "seeds", "backends", "versions", "training_cutoffs"):
        if field not in provenance:
            raise ValueError(f"missing provenance {field}")
    marked_synthetic = stage1_input.get("metadata", {}).get("synthetic") is True
    rows = stage1_input.get("predictions", [])
    if any(row.get("is_synthetic") is not marked_synthetic for row in rows):
        raise ValueError("synthetic provenance mismatch")
    if not marked_synthetic and provenance.get("source_eligibility_status") != "APPROVED_OBSERVED":
        raise ValueError("observed source eligibility not approved")
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=False)
    state_path = output / "RUN_STATE.json"
    state_path.write_text('{"status":"incomplete"}\n', encoding="utf-8")
    payloads = {"predictions_input.json": stage1_input, "candidate_ledger.json": ledger,
                "selection.json": {str(key): value for key, value in selection.items()},
                "calibration.json": {str(key): value for key, value in calibration.items()},
                "provenance.json": provenance}
    for name, value in payloads.items():
        (output / name).write_text(json.dumps(value, sort_keys=True, allow_nan=False, indent=2) + "\n", encoding="utf-8")
    input_sha = _sha((output / "predictions_input.json").read_bytes())
    write_run(output / "predictions_input.json", output / "scorecard")
    verify_stage1_run(output / "scorecard", input_sha)
    chart_files = []
    grouped = defaultdict(list)
    for row in stage1_input.get("predictions", []):
        grouped[(row["series_kind"], row["series_id"], row["unit"], row["source_id"], row["model"], row["backend"], row["split"], row["horizon"])].append(row)
    charts_dir = output / "charts"
    charts_dir.mkdir()
    for group, rows in sorted(grouped.items()):
        digest = _sha(json.dumps(group, ensure_ascii=True).encode())[:16]
        name = f"charts/{digest}.svg"
        label = " / ".join(str(value) for value in group) + f" / source {provenance['source_manifest_sha256']}"
        (output / name).write_text(_svg_chart(rows, label), encoding="utf-8")
        chart_files.append(name)
    names = list(payloads) + ["scorecard/" + name for name in (*ARTIFACTS, "RUN_STATE.json")] + chart_files
    hashes = {name: _sha((output / name).read_bytes()) for name in names}
    state_temp = output / ".RUN_STATE.tmp"
    state_temp.write_text(json.dumps({"status": "complete", "artifact_sha256": hashes, "stage1_input_sha256": input_sha}, sort_keys=True) + "\n", encoding="utf-8")
    os.replace(state_temp, state_path)
    return {"artifact_sha256": hashes, "stage1_input_sha256": input_sha}


def verify_runner_run(output_dir: str | Path, *, expected_run=None) -> bool:
    if expected_run is not None:
        from .observed_admission import RootNativeRunBinding, RootGenericRunBinding
        if isinstance(expected_run, RootGenericRunBinding):
            return _verify_root_generic(output_dir, expected_run)
        if isinstance(expected_run, RootNativeRunBinding):
            return _verify_root_native(output_dir, expected_run)
        return _verify_admitted_fixture(output_dir, expected_run)
    output = Path(output_dir)
    state, _ = strict_json(output / "RUN_STATE.json")
    hashes = state.get("artifact_sha256")
    if state.get("status") != "complete" or not isinstance(hashes, dict):
        raise ValueError("runner incomplete")
    required = {"predictions_input.json", "candidate_ledger.json", "selection.json", "calibration.json", "provenance.json"}
    required.update("scorecard/" + name for name in (*ARTIFACTS, "RUN_STATE.json"))
    if not required.issubset(hashes):
        raise ValueError("runner artifact hash manifest incomplete")
    actual_charts = {"charts/" + path.name for path in (output / "charts").glob("*.svg")}
    declared_charts = {name for name in hashes if name.startswith("charts/")}
    if actual_charts != declared_charts:
        raise ValueError("runner chart hash manifest mismatch")
    for name, expected in hashes.items():
        relative = Path(name)
        if relative.is_absolute() or ".." in relative.parts or _sha((output / relative).read_bytes()) != expected:
            raise ValueError("runner artifact hash mismatch")
    input_sha = _sha((output / "predictions_input.json").read_bytes())
    if state.get("stage1_input_sha256") != input_sha:
        raise ValueError("runner input hash mismatch")
    verify_stage1_run(output / "scorecard", input_sha)
    return True


def _publish_incomplete_run(output_dir: str | Path, ledger: list[dict], reason: str) -> None:
    """Leave a sanitized, explicitly incomplete receipt after a failed fit."""
    if reason not in {"fit_deadline", "fit_unavailable", "backend_changed", "calibration_unavailable", "fit_interrupted", "owned_child_cleanup_unverified"}:
        reason = "fit_unavailable"
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=False)
    (output / "RUN_STATE.json").write_text(json.dumps({"status": "incomplete", "reason": reason}) + "\n", encoding="utf-8")
    sanitized_ledger = [{**item, "status": "not_evaluated", "reason": item.get("reason") or reason}
                        if item.get("status") == "pending" else item for item in ledger]
    (output / "candidate_ledger.json").write_text(json.dumps(sanitized_ledger, sort_keys=True, allow_nan=False, indent=2) + "\n", encoding="utf-8")


def run_offline_evaluation(protocol_path: str | Path, source_manifest_path: str | Path,
                           source_file: str | Path, output_dir: str | Path, as_of: str,
                           *, accepted_run=None) -> dict:
    if accepted_run is None:
        raise ValueError("registered execution binding required; observed execution HOLD")
    from .observed_admission import RootNativeRunBinding, RootGenericRunBinding
    if isinstance(accepted_run, RootGenericRunBinding):
        return _run_root_generic(protocol_path, source_manifest_path, source_file, output_dir, as_of, accepted_run)
    if isinstance(accepted_run, RootNativeRunBinding):
        return _run_root_native(protocol_path, source_manifest_path, source_file, output_dir, as_of, accepted_run)
    return _run_admitted_fixture(protocol_path, source_manifest_path, source_file, output_dir, as_of, accepted_run)


# A fixture run must intercept this PUBLIC boundary in its owning test. No
# product-supplied callback, child command, or real statistical fitting is enabled.
_REAL_FIT_BOUNDARY = fit_in_owned_child
_ADMITTED_SCHEMA = "observed-run-fixture-closure.v1"


def _owned_json(path, max_bytes=8 * 1024 * 1024):
    from .observed_admission import Lifetime, decode
    life = Lifetime()
    try:
        fd = life.open(path, limit=max_bytes)
        raw = life.read(fd, max_bytes)
        value = decode(raw)
        if not isinstance(value, dict):
            raise ValueError("artifact must be object")
        return value, raw
    finally:
        life.close()


def _exclusive_json(path, value):
    from .observed_admission import canonical, _identity
    raw = canonical(value)
    parent_fd = os.open(path.parent, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
    try:
        parent_identity = _identity(os.fstat(parent_fd), True)
        fd = os.open(path.name, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW, 0o600, dir_fd=parent_fd)
        try:
            offset = 0
            while offset < len(raw):
                written = os.write(fd, raw[offset:])
                if written <= 0:
                    raise OSError("short artifact write")
                offset += written
            os.fsync(fd)
            file_identity = _identity(os.fstat(fd))
            if os.fstat(fd).st_nlink != 1:
                raise ValueError("artifact hardlink")
        finally:
            os.close(fd)
    finally:
        os.close(parent_fd)
    if _identity(os.lstat(path.parent), True) != parent_identity or _identity(os.lstat(path)) != file_identity:
        raise ValueError("artifact postclose identity changed")
    return _sha(raw)


def _fixture_ledger():
    rows = []
    for kind, series_id, unit, reason in (
            ("commodity", "copper", "USD/metric tonne", None),
            ("commodity", "aluminum", "USD/metric tonne", None),
            ("fx", "USDINR", "INR/USD", "ecb_originals_absent"),
            ("joined", "CopperINRkg", "INR/kg", "fx_unavailable"),
            ("joined", "AluminiumINRkg", "INR/kg", "fx_unavailable"),
            ("commodity", "iron_ore", "USD/dmtu", "root_unit_review")):
        for model in (("all",) if series_id == "iron_ore" else ("naive", "ets", "gbm")):
            for split in ("validation", "test"):
                for horizon in (1, 3, 6):
                    rows.append({"series_kind": kind, "series_id": series_id, "unit": unit,
                                 "model": model, "split": split, "horizon": horizon,
                                 "status": "pending" if reason is None else "not_evaluated",
                                 "reason": reason, "backend": None, "metrics": None})
    assert len(rows) == 96
    return rows


class _FixturePhases:
    """Host-owned state, never a phase grant decoded from a product request."""

    def __init__(self, output, read, context):
        self.output, self.read, self.context = output, read, context
        self._step = 0
        self._freezes = {}
        self._file_ids = {}
        self.identity = (output.stat().st_dev, output.stat().st_ino, output.stat().st_mode, output.stat().st_uid)

    def check(self):
        self.read.check()
        info = os.lstat(self.output)
        if (info.st_dev, info.st_ino, info.st_mode, info.st_uid) != self.identity:
            raise ValueError("output identity changed")
        if os.path.lexists(self.output / "FAILURE.json"):
            raise ValueError("failure entry dominates")
        for name, (sha, _) in self._freezes.items():
            from .observed_admission import _identity
            if _identity(os.lstat(self.output / name)) != self._file_ids[name]:
                raise ValueError("freeze identity changed")
            doc, raw = _owned_json(self.output / name)
            if _sha(raw) != sha or doc.get("context") != self.context:
                raise ValueError("foreign/stale freeze")

    def freeze(self, split, rows):
        from .observed_admission import canonical
        self.check()
        if (split, self._step) not in {("validation", 0), ("test", 2)}:
            raise ValueError("prediction freeze wrong phase")
        if not rows or any("actual" in row or row["split"] != split for row in rows):
            raise ValueError("freeze must be actual-free")
        counts = defaultdict(set)
        identities = set()
        for row in rows:
            identity = (row["series_id"], row["model"], row["horizon"], row["origin"])
            if identity in identities or row["target"] != _name(_month(row["origin"]) + row["horizon"]):
                raise ValueError("duplicate/foreign prediction identity")
            identities.add(identity)
            counts[(row["series_id"], row["model"], row["horizon"])].add(row["origin"])
        origins = set(_range(MONTHS[split + "_start"], MONTHS[split + "_end"]))
        if any(cells != origins for cells in counts.values()):
            raise ValueError("incomplete prediction cohort")
        name = split + "_prediction_freeze.json"
        raw = canonical(rows)
        doc = {"schema_version": "actual-free-prediction-freeze.v1", "context": self.context,
               "split": split, "prediction_sha256": _sha(raw), "predictions": rows}
        sha = _exclusive_json(self.output / name, doc)
        self._freezes[name] = (sha, raw)
        from .observed_admission import _identity
        self._file_ids[name] = _identity(os.lstat(self.output / name))
        self._step += 1
        self.check()

    def attach(self, split):
        self.check()
        if (split, self._step) not in {("validation", 1), ("test", 3)}:
            raise ValueError("actual attachment wrong phase")
        name = split + "_prediction_freeze.json"
        _, raw = self._freezes[name]
        from .observed_admission import canonical, decode
        grouped = defaultdict(list)
        for row in decode(raw):
            grouped[(row["series_id"], row["model"], row["origin"])].append(row)
        scored = []
        for (sid, _, _), rows in grouped.items():
            frozen = canonical(rows)
            actuals = {row["target"]: self.read.series[sid][row["target"]] for row in rows}
            scored.extend(attach_frozen_actuals(frozen, actuals, expected_prediction_sha256=_sha(frozen)))
        self.check()
        return scored

    def freeze_selection(self, selections, calibration, failed):
        self.check()
        if self._step != 1:
            raise ValueError("selection freeze wrong phase")
        doc = {"schema_version": "validation-selection-freeze.v1", "context": self.context,
               "validation_prediction_sha256": self._freezes["validation_prediction_freeze.json"][0],
               "selection": selections, "calibration": calibration, "failed_cohorts": failed}
        sha = _exclusive_json(self.output / "selection_freeze.json", doc)
        self._freezes["selection_freeze.json"] = (sha, None)
        from .observed_admission import _identity
        self._file_ids["selection_freeze.json"] = _identity(os.lstat(self.output / "selection_freeze.json"))
        self._step = 2
        self.check()


def _run_admitted_fixture(protocol_path, manifest_path, source_file, output_dir, as_of, binding):
    from .observed_admission import (canonical, decode, fixture_run_record, fixture_admission_record,
                                     open_observed_native_snapshot, Lifetime, _identity,
                                     _begin_fixture_run, _capture_fixture_output, _seal_fixture_run, _fail_fixture_run)
    run = fixture_run_record(binding)  # permission before even source stat
    admission = fixture_admission_record(run["admission"])
    if tuple(map(str, (protocol_path, manifest_path, source_file, output_dir))) != (
            run["protocol"], run["manifest"], admission["source"], run["output"]):
        raise ValueError("foreign run/source/control/output binding")
    output = Path(output_dir)
    if os.path.lexists(output) or _identity(os.lstat(output.parent), True) != run["parent_identity"]:
        raise ValueError("output entry or parent identity changed")
    if fit_in_owned_child is _REAL_FIT_BOUNDARY:
        raise ValueError("fixture requires owning test interception; real fitting HOLD")
    controls = Lifetime()
    try:
        values = []
        for path, sha in admission["controls"]:
            fd = controls.open(path, private=True, limit=1024 * 1024)
            raw = controls.read(fd, 1024 * 1024)
            if _sha(raw) != sha:
                raise ValueError("foreign control hash")
            values.append((decode(raw), raw))
    finally:
        controls.close()
    (protocol, protocol_raw), (manifest, manifest_raw) = values
    if (protocol.get("status") != "INFRASTRUCTURE_ONLY_OBSERVED_EXECUTION_HOLD"
            or protocol.get("observed_execution_authorized") is not False
            or manifest.get("is_synthetic") is not True
            or manifest.get("observed_authenticity") is not False
            or manifest.get("observed_execution_authorized") is not False
            or manifest.get("sha256") != admission["grant"]["source_sha256"]
            or manifest.get("bytes") != admission["grant"]["source_bytes"]
            or manifest.get("source_id") != "owned_fixture"):
        raise ValueError("fixture controls required; observed execution HOLD")
    protocol = {**protocol, "protocol_sha256": _sha(protocol_raw),
                "source_sha256": manifest["sha256"], "snapshot_sha256": manifest["sha256"]}
    _check_protocol(protocol, as_of)
    context = {"fixture_id": binding.fixture_id, "source_sha256": manifest["sha256"],
               "protocol_sha256": _sha(protocol_raw), "manifest_sha256": _sha(manifest_raw),
               "code_sha256": dict(admission["code"]), "output": str(output), "as_of": as_of}
    ledger, failed = _fixture_ledger(), {}
    started, elapsed, backends = time.monotonic(), defaultdict(float), {}
    provenance = {"source_manifest_sha256": _sha(manifest_raw), "protocol_sha256": _sha(protocol_raw),
                  "seeds": {}, "backends": {}, "versions": {}, "training_cutoffs": {},
                  "source_slice_sha256": {}, "model_fit_seeds": {}, "source_snapshot_sha256": manifest["sha256"],
                  "source_eligibility_status": "OWNED_SYNTHETIC_FIXTURE", "latest_vintage_proxy": True,
                  "full_parent_snapshot_exposed": True, "physical_test_actuals_untouched": False,
                  "observed_authenticity": False, "observed_execution_authorized": False}
    token = _begin_fixture_run(binding)
    owned_output_identity = None
    try:
        with open_observed_native_snapshot(Path(source_file), admission["raw"], expected_admission=run["admission"]) as read:
            windows = plan_common_origins({sid: dict(s) for sid, s in read.series.items()}, as_of, protocol)
            read.check()
            output.mkdir(mode=0o700, exist_ok=False)
            owned_output_identity = _identity(os.lstat(output), True)
            read.check()
            _capture_fixture_output(binding, token, owned_output_identity)
            _exclusive_json(output / "HOST_PENDING.json", {"schema_version": _ADMITTED_SCHEMA, "context": context,
                                                           "status": "incomplete", "is_synthetic": True})
            host = _FixturePhases(output, read, context)
            _exclusive_json(output / "source_receipt.json", dict(read.receipt))

            def predict(split):
                predictions = []
                for sid in ("copper", "aluminum"):
                    for model in ("naive", "ets", "gbm"):
                        key = (sid, model)
                        if key in failed:
                            continue
                        if (sid, "naive") in failed:
                            failed[key] = "naive_baseline_unavailable"
                            continue
                        rows = []
                        before = time.monotonic()
                        try:
                            for origin in windows[split]:
                                host.check()
                                remaining = min(60.0, 1800 - (time.monotonic() - started),
                                                900 - elapsed[sid] - (time.monotonic() - before))
                                if remaining <= 0:
                                    raise TimeoutError("fit_deadline")
                                cutoff = _name(_month(origin) - 1)
                                training = {m: read.series[sid][m] for m in _range("2016-09", cutoff)}
                                spec = {"series_kind": "commodity", "series_id": sid, "unit": "USD/metric tonne",
                                        "source_id": "owned_fixture", "is_synthetic": True}
                                def bounded(model_name, logs, lead, name, refit):
                                    fit_started = time.monotonic()
                                    fit = fit_in_owned_child(model_name, logs, lead, name, refit,
                                                             remaining_seconds=remaining)
                                    if time.monotonic() - fit_started > remaining:
                                        raise TimeoutError("fit_deadline")
                                    if fit.backend != "fixture":
                                        raise ValueError("unsupported_unpinned_backend")
                                    if key in backends and fit.backend != backends[key]:
                                        raise ValueError("backend_changed")
                                    backends[key] = fit.backend
                                    provenance["backends"][f"{sid}:{model}:{origin}"] = fit.backend
                                    provenance["versions"].update(getattr(fit, "versions", {}))
                                    host.check()
                                    return fit
                                rows.extend(forecast_from_training(training, origin, model, protocol, spec, split, bounded))
                                provenance["seeds"][f"{sid}:{origin}"] = canonical_seed(protocol, sid, origin)
                                provenance["training_cutoffs"][f"{sid}:{origin}"] = cutoff
                                provenance["source_slice_sha256"][f"{sid}:{origin}"] = _sha(canonical(training))
                                provenance["model_fit_seeds"][f"{sid}:{model}"] = 0 if model == "gbm" else None
                        except (ValueError, RuntimeError, TimeoutError) as exc:
                            reason = str(exc) if str(exc) in {"backend_changed", "unsupported_unpinned_backend", "fit_deadline"} else "fit_unavailable"
                            failed[key] = reason
                            if reason == "fit_deadline" or isinstance(exc, OwnedChildCleanupError):
                                raise
                        else:
                            predictions.extend(rows)
                        finally:
                            elapsed[sid] += time.monotonic() - before
                if not predictions:
                    raise ValueError("all native cohorts unavailable")
                return predictions

            host.freeze("validation", predict("validation"))
            validation = host.attach("validation")
            calibration, calibrated = {}, []
            for sid in ("copper", "aluminum"):
                for model in ("naive", "ets", "gbm"):
                    if (sid, model) in failed:
                        continue
                    cohort = []
                    try:
                        for horizon in (1, 3, 6):
                            factor, adjusted = calibrate_intervals([r for r in validation if
                                (r["series_id"], r["model"], r["horizon"]) == (sid, model, horizon)])
                            calibration[f"{sid}:{model}:{horizon}"] = factor
                            cohort.extend(adjusted)
                    except ValueError:
                        failed[(sid, model)] = "calibration_unavailable"
                        calibration = {key: value for key, value in calibration.items()
                                       if not key.startswith(f"{sid}:{model}:")}
                    else:
                        calibrated.extend(cohort)
            calibrated = [r for r in calibrated if (r["series_id"], r["model"]) not in failed
                          and (r["series_id"], "naive") not in failed]
            selection = {f"{sid}:{h}": value for (sid, h), value in select_models(calibrated).items()}
            host.freeze_selection(selection, calibration, {f"{sid}:{m}": reason for (sid, m), reason in failed.items()})
            host.freeze("test", predict("test"))
            test = host.attach("test")
            scored = calibrated + [_scaled_interval(r, calibration[f"{r['series_id']}:{r['model']}:{r['horizon']}"])
                                    for r in test]
            available = {(r["series_id"], r["model"]) for r in test}
            scored = [r for r in scored if (r["series_id"], r["model"]) in available
                      and (r["series_id"], "naive") not in failed]
            for entry in ledger:
                if entry["status"] == "pending":
                    key = (entry["series_id"], entry["model"])
                    if key in failed or (key[0], "naive") in failed:
                        entry.update(status="not_evaluated", reason=failed.get(key, "naive_baseline_unavailable"))
                    elif key in available:
                        entry.update(status="scored", backend=backends[key], reason=None)
                    else:
                        entry.update(status="not_evaluated", reason="missing_cohort")
            if not scored:
                raise ValueError("all scored cohorts unavailable")
            _exclusive_json(output / "candidate_ledger.json", {"entries": ledger, "provider_dl": "out_of_scope"})
            _exclusive_json(output / "scored_rows.json", {"context": context, "predictions": scored})
            # Existing schema is unchanged, explicitly nested and fixture marked.
            doc = {"schema_version": "forecast-scorecard.v1", "metadata": {"synthetic": True,
                    "latest_vintage_proxy": True, "observed_execution_authorized": False}, "predictions": scored}
            publish_runner(output / "legacy", doc, ledger, selection, calibration, provenance)
            verify_runner_run(output / "legacy")
            host.check()
        # Source/control/code close and postclose checks succeeded before closure.
        artifacts = _admitted_artifacts(output)
        read.lifetime.check(postclose=True)
        info = os.lstat(output)
        if (info.st_dev, info.st_ino, info.st_mode, info.st_uid) != host.identity or os.path.lexists(output / "FAILURE.json"):
            raise ValueError("late output/failure identity changed")
        closure = {"schema_version": _ADMITTED_SCHEMA, "status": "fixture_complete", "context": context,
                   "is_synthetic": True, "observed_authenticity": False, "observed_execution_authorized": False,
                   "full_parent_snapshot_exposed": True, "physical_test_actuals_untouched": False,
                   "close_postclose_verified": True, "artifact_sha256": artifacts}
        _exclusive_json(output / "ADMITTED_COMPLETE.json", closure)
        _seal_fixture_run(binding, token, canonical(closure))
        _verify_admitted_fixture(output, binding)
        return closure
    except BaseException as exc:
        _fail_fixture_run(binding, token)
        try:
            still_owned = (owned_output_identity is not None
                           and _identity(os.lstat(output), True) == owned_output_identity
                           and _identity(os.lstat(output.parent), True) == run["parent_identity"])
        except OSError:
            still_owned = False
        if still_owned and not os.path.lexists(output / "FAILURE.json"):
            reason = ("fit_interrupted" if isinstance(exc, (InterruptedError, KeyboardInterrupt))
                      else "owned_child_cleanup_unverified" if isinstance(exc, OwnedChildCleanupError)
                      else "fit_deadline" if isinstance(exc, TimeoutError)
                      else "admitted_execution_unverified")
            _exclusive_json(output / "FAILURE.json", {"schema_version": _ADMITTED_SCHEMA, "status": "incomplete",
                                                      "reason": reason})
            if not os.path.lexists(output / "candidate_ledger.json"):
                interrupted_ledger = [{**item, "status": "not_evaluated", "reason": item["reason"] or reason}
                                      for item in ledger]
                _exclusive_json(output / "candidate_ledger.json", {"entries": interrupted_ledger,
                                                                   "provider_dl": "out_of_scope"})
        if isinstance(exc, (InterruptedError, KeyboardInterrupt)):
            raise RuntimeError("runner incomplete: fit_interrupted") from None
        if isinstance(exc, OwnedChildCleanupError):
            raise RuntimeError("runner incomplete: owned_child_cleanup_unverified") from None
        raise


def _admitted_artifacts(output):
    from .observed_admission import Lifetime
    hashes = {}
    life = Lifetime()
    try:
        for path in sorted(output.rglob("*")):
            if path.is_symlink():
                raise ValueError("artifact symlink")
            if path.is_dir():
                continue
            if path.name == "ADMITTED_COMPLETE.json" and path.parent == output:
                continue
            fd = life.open(path, limit=16 * 1024 * 1024)
            hashes[str(path.relative_to(output))] = _sha(life.read(fd, 16 * 1024 * 1024))
    finally:
        life.close()
    return hashes


def _verify_admitted_fixture(output_dir, binding):
    from .observed_admission import fixture_run_record, _fixture_completion_identity, _RunDirectoryAnchor
    run = fixture_run_record(binding)
    output = Path(output_dir)
    if str(output) != run["output"]:
        raise ValueError("foreign output binding")
    if os.path.lexists(output / "FAILURE.json"):
        raise ValueError("failure entry dominates complete")
    identities = _fixture_completion_identity(binding)
    with _RunDirectoryAnchor(output, identities) as anchor:
        result = _verify_admitted_fixture_artifacts(output, binding)
        anchor.check()  # after final artifact/closure checks, while still open
    return result  # close/postclose uncertainty cannot return success


def _verify_admitted_fixture_artifacts(output_dir, binding):
    from .observed_admission import fixture_run_record, fixture_admission_record, _identity, _fixture_completion_sha
    run = fixture_run_record(binding)
    admission = fixture_admission_record(run["admission"])
    output = Path(output_dir)
    if str(output) != run["output"] or output.is_symlink():
        raise ValueError("foreign output binding")
    if os.path.lexists(output / "FAILURE.json"):
        raise ValueError("failure entry dominates complete")
    for path, identity, directory in admission["identities"]:
        if _identity(os.lstat(path), directory) != identity:
            raise ValueError("source/control/code identity changed")
    closure, closure_raw = _owned_json(output / "ADMITTED_COMPLETE.json")
    if _sha(closure_raw) != _fixture_completion_sha(binding):
        raise ValueError("host-issued closure bytes mismatch")
    if (closure.get("schema_version") != _ADMITTED_SCHEMA or closure.get("status") != "fixture_complete"
            or closure.get("is_synthetic") is not True or closure.get("observed_authenticity") is not False
            or closure.get("observed_execution_authorized") is not False
            or closure.get("close_postclose_verified") is not True
            or closure.get("full_parent_snapshot_exposed") is not True
            or closure.get("physical_test_actuals_untouched") is not False):
        raise ValueError("admitted outer closure incomplete")
    context = closure.get("context", {})
    if (context.get("fixture_id") != binding.fixture_id or context.get("source_sha256") != admission["grant"]["source_sha256"]
            or context.get("code_sha256") != dict(admission["code"])
            or context.get("protocol_sha256") != dict(admission["controls"])[run["protocol"]]
            or context.get("manifest_sha256") != dict(admission["controls"])[run["manifest"]]
            or context.get("output") != str(output) or context.get("as_of") != "2026-10-08"):
        raise ValueError("foreign closure context")
    hashes = _admitted_artifacts(output)
    required = {"HOST_PENDING.json", "validation_prediction_freeze.json", "selection_freeze.json",
                "test_prediction_freeze.json", "candidate_ledger.json", "scored_rows.json", "source_receipt.json"}
    if not required.issubset(hashes) or hashes != closure.get("artifact_sha256"):
        raise ValueError("admitted closure artifact mismatch")
    for name in ("HOST_PENDING.json", "validation_prediction_freeze.json", "selection_freeze.json",
                 "test_prediction_freeze.json", "scored_rows.json"):
        doc, _ = _owned_json(output / name)
        if doc.get("context") != context:
            raise ValueError("foreign causal artifact")
    selection, _ = _owned_json(output / "selection_freeze.json")
    if selection.get("validation_prediction_sha256") != hashes["validation_prediction_freeze.json"]:
        raise ValueError("selection does not bind validation freeze")
    verify_runner_run(output / "legacy")
    if _admitted_artifacts(output) != hashes or os.path.lexists(output / "FAILURE.json"):
        raise ValueError("late closure artifact/failure change")
    return True


def _legacy_run_offline_evaluation(protocol_path, source_manifest_path, source_file, output_dir, as_of):
    """Compatibility name: unbound direct execution has no authority."""
    raise ValueError("legacy runner unavailable")


def _run_offline_evaluation_inner(protocol_path: str | Path, source_manifest_path: str | Path,
                                  source_file: str | Path, output_dir: str | Path, as_of: str) -> dict:
    """Compatibility name: only accepted_run-bound evaluation may execute."""
    raise ValueError("legacy runner unavailable")


# Explicit native transport. Legacy --fit-child payload/error behavior is retained.
class NativeFitUnavailable(RuntimeError):
    def __init__(self, reason, request_sha256, backend):
        super().__init__(reason)
        self._outcome = (reason, 'backend_precheck', request_sha256, backend)

    reason = property(lambda self: self._outcome[0])
    stage = property(lambda self: self._outcome[1])
    request_sha256 = property(lambda self: self._outcome[2])
    backend = property(lambda self: self._outcome[3])


class NativeChildProtocolError(RuntimeError):
    pass


class NativeChildExecutionError(RuntimeError):
    pass


class NativeBackendIntegrityError(RuntimeError):
    pass


_NATIVE_VERSION_KEYS = frozenset(('numpy', 'pandas', 'openpyxl', 'scikit-learn', 'statsmodels', 'lightgbm'))


def _native_versions(value):
    from collections.abc import Mapping
    if (not isinstance(value, Mapping) or set(value) != _NATIVE_VERSION_KEYS or
            any(not isinstance(v, str) or not 0 < len(v) <= 128 or any(ord(c) < 32 for c in v) for v in value.values())):
        raise NativeChildProtocolError('native exact six versions required')
    return dict(value)


def _native_finite(value):
    return type(value) in (int, float) and math.isfinite(value)


def _native_request(raw):
    if not isinstance(raw, bytes) or len(raw) > 1024 * 1024:
        raise NativeChildProtocolError('native request bound')
    try:
        req = json.loads(raw, object_pairs_hook=_unique_pairs, parse_constant=_no_constant)
        if not isinstance(req, dict) or set(req) != {'schema_version', 'model', 'name', 'months', 'values', 'expected_versions'}:
            raise ValueError('request keys')
        if req['schema_version'] != 'observed-native-fit-request.v1' or req['model'] not in ('naive', 'ets', 'gbm'):
            raise ValueError('request version/model')
        if not isinstance(req['name'], str) or not 0 < len(req['name']) <= 256 or any(ord(c) < 32 for c in req['name']):
            raise ValueError('request name')
        labels, values = req['months'], req['values']
        if (not isinstance(labels, list) or not 60 <= len(labels) <= 113 or
                labels != _range('2016-09', _name(_month('2016-09') + len(labels) - 1)) or
                not isinstance(values, list) or len(values) != len(labels) or not all(_native_finite(v) for v in values)):
            raise ValueError('past log-series shape')
        _native_versions(req['expected_versions'])
        return req
    except (ValueError, TypeError, OverflowError) as exc:
        raise NativeChildProtocolError('native request invalid') from exc


def _decode_native_fit_result(raw, *, request_raw, expected_model, expected_training_count, expected_versions):
    from types import SimpleNamespace
    expected_versions = _native_versions(expected_versions)
    try:
        if not isinstance(raw, bytes) or len(raw) > 1024 * 1024:
            raise ValueError('result bound')
        value = json.loads(raw, object_pairs_hook=_unique_pairs, parse_constant=_no_constant)
        if not isinstance(value, dict):
            raise ValueError('result object')
        status = value.get('status')
        common = {'schema_version', 'status', 'request_sha256', 'model'}
        extras = {'ok': {'drift', 'residuals', 'noise_scale', 'backend', 'versions', 'model_fit_attempted'},
                  'unavailable': {'reason', 'stage', 'backend', 'expected_backend', 'versions', 'model_fit_attempted'},
                  'failed': {'reason', 'stage'}}
        if not isinstance(status, str) or status not in extras or set(value) != common | extras[status]:
            raise ValueError('result exact status/key set')
        if (value['schema_version'] != 'observed-native-fit-result.v1' or value['request_sha256'] != _sha(request_raw)
                or value['model'] != expected_model):
            raise ValueError('result foreign job/version')
        if status == 'failed':
            if (value['reason'], value['stage']) not in {
                    ('protocol_fault', 'request_validation'), ('runtime_binding_mismatch', 'runtime_validation'),
                    ('invalid_fit_result', 'result_validation'), ('unexpected_child_exception', 'execution')}:
                raise ValueError('failed reason/stage invalid')
            raise NativeChildExecutionError(value['reason'])
        versions = _native_versions(value['versions'])
        if versions != expected_versions:
            raise NativeBackendIntegrityError('native version mismatch')
        if status == 'unavailable':
            if (expected_model, value['reason'], value['stage'], value['backend'], value['expected_backend']) != (
                    'gbm', 'unsupported_unpinned_backend', 'backend_precheck', 'lightgbm', 'sklearn') or value['model_fit_attempted'] is not False:
                raise ValueError('unavailable causal tuple invalid')
            raise NativeFitUnavailable(value['reason'], value['request_sha256'], value['backend'])
        if value['backend'] != {'naive': 'naive', 'ets': 'ets', 'gbm': 'sklearn'}[expected_model]:
            raise NativeBackendIntegrityError('native backend mismatch')
        drift, residuals, noise = value['drift'], value['residuals'], value['noise_scale']
        if (not isinstance(drift, list) or len(drift) != 7 or not all(_native_finite(v) for v in drift)
                or not isinstance(residuals, list) or not 12 <= len(residuals) <= min(120, expected_training_count)
                or not all(_native_finite(v) for v in residuals) or not _native_finite(noise) or noise <= 0
                or value['model_fit_attempted'] is not True):
            raise ValueError('result array/noise shape invalid')
        return SimpleNamespace(drift=np.asarray(drift), residuals=np.asarray(residuals), noise_scale=noise,
                               backend=value['backend'], versions=versions)
    except (ValueError, TypeError, OverflowError, KeyError) as exc:
        raise NativeChildProtocolError('native malformed result') from exc


def _fit_in_native_child(model, logs, horizon, name, refit, remaining, versions):
    if horizon != 7 or refit != 1 or model not in ('naive', 'ets', 'gbm'):
        raise NativeChildProtocolError('native fit protocol mismatch')
    payload = json.dumps({'schema_version': 'observed-native-fit-request.v1', 'model': model, 'name': name,
                          'months': [v.strftime('%Y-%m') for v in logs.index],
                          'values': [float(v) for v in logs.values], 'expected_versions': _native_versions(versions)}, allow_nan=False).encode()
    request = _native_request(payload)  # validate before transport; no fallback
    response = run_bounded_child([sys.executable, '-m', 'procure.evaluation.observed_runner', '--native-fit-child-v1'],
                                 payload, timeout=min(60., remaining))
    return _decode_native_fit_result(response, request_raw=payload, expected_model=model,
                                     expected_training_count=len(request['months']), expected_versions=versions)


def _native_fit_child_main():
    # Entire native body, including import/serialization, has hard exception semantics.
    raw = sys.stdin.buffer.read(1024 * 1024 + 1)
    model = None
    common = {'schema_version': 'observed-native-fit-result.v1', 'request_sha256': _sha(raw), 'model': None}
    try:
        req = _native_request(raw)
        model = req['model']; common['model'] = model
    except Exception:
        response = {**common, 'status': 'failed', 'reason': 'protocol_fault', 'stage': 'request_validation'}
    else:
        try:
            from importlib.metadata import version
            versions = {k: version(k) for k in _NATIVE_VERSION_KEYS}
            if versions != req['expected_versions']:
                response = {**common, 'status': 'failed', 'reason': 'runtime_binding_mismatch', 'stage': 'runtime_validation'}
            else:
                from procure.forecast import models
                if model == 'gbm' and models.GBM_BACKEND == 'lightgbm':
                    response = {**common, 'status': 'unavailable', 'reason': 'unsupported_unpinned_backend',
                                'stage': 'backend_precheck', 'backend': 'lightgbm', 'expected_backend': 'sklearn',
                                'versions': versions, 'model_fit_attempted': False}
                else:
                    if model == 'gbm' and models.GBM_BACKEND != 'sklearn':
                        raise NativeBackendIntegrityError('unknown native backend')
                    series = pd.Series(req['values'], index=pd.to_datetime([m + '-01' for m in req['months']]))
                    models.clear_cache()
                    fit = models.fit_series(model, series, 7, name=req['name'], refit_every=1)
                    if not _native_finite(fit.noise_scale) or fit.noise_scale <= 0:
                        raise NativeChildProtocolError('native fitted noise invalid')
                    response = {**common, 'status': 'ok', 'drift': fit.drift.tolist(),
                                'residuals': np.asarray(fit.residuals).tolist(), 'noise_scale': float(fit.noise_scale),
                                'backend': models.GBM_BACKEND if model == 'gbm' else model,
                                'versions': {k: version(k) for k in _NATIVE_VERSION_KEYS}, 'model_fit_attempted': True}
                    try:
                        _decode_native_fit_result(json.dumps(response, allow_nan=False).encode(), request_raw=raw,
                                                  expected_model=model, expected_training_count=len(req['months']), expected_versions=versions)
                    except NativeChildProtocolError:
                        response = {**common, 'status': 'failed', 'reason': 'invalid_fit_result', 'stage': 'result_validation'}
        except Exception:
            response = {**common, 'status': 'failed', 'reason': 'unexpected_child_exception', 'stage': 'execution'}
    try:
        sys.stdout.write(json.dumps(response, allow_nan=False))
        return 0
    except Exception:
        return 1  # serialization/stdout uncertainty never unavailable

class _NativeBudgetUnavailable(RuntimeError):
    pass


class _NativeBudget:
    def __init__(self):
        self.started = time.monotonic()
        self.series = defaultdict(float)
        self.models = defaultdict(float)
        self.phases = defaultdict(float)
        self.attempts = []
        self.seen = set()

    def check(self):
        if time.monotonic() - self.started >= 1800 or any(v >= 900 for v in self.series.values()):
            raise TimeoutError('native hard deadline')

    def charge(self, sid, model, split, elapsed):
        if elapsed < 0:
            raise RuntimeError('native monotonic clock reversed')
        self.series[sid] += elapsed
        self.models[sid, model] += elapsed
        self.phases[sid, model, split] += elapsed
        self.check()
        caps = {'naive': 120, 'ets': 360, 'gbm': 360}
        phases = {'validation': {'naive': 48, 'ets': 144, 'gbm': 144}, 'test': {'naive': 72, 'ets': 216, 'gbm': 216}}
        if self.models[sid, model] > caps[model] or self.phases[sid, model, split] > phases[split][model]:
            raise TimeoutError('native actual model/phase overrun')

    def remaining(self, sid, model, split):
        self.check()
        caps = {'naive': 120, 'ets': 360, 'gbm': 360}
        phase = {'validation': {'naive': 48, 'ets': 144, 'gbm': 144}, 'test': {'naive': 72, 'ets': 216, 'gbm': 216}}
        seconds = min(60., caps[model] - self.models[sid, model] - 3,
                      phase[split][model] - self.phases[sid, model, split] - 3,
                      900 - self.series[sid] - 3, 1800 - (time.monotonic() - self.started) - 33)
        if seconds <= 0:
            raise _NativeBudgetUnavailable('model_phase_budget_unavailable')
        return seconds

    def begin_attempt(self, sid, model, split, origin, timeout):
        key = (sid, model, split, origin)
        if key in self.seen or len(self.seen) >= 288:
            raise RuntimeError('native attempt reuse/ceiling')
        self.seen.add(key)
        item = {'series_id': sid, 'model': model, 'split': split, 'origin': origin,
                'planned_timeout': timeout, 'elapsed_seconds': None, 'status': 'running',
                'reason': None, 'expected_backend': {'naive': 'naive', 'ets': 'ets', 'gbm': 'sklearn'}[model],
                'backend': None, 'cleanup_verified': False}
        self.attempts.append(item)
        return item


def _native_attempt_document(budget, context, failed, terminal=None):
    skipped = []
    for split in ('validation', 'test'):
        for sid in ('copper', 'aluminum'):
            for model in ('naive', 'ets', 'gbm'):
                for origin in _range(MONTHS[split + '_start'], MONTHS[split + '_end']):
                    if (sid, model, split, origin) not in budget.seen:
                        skipped.append({'series_id': sid, 'model': model, 'split': split, 'origin': origin,
                            'status': 'skipped', 'reason': terminal or failed.get((sid, model), 'naive_baseline_unavailable'),
                            'expected_backend': {'naive': 'naive', 'ets': 'ets', 'gbm': 'sklearn'}[model],
                            'backend': None, 'planned_timeout': None, 'elapsed_seconds': None, 'cleanup_verified': False})
    return {'context': context, 'attempts': budget.attempts, 'skipped': skipped,
            'series_seconds': dict(budget.series), 'elapsed_seconds': time.monotonic() - budget.started}


class _NativePhases(_FixturePhases):
    def __init__(self, output, read, context, run, budget, output_anchor):
        super().__init__(output, read, context)
        self.run, self.budget = run, budget
        self.output_anchor = output_anchor
        self.attached = set()

    def check(self):
        super().check()
        self.output_anchor.check()
        self.budget.check()
        from .observed_admission import _native_alias_guard, _identity
        _native_alias_guard(self.run['control']['runtime'])
        if _identity(os.lstat(self.output.parent), True) != self.run['parent_identity']:
            raise ValueError('native output parent identity changed')

    def freeze(self, split, rows):
        groups = defaultdict(set)
        from .forecast_scorecard import FIELDS
        for row in rows:
            if (set(row) != FIELDS - {'actual'} or row['series_id'] not in ('copper', 'aluminum')
                    or row['model'] not in ('naive', 'ets', 'gbm') or row['horizon'] not in (1, 3, 6)
                    or type(row['horizon']) is not int or row['unit'] != 'USD/metric tonne'
                    or row['training_cutoff'] != _name(_month(row['origin']) - 1)):
                raise ValueError('native freeze row shape/calendar mismatch')
            groups[row['series_id'], row['model']].add(row['horizon'])
        if not groups or any(h != {1, 3, 6} for h in groups.values()):
            raise ValueError('native complete three-horizon cohort required')
        if split == 'test' and 'validation' not in self.attached:
            raise ValueError('native validation attachment absent')
        super().freeze(split, rows)

    def attach(self, split):
        if split in self.attached:
            raise ValueError('native actual attachment already consumed')
        rows = super().attach(split)
        self.attached.add(split)
        return rows

    def freeze_selection(self, selections, calibration, failed):
        if 'validation' not in self.attached:
            raise ValueError('native selection before validation attachment')
        super().freeze_selection(selections, calibration, failed)


def _native_artifacts(output):
    from .observed_admission import _NativeLifetime
    life = _NativeLifetime(); hashes = {}; total = 0
    try:
        paths = sorted(output.rglob('*'))
        if len(paths) > 132:
            raise ValueError('native artifact count bound')
        for path in paths:
            info = os.lstat(path)
            import stat
            if stat.S_ISDIR(info.st_mode):
                if str(path.relative_to(output)) not in {'legacy', 'legacy/charts', 'legacy/scorecard'}:
                    raise ValueError('native unexpected artifact directory')
                continue
            if path.name == 'NATIVE_COMPLETE.json' and path.parent == output:
                continue
            fd = life.open(path, limit=16 * 1024 * 1024)
            raw = life.read(fd, 16 * 1024 * 1024)
            total += len(raw)
            if total > 64 * 1024 * 1024 or len(hashes) >= 128:
                raise ValueError('native artifact byte/file bound')
            hashes[str(path.relative_to(output))] = _sha(raw)
        life.check()
    finally:
        life.close()
    return hashes


def _native_expected_artifacts(rows):
    names = {'HOST_PENDING.json', 'source_receipt.json', 'validation_prediction_freeze.json', 'selection_freeze.json',
             'test_prediction_freeze.json', 'candidate_ledger.json', 'scored_rows.json', 'attempts.json',
             'legacy/RUN_STATE.json', 'legacy/predictions_input.json', 'legacy/candidate_ledger.json',
             'legacy/selection.json', 'legacy/calibration.json', 'legacy/provenance.json'}
    names.update('legacy/scorecard/' + n for n in (*ARTIFACTS, 'RUN_STATE.json'))
    for r in rows:
        group = tuple(r[k] for k in ('series_kind', 'series_id', 'unit', 'source_id', 'model', 'backend', 'split', 'horizon'))
        names.add('legacy/charts/' + _sha(json.dumps(group, ensure_ascii=True).encode())[:16] + '.svg')
    return names


def _verify_root_native(output_dir, binding):
    from contextlib import closing
    from .observed_admission import (native_run_record, _native_completion, _RunDirectoryAnchor,
                                     _native_pins_lifetime, _fail_native_run, _native_state)
    run = native_run_record(binding)  # original issued live identity before paths
    ids, issued_sha = _native_completion(binding)
    output = Path(output_dir)
    if str(output) != run['output']:
        raise ValueError('native foreign output')  # a foreign argument cannot invalidate the original
    try:
        if os.path.lexists(output / 'FAILURE.json'):
            raise ValueError('native failure entry')
        # Metadata-only source; never workbook body. Close uncertainty is terminal.
        with closing(_native_pins_lifetime(run)) as pins, _RunDirectoryAnchor(output, ids) as anchor:
            closure, raw = _owned_json(output / 'NATIVE_COMPLETE.json')
            c = run['control']; observed = c['profile'] == 'native_observed'
            statuses = ('native_complete', 'native_complete_with_unavailable') if observed else (
                'native_fixture_complete', 'native_fixture_complete_with_unavailable')
            if (_sha(raw) != issued_sha or closure.get('schema_version') != 'root-native-host-closure.v1'
                    or closure.get('status') not in statuses or closure.get('context') != _native_context(run)
                    or closure.get('is_synthetic') is not (not observed) or closure.get('observed_authenticity') is not observed
                    or closure.get('observed_execution_authorized') is not observed
                    or closure.get('full_parent_snapshot_exposed') is not True
                    or closure.get('physical_test_actuals_untouched') is not False
                    or closure.get('close_postclose_verified') is not True):
                raise ValueError('native closure schema/issued context mismatch')
            hashes = _native_artifacts(output)
            scored, _ = _owned_json(output / 'scored_rows.json')
            if set(hashes) != _native_expected_artifacts(scored['predictions']) or hashes != closure['artifact_sha256']:
                raise ValueError('native exact artifact set/hash mismatch')
            context = closure['context']
            for name in ('HOST_PENDING.json', 'source_receipt.json', 'validation_prediction_freeze.json',
                         'selection_freeze.json', 'test_prediction_freeze.json', 'scored_rows.json', 'attempts.json'):
                doc, _ = _owned_json(output / name)
                if doc.get('context') != context:
                    raise ValueError('native foreign artifact context')
            selection, _ = _owned_json(output / 'selection_freeze.json')
            ledger, _ = _owned_json(output / 'candidate_ledger.json')
            attempts, _ = _owned_json(output / 'attempts.json')
            history = attempts['attempts'] + attempts['skipped']
            planned = {(s, m, split, o) for s in ('copper', 'aluminum') for m in ('naive', 'ets', 'gbm')
                       for split in ('validation', 'test') for o in _range(MONTHS[split + '_start'], MONTHS[split + '_end'])}
            if (len(history) != 288 or {(r['series_id'],r['model'],r['split'],r['origin']) for r in history} != planned
                    or any(r['status'] not in ('ok','unavailable') or r['cleanup_verified'] is not True for r in attempts['attempts'])
                    or any(r['status'] != 'skipped' or r['elapsed_seconds'] is not None or r['cleanup_verified'] is not False
                           for r in attempts['skipped'])):
                raise ValueError('native exact attempt/skip agreement required')
            if (selection['validation_prediction_sha256'] != hashes['validation_prediction_freeze.json']
                    or len(ledger['entries']) != 96 or len(attempts['attempts']) > 288
                    or closure['candidate_ledger_sha256'] != hashes['candidate_ledger.json']
                    or closure['attempts_sha256'] != hashes['attempts.json']
                    or closure['validation_prediction_sha256'] != hashes['validation_prediction_freeze.json']
                    or closure['selection_sha256'] != hashes['selection_freeze.json']
                    or closure['test_prediction_sha256'] != hashes['test_prediction_freeze.json']):
                raise ValueError('native cohort/ledger/phase receipt mismatch')
            # Fresh verifier consumes existing artifacts, never re-fits/re-scores.
            verify_runner_run(output / 'legacy')
            if _native_artifacts(output) != hashes or os.path.lexists(output / 'FAILURE.json'):
                raise ValueError('native late artifact/failure drift')
            pins.check(); anchor.check()
        if _native_completion(binding) != (ids, issued_sha):
            raise ValueError('native live issued seal changed after close')
    except BaseException:
        _fail_native_run(binding, _native_state(binding)['token'])
        raise
    return True


def _native_context(run):
    from .observed_admission import _native_plain, canonical
    c = run['control']
    return {'profile': c['profile'], 'acceptance_id': c['acceptance_id'], 'run_id': c['run_id'],
            'root_control_sha256': run['root_pin']['sha256'], 'source_sha256': c['source']['file']['sha256'],
            'protocol_sha256': c['controls']['protocol']['sha256'], 'manifest_sha256': c['controls']['manifest']['sha256'],
            'code_sha256': {p['repository_path']: p['file']['sha256'] for p in c['code']},
            'runtime_pin_sha256': _sha(canonical(_native_plain(c['runtime']))), 'output': run['output'],
            'as_of': '2026-10-08', 'analysis_as_of_utc': c['analysis_as_of_utc'],
            'budget_policy_sha256': _sha(canonical(_native_plain(c['budgets'])))}


def _run_root_native(protocol_path, manifest_path, source_file, output_dir, as_of, binding):
    from contextlib import ExitStack
    from .observed_admission import (native_run_record, _native_plain, canonical, decode, _identity,
                                     _begin_native_run, _capture_native_output, _seal_native_run, _fail_native_run,
                                     open_root_native_snapshot, _RunDirectoryAnchor)
    run = native_run_record(binding)
    c = run['control']; observed = c['profile'] == 'native_observed'
    if tuple(map(str, (protocol_path, manifest_path, source_file, output_dir, as_of))) != (
            c['controls']['protocol']['path'], c['controls']['manifest']['path'], c['source']['file']['path'], run['output'], '2026-10-08'):
        raise ValueError('native foreign run/source/control/output arguments')
    if (observed and fit_in_owned_child is not _REAL_FIT_BOUNDARY) or (not observed and fit_in_owned_child is _REAL_FIT_BOUNDARY):
        raise ValueError('native mode/fitter identity mismatch before source body')
    token = _begin_native_run(binding)
    output = Path(run['output']); budget = _NativeBudget(); made = None
    ledger, failed, backends = _fixture_ledger(), {}, {}
    context = _native_context(run)
    versions = {p['distribution']: p['version'] for p in c['runtime']['packages']}
    provenance = {'source_manifest_sha256': c['controls']['manifest']['sha256'], 'protocol_sha256': c['controls']['protocol']['sha256'],
                  'seeds': {}, 'backends': {}, 'versions': versions, 'training_cutoffs': {}, 'source_slice_sha256': {},
                  'source_snapshot_sha256': c['source']['file']['sha256'], 'model_fit_seeds': {},
                  'source_eligibility_status': 'APPROVED_OBSERVED' if observed else 'OWNED_SYNTHETIC_FIXTURE',
                  'full_parent_snapshot_exposed': True, 'physical_test_actuals_untouched': False,
                  'observed_authenticity': observed, 'observed_execution_authorized': observed, 'latest_vintage_proxy': True}
    try:
        if os.path.lexists(output) or _identity(os.lstat(output.parent), True) != run['parent_identity']:
            raise ValueError('native output absence/parent drift')
        with open_root_native_snapshot(binding) as read, ExitStack() as output_lifetime:
            p = {**PROTOCOL_FIELDS, 'status': 'APPROVED_OBSERVED_EXECUTION' if observed else 'INFRASTRUCTURE_ONLY_OBSERVED_EXECUTION_HOLD',
                 'observed_execution_authorized': observed, 'protocol_sha256': c['controls']['protocol']['sha256'],
                 'source_sha256': c['source']['file']['sha256'], 'snapshot_sha256': c['source']['file']['sha256']}
            windows = plan_common_origins({sid: dict(s) for sid, s in read.series.items()}, as_of, p)
            budget.check(); read.check()
            output.mkdir(mode=0o700, exist_ok=False); made = _identity(os.lstat(output), True)
            _capture_native_output(binding, token, made)
            output_anchor = output_lifetime.enter_context(_RunDirectoryAnchor(output, (run['parent_identity'], made)))
            host = _NativePhases(output, read, context, run, budget, output_anchor)
            _exclusive_json(output / 'HOST_PENDING.json', {'schema_version': 'root-native-host-closure.v1', 'context': context,
                                                         'status': 'incomplete', 'is_synthetic': not observed})
            _exclusive_json(output / 'source_receipt.json', {**dict(read.receipt), 'context': context})

            def predict(split):
                predictions = []
                for sid in ('copper', 'aluminum'):
                    for model in ('naive', 'ets', 'gbm'):
                        key = (sid, model)
                        if key in failed:
                            continue
                        if (sid, 'naive') in failed:
                            failed[key] = 'naive_baseline_unavailable'; continue
                        rows = []
                        for origin in windows[split]:
                            before = time.monotonic(); item = None; precharged = 0.
                            try:
                                host.check()
                                budget.remaining(sid, model, split)
                                cutoff = _name(_month(origin) - 1)
                                training = {m: read.series[sid][m] for m in _range('2016-09', cutoff)}
                                spec = {'series_kind': 'commodity', 'series_id': sid, 'unit': 'USD/metric tonne',
                                        'source_id': c['source']['source_id'], 'is_synthetic': not observed}
                                def bounded(m, logs, lead, name, refit):
                                    nonlocal item, precharged
                                    # Charge preparation before determining the launch timeout.
                                    precharged = time.monotonic() - before
                                    budget.charge(sid, model, split, precharged)
                                    seconds = budget.remaining(sid, model, split)
                                    item = budget.begin_attempt(sid, model, split, origin, seconds)
                                    started = time.monotonic()
                                    try:
                                        fit = fit_in_owned_child(m, logs, lead, name, refit, remaining_seconds=seconds,
                                                                 result_contract='native-v1', expected_versions=versions)
                                    except NativeFitUnavailable as exc:
                                        host.check()
                                        if key in backends:
                                            raise NativeBackendIntegrityError('native later backend precheck changed') from exc
                                        item.update(status='unavailable', reason=exc.reason, backend=exc.backend, cleanup_verified=True)
                                        raise
                                    finally:
                                        item['elapsed_seconds'] = time.monotonic() - started
                                        if item['elapsed_seconds'] > seconds:
                                            raise TimeoutError('native late fit beyond deadline')
                                    host.check()
                                    expected = {'naive': 'naive', 'ets': 'ets', 'gbm': 'sklearn'}[model] if observed else 'fixture'
                                    if fit.backend != expected or fit.versions != versions or (key in backends and backends[key] != fit.backend):
                                        raise NativeBackendIntegrityError('native returned backend/version drift')
                                    backends[key] = fit.backend
                                    item.update(status='ok', backend=fit.backend, cleanup_verified=True)
                                    provenance['backends'][f'{sid}:{model}:{origin}'] = fit.backend
                                    return fit
                                rows.extend(forecast_from_training(training, origin, model, p, spec, split, bounded))
                                provenance['seeds'][f'{sid}:{origin}'] = canonical_seed(p, sid, origin)
                                provenance['training_cutoffs'][f'{sid}:{origin}'] = cutoff
                                provenance['source_slice_sha256'][f'{sid}:{origin}'] = _sha(canonical(training))
                                provenance['model_fit_seeds'][f'{sid}:{model}'] = 0 if model == 'gbm' else None
                            except (NativeFitUnavailable, _NativeBudgetUnavailable) as exc:
                                failed[key] = exc.reason if isinstance(exc, NativeFitUnavailable) else str(exc)
                                break
                            except BaseException:
                                if item is not None:
                                    item.update(status='failed', reason='native_integrity_unverified', cleanup_verified=False)
                                raise
                            finally:
                                # Charge the remaining wall interval once, including post-fit checks/processing.
                                budget.charge(sid, model, split, time.monotonic() - before - precharged)
                        else:
                            predictions.extend(rows)
                if not predictions:
                    raise RuntimeError('native all full cohorts unavailable; HOLD')
                return predictions

            host.freeze('validation', predict('validation'))
            validation = host.attach('validation'); calibration, calibrated = {}, []
            for sid in ('copper', 'aluminum'):
                for model in ('naive', 'ets', 'gbm'):
                    key = (sid, model)
                    if key in failed or (sid, 'naive') in failed:
                        continue
                    start = time.monotonic(); cohort = []
                    host.check()
                    try:
                        for horizon in (1, 3, 6):
                            factor, adjusted = calibrate_intervals([r for r in validation if (r['series_id'], r['model'], r['horizon']) == (sid, model, horizon)])
                            calibration[f'{sid}:{model}:{horizon}'] = factor; cohort.extend(adjusted)
                    except ValueError:
                        failed[key] = 'calibration_unavailable'
                        calibration = {k: v for k, v in calibration.items() if not k.startswith(f'{sid}:{model}:')}
                    else:
                        calibrated.extend(cohort)
                    finally:
                        budget.series[sid] += time.monotonic() - start; budget.check()
                    host.check()
            calibrated = [r for r in calibrated if (r['series_id'], r['model']) not in failed and (r['series_id'], 'naive') not in failed]
            selection = {f'{sid}:{h}': value for (sid, h), value in select_models(calibrated).items()}
            host.freeze_selection(selection, calibration, {f'{sid}:{m}': reason for (sid, m), reason in failed.items()})
            host.freeze('test', predict('test'))
            test = host.attach('test')
            available = {(r['series_id'], r['model']) for r in test}
            scored = [r for r in calibrated if (r['series_id'], r['model']) in available and (r['series_id'], 'naive') in available]
            scored.extend(_scaled_interval(r, calibration[f"{r['series_id']}:{r['model']}:{r['horizon']}"]) for r in test
                          if (r['series_id'], 'naive') in available)
            if not scored:
                raise RuntimeError('native no complete naive baseline; HOLD')
            for entry in ledger:
                if entry['status'] == 'pending':
                    key = (entry['series_id'], entry['model'])
                    if key in available and (key[0], 'naive') in available:
                        entry.update(status='scored', backend=backends[key], reason=None)
                    else:
                        entry.update(status='not_evaluated', backend=None, reason=failed.get(key, 'naive_baseline_unavailable'))
            _exclusive_json(output / 'candidate_ledger.json', {'entries': ledger, 'provider_dl': 'out_of_scope'})
            _exclusive_json(output / 'scored_rows.json', {'context': context, 'predictions': scored})
            _exclusive_json(output / 'attempts.json', _native_attempt_document(budget, context, failed))
            doc = {'schema_version': 'forecast-scorecard.v1', 'metadata': {'synthetic': not observed,
                   'latest_vintage_proxy': True, 'observed_execution_authorized': observed}, 'predictions': scored}
            publish_runner(output / 'legacy', doc, ledger, selection, calibration, provenance)
            verify_runner_run(output / 'legacy'); host.check()
        read.lifetime.check(postclose=True); budget.check()
        artifacts = _native_artifacts(output)
        if set(artifacts) != _native_expected_artifacts(scored):
            raise ValueError('native publication artifact set mismatch')
        unavailable = [{'series_id': sid, 'model': model, 'reason': reason} for (sid, model), reason in sorted(failed.items())]
        selected_missing = [{'series_id': sid, 'horizon': int(h), 'model': model,
                             'reason': failed.get((sid, model), 'naive_baseline_unavailable')}
                            for k, model in selection.items() for sid, h in [k.split(':')] if (sid, model) not in available]
        status = ('native_complete' if observed else 'native_fixture_complete') + ('_with_unavailable' if unavailable else '')
        closure = {'schema_version': 'root-native-host-closure.v1', 'status': status, 'context': context,
                   'is_synthetic': not observed, 'observed_authenticity': observed, 'observed_execution_authorized': observed,
                   'full_parent_snapshot_exposed': True, 'physical_test_actuals_untouched': False,
                   'close_postclose_verified': True, 'artifact_sha256': artifacts,
                   'complete_cohorts': [{'series_id': sid, 'model': m} for sid, m in sorted(available) if (sid, 'naive') in available],
                   'unavailable_cohorts': unavailable, 'selected_unavailable': selected_missing,
                   'candidate_ledger_sha256': artifacts['candidate_ledger.json'], 'attempts_sha256': artifacts['attempts.json'],
                   'validation_prediction_sha256': artifacts['validation_prediction_freeze.json'],
                   'selection_sha256': artifacts['selection_freeze.json'], 'test_prediction_sha256': artifacts['test_prediction_freeze.json']}
        budget.check()
        _exclusive_json(output / 'NATIVE_COMPLETE.json', closure)
        _seal_native_run(binding, token, canonical(closure)); budget.check()
        _verify_root_native(output, binding); budget.check()
        return closure
    except BaseException as exc:
        _fail_native_run(binding, token)
        try:
            owned = made is not None and _identity(os.lstat(output), True) == made and _identity(os.lstat(output.parent), True) == run['parent_identity']
        except OSError:
            owned = False
        if owned and not os.path.lexists(output / 'FAILURE.json'):
            reason = 'owned_child_cleanup_unverified' if isinstance(exc, OwnedChildCleanupError) else 'fit_deadline' if isinstance(exc, TimeoutError) else 'native_integrity_unverified'
            _exclusive_json(output / 'FAILURE.json', {'schema_version': 'root-native-host-closure.v1', 'status': 'incomplete', 'reason': reason})
            if not os.path.lexists(output / 'candidate_ledger.json'):
                _exclusive_json(output / 'candidate_ledger.json', {'entries': [{**e, 'status': 'not_evaluated', 'backend': None,
                                                      'reason': e['reason'] or reason} for e in ledger], 'provider_dl': 'out_of_scope'})
            if not os.path.lexists(output / 'attempts.json'):
                _exclusive_json(output / 'attempts.json', _native_attempt_document(budget, context, failed, reason))
        raise


def _run_root_generic(protocol_path, manifest_path, source_file, output_dir, as_of, binding):
    from contextlib import ExitStack
    from .observed_admission import (native_run_record, _native_plain, canonical, decode, _identity,
                                     _begin_native_run, _capture_native_output, _seal_native_run, _fail_native_run,
                                     open_root_generic_snapshot, _RunDirectoryAnchor)
    run = native_run_record(binding)
    c = run['control']; observed = c['profile'] == 'generic-observed-v1'
    if (tuple(map(str, (protocol_path, manifest_path, output_dir, as_of))) != (
            c['controls']['protocol']['path'], c['controls']['manifest']['path'], run['output'], '2026-10-09')
            or not isinstance(source_file, dict) or source_file != {k:v['file']['path'] for k,v in c['sources'].items()}):
        raise ValueError('native foreign run/source/control/output arguments')
    if (observed and fit_in_owned_child is not _REAL_FIT_BOUNDARY) or (not observed and fit_in_owned_child is _REAL_FIT_BOUNDARY):
        raise ValueError('native mode/fitter identity mismatch before source body')
    token = _begin_native_run(binding)
    output = Path(run['output']); budget = _GenericBudget(c['selected_ids']); made = None
    ledger, failed, backends = [], {}, {}
    context = _generic_context(run)
    versions = {p['distribution']: p['version'] for p in c['runtime']['packages']}
    provenance = {'source_manifest_sha256': c['controls']['manifest']['sha256'], 'protocol_sha256': c['controls']['protocol']['sha256'],
                  'seeds': {}, 'backends': {}, 'versions': versions, 'training_cutoffs': {}, 'source_slice_sha256': {},
                  'source_snapshot_sha256': c['catalog']['sha256'], 'model_fit_seeds': {},
                  'source_eligibility_status': 'APPROVED_OBSERVED' if observed else 'OWNED_SYNTHETIC_FIXTURE',
                  'full_parent_snapshot_exposed': True, 'physical_test_actuals_untouched': False,
                  'observed_authenticity': observed, 'observed_execution_authorized': observed, 'latest_vintage_proxy': True}
    try:
        ledger = _generic_ledger(run)
        if os.path.lexists(output) or _identity(os.lstat(output.parent), True) != run['parent_identity']:
            raise ValueError('native output absence/parent drift')
        with open_root_generic_snapshot(binding) as read, ExitStack() as output_lifetime:
            p = {**PROTOCOL_FIELDS, 'as_of':'2026-10-09', 'profile':c['profile'], 'status': 'APPROVED_OBSERVED_EXECUTION' if observed else 'INFRASTRUCTURE_ONLY_OBSERVED_EXECUTION_HOLD',
                 'observed_execution_authorized': observed, 'protocol_sha256': c['controls']['protocol']['sha256'],
                 'source_sha256': c['catalog']['sha256'], 'snapshot_sha256': c['catalog']['sha256']}
            windows = plan_common_origins({sid: dict(s) for sid, s in read.series.items()}, as_of, p)
            budget.check(); read.check()
            output.mkdir(mode=0o700, exist_ok=False); made = _identity(os.lstat(output), True)
            _capture_native_output(binding, token, made)
            output_anchor = output_lifetime.enter_context(_RunDirectoryAnchor(output, (run['parent_identity'], made)))
            host = _GenericPhases(output, read, context, run, budget, output_anchor)
            host.binding = binding
            _exclusive_json(output / 'HOST_PENDING.json', {'schema_version': 'root-generic-host-closure.v1', 'context': context,
                                                         'status': 'incomplete', 'is_synthetic': not observed})
            _exclusive_json(output / 'source_receipt.json', {**dict(read.receipt), 'context': context})

            def predict(split):
                predictions = []
                for sid in c['selected_ids']:
                    for model in ('naive', 'ets', 'gbm'):
                        key = (sid, model)
                        if key in failed:
                            continue
                        if (sid, 'naive') in failed:
                            failed[key] = 'naive_baseline_unavailable'; continue
                        rows = []
                        for origin in windows[split]:
                            before = time.monotonic(); item = None; precharged = 0.
                            try:
                                host.check()
                                budget.remaining(sid, model, split)
                                cutoff = _name(_month(origin) - 1)
                                training = {m: read.series[sid][m] for m in _range('2016-09', cutoff)}
                                record = next(r for r in read.receipt['catalog'] if r['series_id']==sid)
                                spec = {'series_kind': record['series_kind'], 'series_id': sid, 'unit': record['quoted_unit'],
                                        'source_id': record['source_id'], 'is_synthetic': not observed}
                                def bounded(m, logs, lead, name, refit):
                                    nonlocal item, precharged
                                    # Charge preparation before determining the launch timeout.
                                    precharged = time.monotonic() - before
                                    budget.charge(sid, model, split, precharged)
                                    seconds = budget.remaining(sid, model, split)
                                    item = budget.begin_attempt(sid, model, split, origin, seconds)
                                    started = time.monotonic()
                                    try:
                                        fit = fit_in_owned_child(m, logs, lead, name, refit, remaining_seconds=seconds,
                                                                 result_contract='native-v1', expected_versions=versions)
                                    except NativeFitUnavailable as exc:
                                        host.check()
                                        if key in backends:
                                            raise NativeBackendIntegrityError('native later backend precheck changed') from exc
                                        item.update(status='unavailable', reason=exc.reason, backend=exc.backend, cleanup_verified=True)
                                        raise
                                    finally:
                                        item['elapsed_seconds'] = time.monotonic() - started
                                        if item['elapsed_seconds'] > seconds:
                                            raise TimeoutError('native late fit beyond deadline')
                                    host.check()
                                    expected = {'naive': 'naive', 'ets': 'ets', 'gbm': 'sklearn'}[model] if observed else 'fixture'
                                    if fit.backend != expected or fit.versions != versions or (key in backends and backends[key] != fit.backend):
                                        raise NativeBackendIntegrityError('native returned backend/version drift')
                                    backends[key] = fit.backend
                                    item.update(status='ok', backend=fit.backend, cleanup_verified=True)
                                    provenance['backends'][f'{sid}:{model}:{origin}'] = fit.backend
                                    return fit
                                rows.extend(forecast_from_training(training, origin, model, p, spec, split, bounded))
                                provenance['seeds'][f'{sid}:{origin}'] = canonical_seed(p, sid, origin)
                                provenance['training_cutoffs'][f'{sid}:{origin}'] = cutoff
                                provenance['source_slice_sha256'][f'{sid}:{origin}'] = _sha(canonical(training))
                                provenance['model_fit_seeds'][f'{sid}:{model}'] = 0 if model == 'gbm' else None
                            except (NativeFitUnavailable, _NativeBudgetUnavailable) as exc:
                                failed[key] = exc.reason if isinstance(exc, NativeFitUnavailable) else str(exc)
                                break
                            except BaseException:
                                if item is not None:
                                    item.update(status='failed', reason='native_integrity_unverified', cleanup_verified=False)
                                raise
                            finally:
                                # Charge the remaining wall interval once, including post-fit checks/processing.
                                budget.charge(sid, model, split, time.monotonic() - before - precharged)
                        else:
                            predictions.extend(rows)
                if not predictions:
                    raise RuntimeError('native all full cohorts unavailable; HOLD')
                return predictions

            host.freeze('validation', predict('validation'))
            validation = host.attach('validation'); calibration, calibrated = {}, []
            for sid in c['selected_ids']:
                for model in ('naive', 'ets', 'gbm'):
                    key = (sid, model)
                    if key in failed or (sid, 'naive') in failed:
                        continue
                    start = time.monotonic(); cohort = []
                    host.check()
                    try:
                        for horizon in (1, 3, 6):
                            factor, adjusted = calibrate_intervals([r for r in validation if (r['series_id'], r['model'], r['horizon']) == (sid, model, horizon)])
                            calibration[f'{sid}:{model}:{horizon}'] = factor; cohort.extend(adjusted)
                    except ValueError as exc:
                        if str(exc) != 'calibration_unavailable':
                            raise
                        failed[key] = 'calibration_unavailable'
                        calibration = {k: v for k, v in calibration.items() if not k.startswith(f'{sid}:{model}:')}
                    else:
                        calibrated.extend(cohort)
                    finally:
                        budget.series[sid] += time.monotonic() - start; budget.check()
                    host.check()
            calibrated = [r for r in calibrated if (r['series_id'], r['model']) not in failed and (r['series_id'], 'naive') not in failed]
            selection = {f'{sid}:{h}': value for (sid, h), value in select_models(calibrated).items()}
            host.freeze_selection(selection, calibration, {f'{sid}:{m}': reason for (sid, m), reason in failed.items()})
            host.freeze('test', predict('test'))
            test = host.attach('test')
            available = {(r['series_id'], r['model']) for r in test}
            scored = [r for r in calibrated if (r['series_id'], r['model']) in available and (r['series_id'], 'naive') in available]
            scored.extend(_scaled_interval(r, calibration[f"{r['series_id']}:{r['model']}:{r['horizon']}"]) for r in test
                          if (r['series_id'], 'naive') in available)
            if not scored:
                raise RuntimeError('native no complete naive baseline; HOLD')
            for entry in ledger:
                if entry['status'] == 'pending':
                    key = (entry['series_id'], entry['model'])
                    if key in available and (key[0], 'naive') in available:
                        entry.update(status='scored', backend=backends[key], reason=None)
                    else:
                        entry.update(status='not_evaluated', backend=None, reason=failed.get(key, 'naive_baseline_unavailable'))
            _exclusive_json(output / 'candidate_ledger.json', {'entries': ledger, 'provider_dl': 'out_of_scope'})
            _exclusive_json(output / 'scored_rows.json', {'context': context, 'predictions': scored})
            _exclusive_json(output / 'attempts.json', _generic_attempt_document(budget, context, failed))
            doc = {'schema_version': 'forecast-scorecard.v1', 'metadata': {'synthetic': not observed,
                   'latest_vintage_proxy': True, 'observed_execution_authorized': observed}, 'predictions': scored}
            publish_runner(output / 'legacy', doc, ledger, selection, calibration, provenance)
            verify_runner_run(output / 'legacy'); host.check()
        read.lifetime.check(postclose=True); budget.check()
        artifacts = _native_artifacts(output)
        if set(artifacts) != _native_expected_artifacts(scored):
            raise ValueError('native publication artifact set mismatch')
        unavailable = [{'series_id': sid, 'model': model, 'reason': reason} for (sid, model), reason in sorted(failed.items())]
        selected_missing = [{'series_id': sid, 'horizon': int(h), 'model': model,
                             'reason': failed.get((sid, model), 'naive_baseline_unavailable')}
                            for k, model in selection.items() for sid, h in [k.rsplit(':',1)] if (sid, model) not in available]
        status = ('generic_complete' if observed else 'generic_fixture_complete') + ('_with_unavailable' if unavailable else '')
        closure = {'schema_version': 'root-generic-host-closure.v1', 'status': status, 'context': context,
                   'is_synthetic': not observed, 'observed_authenticity': observed, 'observed_execution_authorized': observed,
                   'full_parent_snapshot_exposed': True, 'physical_test_actuals_untouched': False,
                   'close_postclose_verified': True, 'artifact_sha256': artifacts,
                   'complete_cohorts': [{'series_id': sid, 'model': m} for sid, m in sorted(available) if (sid, 'naive') in available],
                   'unavailable_cohorts': unavailable, 'selected_unavailable': selected_missing,
                   'candidate_ledger_sha256': artifacts['candidate_ledger.json'], 'attempts_sha256': artifacts['attempts.json'],
                   'validation_prediction_sha256': artifacts['validation_prediction_freeze.json'],
                   'selection_sha256': artifacts['selection_freeze.json'], 'test_prediction_sha256': artifacts['test_prediction_freeze.json']}
        budget.check()
        _exclusive_json(output / 'NATIVE_COMPLETE.json', closure)
        _seal_native_run(binding, token, canonical(closure)); budget.check()
        _verify_root_generic(output, binding); budget.check()
        return closure
    except BaseException as exc:
        _fail_native_run(binding, token)
        try:
            owned = made is not None and _identity(os.lstat(output), True) == made and _identity(os.lstat(output.parent), True) == run['parent_identity']
        except OSError:
            owned = False
        if owned and not os.path.lexists(output / 'FAILURE.json'):
            reason = 'owned_child_cleanup_unverified' if isinstance(exc, OwnedChildCleanupError) else 'fit_deadline' if isinstance(exc, TimeoutError) else 'native_integrity_unverified'
            _exclusive_json(output / 'FAILURE.json', {'schema_version': 'root-generic-host-closure.v1', 'status': 'incomplete', 'reason': reason})
            if not os.path.lexists(output / 'candidate_ledger.json'):
                _exclusive_json(output / 'candidate_ledger.json', {'entries': [{**e, 'status': 'not_evaluated', 'backend': None,
                                                      'reason': e['reason'] or reason} for e in ledger], 'provider_dl': 'out_of_scope'})
            if not os.path.lexists(output / 'attempts.json'):
                _exclusive_json(output / 'attempts.json', _generic_attempt_document(budget, context, failed, reason))
        raise





def _generic_context(run):
    from .observed_admission import canonical,_native_plain
    c=run['control']
    return {'profile':c['profile'],'acceptance_id':c['acceptance_id'],'run_id':c['run_id'],
            'root_control_sha256':run['root_pin']['sha256'],
            'source_sha256':{k:s['file']['sha256'] for k,s in c['sources'].items()},
            'catalog_sha256':c['catalog']['sha256'],'selected_ids':list(c['selected_ids']),'batch_index':c['batch_index'],
            'protocol_sha256':c['controls']['protocol']['sha256'],'manifest_sha256':c['controls']['manifest']['sha256'],
            'code_sha256':{p['repository_path']:p['file']['sha256'] for p in c['code']},
            'runtime_pin_sha256':_sha(canonical(_native_plain(c['runtime']))),'output':run['output'],
            'as_of':'2026-10-09','analysis_as_of_utc':c['analysis_as_of_utc'],
            'budget_policy_sha256':_sha(canonical(_native_plain(c['budgets'])))}


def _generic_ledger(run):
    # Derived from bound catalog; body eligibility is independently reconstructed
    # before any fit. A ledger cell is model x split x horizon, not an origin.
    from .observed_admission import decode
    life=__import__('contextlib').closing
    from .observed_admission import _generic_pins_lifetime,_native_pin_open
    with life(_generic_pins_lifetime(run)) as pins:
        records=decode(_native_pin_open(pins,run['control']['catalog'],private=True))['series']
    selected=set(run['control']['selected_ids']);specs={r['series_id']:r for r in records}
    if not selected.issubset(specs):raise ValueError('generic selected catalog ID absent')
    return [{'series_kind':specs[sid]['series_kind'],'series_id':sid,'unit':specs[sid]['quoted_unit'],
             'model':model,'split':split,'horizon':h,'status':'pending','reason':None,'backend':None,'metrics':None}
            for sid in run['control']['selected_ids'] for model in ('naive','ets','gbm')
            for split in ('validation','test') for h in (1,3,6)]


class _GenericBudget(_NativeBudget):
    def __init__(self,selected):
        super().__init__();self.selected=tuple(selected)
        if not 1<=len(self.selected)<=2 or len(set(self.selected))!=len(self.selected):raise ValueError('generic initial processing batch bound')

    def begin_attempt(self,sid,model,split,origin,timeout):
        if (sid not in self.selected or model not in ('naive','ets','gbm') or split not in ('validation','test')
                or origin not in _range(MONTHS[split+'_start'],MONTHS[split+'_end'])
                or type(timeout) not in (int,float) or not math.isfinite(timeout) or not 0<timeout<=60
                or sum(k[0]==sid for k in self.seen)>=144
                or len(self.seen)>=144*len(self.selected)):
            raise RuntimeError('generic invalid/repeated/over-budget attempt')
        return super().begin_attempt(sid,model,split,origin,timeout)


def _generic_attempt_document(budget,context,failed,terminal=None):
    skipped=[]
    for sid in budget.selected:
        for model in ('naive','ets','gbm'):
            for split in ('validation','test'):
                for origin in _range(MONTHS[split+'_start'],MONTHS[split+'_end']):
                    if (sid,model,split,origin) not in budget.seen:
                        skipped.append({'series_id':sid,'model':model,'split':split,'origin':origin,'status':'skipped',
                            'reason':terminal or failed.get((sid,model),'naive_baseline_unavailable'),
                            'expected_backend':{'naive':'naive','ets':'ets','gbm':'sklearn'}[model],
                            'backend':None,'planned_timeout':None,'elapsed_seconds':None,'cleanup_verified':False})
    return {'context':context,'attempts':budget.attempts,'skipped':skipped,'series_seconds':dict(budget.series),
            'elapsed_seconds':time.monotonic()-budget.started}


class _GenericPhases(_NativePhases):
    def check(self):
        super().check()
        from .observed_admission import _native_utc,datetime,timezone,native_run_record,_native_state
        c=self.run['control']
        if datetime.now(timezone.utc)>=_native_utc(c['expires_at_utc']):raise ValueError('generic host expired during phase')
        # Recheck active host, not merely copied immutable fields.
        native_run_record(self.binding)
        if _native_state(self.binding)['state']!='RUNNING':raise ValueError('generic live host phase failed')

    def freeze(self,split,rows):
        from .forecast_scorecard import FIELDS
        specs={r['series_id']:r for r in self.read.receipt['catalog'] if r['series_id'] in self.run['control']['selected_ids']}
        groups=defaultdict(set)
        for row in rows:
            sid=row.get('series_id');rec=specs.get(sid)
            if (rec is None or set(row)!=FIELDS-{'actual'} or row['model'] not in ('naive','ets','gbm')
                    or type(row['horizon']) is not int or row['horizon'] not in (1,3,6)
                    or row['unit']!=rec['quoted_unit'] or row['series_kind']!=rec['series_kind']
                    or row['source_id']!=rec['source_id'] or row['is_synthetic'] is not (self.run['control']['profile']=='generated-fixture-test-v1')
                    or row['training_cutoff']!=_name(_month(row['origin'])-1)):
                raise ValueError('generic exact row/native metadata mismatch')
            groups[sid,row['model']].add(row['horizon'])
        if not groups or any(h!={1,3,6} for h in groups.values()):raise ValueError('generic full horizon cohort required')
        if split=='test' and 'validation' not in self.attached:raise ValueError('generic validation attachment absent')
        _FixturePhases.freeze(self,split,rows)


def _verify_root_generic(output_dir,binding):
    """Consume sealed artifacts and metadata only; never source bodies or refits."""
    from contextlib import closing
    from .observed_admission import (native_run_record,_native_completion,_native_state,_fail_native_run,
                                     _generic_pins_lifetime,_RunDirectoryAnchor)
    run=native_run_record(binding);ids,issued_sha=_native_completion(binding);output=Path(output_dir)
    if str(output)!=run['output']:raise ValueError('generic foreign output')
    try:
        if os.path.lexists(output/'FAILURE.json'):raise ValueError('generic failure dominates')
        with closing(_generic_pins_lifetime(run)) as pins,_RunDirectoryAnchor(output,ids) as anchor:
            closure,raw=_owned_json(output/'NATIVE_COMPLETE.json');c=run['control'];observed=c['profile']=='generic-observed-v1'
            statuses=('generic_complete','generic_complete_with_unavailable') if observed else ('generic_fixture_complete','generic_fixture_complete_with_unavailable')
            if (_sha(raw)!=issued_sha or closure.get('schema_version')!='root-generic-host-closure.v1'
                    or closure.get('status') not in statuses or closure.get('context')!=_generic_context(run)
                    or closure.get('is_synthetic') is not (not observed) or closure.get('observed_authenticity') is not observed
                    or closure.get('observed_execution_authorized') is not observed or closure.get('close_postclose_verified') is not True
                    or closure.get('full_parent_snapshot_exposed') is not True or closure.get('physical_test_actuals_untouched') is not False):
                raise ValueError('generic closure/context mismatch')
            hashes=_native_artifacts(output);scored,_=_owned_json(output/'scored_rows.json')
            if set(hashes)!=_native_expected_artifacts(scored['predictions']) or hashes!=closure['artifact_sha256']:
                raise ValueError('generic exact artifact set/hash mismatch')
            for name in ('HOST_PENDING.json','source_receipt.json','validation_prediction_freeze.json','selection_freeze.json',
                         'test_prediction_freeze.json','scored_rows.json','attempts.json'):
                doc,_=_owned_json(output/name)
                if doc.get('context')!=closure['context']:raise ValueError('generic foreign artifact context')
            ledger,_=_owned_json(output/'candidate_ledger.json');selection,_=_owned_json(output/'selection_freeze.json')
            attempts,_=_owned_json(output/'attempts.json');history=attempts['attempts']+attempts['skipped']
            planned={(s,m,split,o) for s in c['selected_ids'] for m in ('naive','ets','gbm') for split in ('validation','test')
                      for o in _range(MONTHS[split+'_start'],MONTHS[split+'_end'])}
            cells={(s,m,split,h) for s in c['selected_ids'] for m in ('naive','ets','gbm') for split in ('validation','test')for h in (1,3,6)}
            if (len(history)!=len(planned) or {(r['series_id'],r['model'],r['split'],r['origin'])for r in history}!=planned
                    or any(r['status'] not in ('ok','unavailable') or r['cleanup_verified'] is not True for r in attempts['attempts'])
                    or any(r['status']!='skipped' or r['elapsed_seconds'] is not None or r['cleanup_verified'] is not False for r in attempts['skipped'])
                    or len(ledger['entries'])!=len(cells) or {(r['series_id'],r['model'],r['split'],r['horizon'])for r in ledger['entries']}!=cells):
                raise ValueError('generic dynamic complete ledger/attempt mismatch')
            for key,name in [('candidate_ledger_sha256','candidate_ledger.json'),('attempts_sha256','attempts.json'),
                             ('validation_prediction_sha256','validation_prediction_freeze.json'),('selection_sha256','selection_freeze.json'),
                             ('test_prediction_sha256','test_prediction_freeze.json')]:
                if closure[key]!=hashes[name]:raise ValueError('generic phase/ledger hash mismatch')
            if selection['validation_prediction_sha256']!=hashes['validation_prediction_freeze.json']:
                raise ValueError('generic selection must bind validation')
            verify_runner_run(output/'legacy')
            if _native_artifacts(output)!=hashes or os.path.lexists(output/'FAILURE.json'):raise ValueError('generic late output drift')
            pins.check();anchor.check()
        if _native_completion(binding)!=(ids,issued_sha):raise ValueError('generic live seal changed after close')
    except BaseException:
        _fail_native_run(binding,_native_state(binding)['token']);raise
    return True

if __name__ == '__main__' and sys.argv[1:] == ['--fit-child']:
    raise SystemExit(_fit_child_main())
if __name__ == '__main__' and sys.argv[1:] == ['--native-fit-child-v1']:
    raise SystemExit(_native_fit_child_main())
