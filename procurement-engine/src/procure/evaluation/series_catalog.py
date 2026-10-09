"""Bounded WB/BIS catalog projection; this pure parser grants no execution authority.

Callers of the host seam must hold ROOT-issued identity before supplying original
bytes. Latest-vintage m+1 availability is a retrospective proxy, never PIT proof.
Every source line remains visible; processing batches do not cap the universe.
"""
from __future__ import annotations

import csv
import hashlib
import io
import math
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from types import MappingProxyType

MAX_WB_BYTES = 12 * 1024 * 1024
MAX_BIS_BYTES = 32 * 1024 * 1024
MAX_BIS_ROWS = 200000
PROFILES = ('generic-observed-v1', 'generated-fixture-test-v1')
AVAILABILITY = 'month_m_available_at_end_of_m_plus_1; retrospective latest-vintage proxy; not PIT'
WINDOW = tuple(f'{n // 12:04d}-{n % 12 + 1:02d}' for n in range(2016*12+8, 2026*12+8))
BIS_REQUIRED = frozenset(('FREQ','REF_AREA','CURRENCY','COLLECTION','TIME_PERIOD','OBS_VALUE','OBS_STATUS','UNIT_MULT'))
# SDMX attributes have explicit bounded support, not an arbitrary field bag.
BIS_OPTIONAL = frozenset(('STRUCTURE','STRUCTURE_ID','ACTION','OBS_CONF','DECIMALS','TITLE',
                          'TITLE_COMPL','SOURCE','COMMENT','BREAKS','UNIT','TIME_FORMAT'))


def canonical_catalog(records):
    from .observed_admission import canonical
    records = list(records)
    ids = [r['series_id'] for r in records]
    if len(set(ids)) != len(ids):
        raise ValueError('catalog duplicate stable ID')
    return canonical({'schema_version':'forecast-series-catalog.v1',
                      'series':sorted(records, key=lambda r:r['series_id'])})


def finite_batches(records):
    """Return a finite plan, never execute or automatically chain controls."""
    canonical_catalog(records)
    ids = sorted(r['series_id'] for r in records if r['forecast_eligible'])
    return tuple(tuple(ids[i:i+2]) for i in range(0,len(ids),2))


@dataclass(frozen=True)
class CatalogSnapshot:
    records: tuple
    values: MappingProxyType
    canonical_bytes: bytes


def _month(label, compact=False):
    pattern = r'([0-9]{4})M(0[1-9]|1[0-2])' if compact else r'([0-9]{4})-(0[1-9]|1[0-2])'
    m = re.fullmatch(pattern,label) if isinstance(label,str) else None
    if m is None or int(m[1]) < 1:
        raise ValueError('source monthly calendar label invalid')
    return m[1]+'-'+m[2]


def _text(value, width=512):
    if not isinstance(value,str) or not value.strip() or len(value)>width or any(ord(c)<32 for c in value):
        raise ValueError('source field type/width invalid')
    return value.strip()


def _source(source, raw, profile):
    from .observed_admission import _native_utc, _generic_source_pin, _native_keys
    _native_keys(source,('file','provider','source_id','source_url','retrieved_utc','rights_disposition_sha256','source_disposition_sha256'))
    if source['provider'] not in ('World Bank','BIS'):
        raise ValueError('catalog provider profile mismatch')
    _generic_source_pin(source['file'],MAX_WB_BYTES if source['provider']=='World Bank' else MAX_BIS_BYTES)
    if profile not in PROFILES:
        raise ValueError('catalog profile invalid')
    pin = source['file']
    if len(raw)!=pin['bytes'] or hashlib.sha256(raw).hexdigest()!=pin['sha256']:
        raise ValueError('catalog original byte binding mismatch')
    if _native_utc(source['retrieved_utc']).date().isoformat()>'2026-10-09':
        raise ValueError('catalog retrieval after generic as-of')
    for key in ('provider','source_id','source_url'):
        _text(source[key])
    for key in ('rights_disposition_sha256','source_disposition_sha256'):
        if re.fullmatch('[0-9a-f]{64}',source[key]) is None:
            raise ValueError('catalog rights/source disposition hash required')


def _base(source, profile, sid, label, kind, measure, unit, currency, locator, evidence):
    return {'schema_version':'SeriesSpecV1', 'series_id':sid, 'profile_id':profile,
            'source_profile_id':'wb-monthly-v1' if kind=='commodity' else 'bis-xru-monthly-average-v1',
            'original':source['file'], 'provider':source['provider'], 'source_id':source['source_id'],
            'source_url':source['source_url'], 'rights_disposition_sha256':source['rights_disposition_sha256'],
            'source_disposition_sha256':source['source_disposition_sha256'], 'locator':locator,
            'header_unit_evidence':evidence, 'label':label, 'series_kind':kind, 'measure_type':measure,
            'quoted_unit':unit, 'currency':currency, 'base_currency':'USD' if kind=='fx' else None,
            'unit_multiplier':'0', 'frequency':'M', 'required_window':{'start':'2016-09','end':'2026-08','months':120},
            'retrieved_utc':source['retrieved_utc'], 'analysis_as_of':'2026-10-09',
            'publication_availability_policy':AVAILABILITY, 'vintage_status':'latest_vintage_historical_publication_unknown',
            'full_parent_snapshot_exposed':True, 'physical_test_actuals_untouched':False,
            'semantic_review_flags':[], 'catalog_visible':True}


def _finish(record, observations, states, *, flags=(), reason=None, duplicates=0):
    valid = {m:v for m,v in observations.items() if states.get(m)=='valid'}
    mask=[states.get(m,'missing') for m in WINDOW]
    selected={m:valid[m] for m in WINDOW if m in valid}
    distinct=len(set(selected.values()))
    record.update(source_first_period=min(states) if states else None, source_last_period=max(states) if states else None,
                  first_valid_period=min(valid) if valid else None, last_valid_period=max(valid) if valid else None,
                  required_window_mask=mask, missing_count=mask.count('missing'), invalid_count=mask.count('invalid'),
                  status_count=mask.count('status'), duplicate_count=duplicates, distinct_window_values=distinct,
                  constant=distinct==1 and len(selected)==120, near_constant=0<distinct<=3,
                  semantic_review_flags=sorted(set(flags)))
    if reason is None and any(m!='valid' for m in mask):
        reason='incomplete_or_invalid_required_window'
    record.update(observation_eligible=all(m=='valid' for m in mask), forecast_eligible=reason is None,
                  eligibility_reason=reason)
    return selected


def _wb_unit(unit):
    if not isinstance(unit,str):return None,None,'unit_evidence_absent'
    text=_text(unit).strip('()').strip()
    if re.fullmatch(r'2010\s*=\s*100',text):return 'index','2010=100',None
    # Unit vocabulary, never a material-name allowlist. Preserve native quoted scale.
    m=re.fullmatch(r'(\$|cents)/([A-Za-z0-9 .]+)',text)
    if m and m[2] in {'mt','kg','bbl','mmbtu','oz','troy oz','cu m','cubic meter','m3','sheet','lb','bu','dmtu','dry kg'}:
        return 'price',text,'physical_unit_interpretation_pending' if m[2]=='dmtu' else None
    return 'price',text,'unit_interpretation_unavailable'


def _wb(raw, source, profile):
    from .observed_admission import _zip_guard, _native
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    if len(raw)>MAX_WB_BYTES:raise ValueError('WB source byte bound')
    _zip_guard(raw)
    stream=io.BytesIO(raw); book=load_workbook(stream,read_only=True,data_only=False,keep_links=False)
    records,values=[],{}
    try:
        for sheet_name,headers,start in [('Monthly Prices',(5,),7),('Monthly Indices',(6,7,8),9)]:
            if sheet_name not in book.sheetnames:raise ValueError('WB required sheet absent')
            sheet=book[sheet_name]
            if not sheet.max_row or not sheet.max_column or sheet.max_row*sheet.max_column>100000 or sheet.max_row>1006:
                raise ValueError('WB sheet expansion bound')
            # Iterate once; native same-byte book, bounded by ZIP and sheet limits.
            rows=list(sheet.iter_rows())
            index_evidence=[]
            for row in rows[:start-1]:
                for cell in row:
                    if cell.data_type not in {'f','e'} and isinstance(cell.value,str) and re.search(r'2010\s*=\s*100',cell.value):
                        index_evidence.append({'cell':cell.coordinate,'text':_text(cell.value)})
            named=[]
            for col in range(2,sheet.max_column+1):
                evidence=[]
                for r in headers:
                    cell=rows[r-1][col-1]
                    if cell.value is not None:
                        evidence.append({'cell':cell.coordinate,'text':_text(cell.value), 'data_type':cell.data_type})
                if not evidence:continue
                label=' / '.join(e['text'] for e in evidence)
                if len(label)>512:raise ValueError('WB header assembled field bound')
                header_bad=any(e['data_type'] in {'f','e'} for e in evidence)
                if sheet_name=='Monthly Prices':
                    cell=rows[5][col-1];measure,unit,reason=_wb_unit(cell.value)
                    measure=measure or 'price'
                    unit_evidence=[{'cell':cell.coordinate,'text':str(cell.value),'data_type':cell.data_type}]
                    if cell.data_type in {'f','e'}:reason='unit_evidence_ambiguous'
                else:
                    measure,unit='index','2010=100';reason=None if index_evidence else 'index_base_evidence_absent'
                    unit_evidence=index_evidence
                if header_bad:reason='header_evidence_ambiguous'
                column=get_column_letter(col)
                sid=f'WB:wb-monthly-v1:{sheet_name.replace(" ","_")}:{column}'
                rec=_base(source,profile,sid,label,'commodity',measure,unit or 'unavailable','USD' if measure=='price' else None,
                          {'sheet':sheet_name,'column':column,'header_rows':list(headers)},
                          {'headers':evidence,'unit':unit_evidence})
                named.append((col,rec,reason,{},{}))
            prior=None
            for row in rows[start-1:]:
                label=row[0].value
                # A trailing blank row is not a monthly data row; reject hidden values.
                if label is None:
                    if any(c.value is not None for c in row):raise ValueError('WB data without month')
                    continue
                month=_month(label,True); index=int(month[:4])*12+int(month[5:])-1
                if prior is not None and index!=prior+1:raise ValueError('WB duplicate/nonconsecutive source months')
                prior=index
                for col,rec,reason,obs,states in named:
                    cell=row[col-1]
                    if cell.value is None:states[month]='missing'
                    else:
                        try:obs[month]=_native(cell);states[month]='valid'
                        except ValueError:states[month]='invalid'
            for col,rec,reason,obs,states in named:
                values[rec['series_id']]=_finish(rec,obs,states,reason=reason)
                records.append(rec)
    finally:
        book.close();stream.close()
    return records,values


def _bis(raw,source,profile,dispositions):
    if len(raw)>MAX_BIS_BYTES:raise ValueError('BIS source byte bound')
    try:text=raw.decode('utf-8-sig',errors='strict')
    except UnicodeError as exc:raise ValueError('BIS UTF-8 required') from exc
    reader=csv.reader(io.StringIO(text,newline=''),strict=True)
    try:header=next(reader)
    except StopIteration:raise ValueError('BIS schema absent')
    if len(header)!=len(set(header)) or not BIS_REQUIRED.issubset(header) or set(header)-BIS_REQUIRED-BIS_OPTIONAL:
        raise ValueError('BIS fixed SDMX schema invalid')
    grouped={}; area_codes={}
    for count,row in enumerate(reader,1):
        if count>MAX_BIS_ROWS:raise ValueError('BIS data row bound')
        if len(row)!=len(header):raise ValueError('BIS row schema width mismatch')
        for key,value in zip(header,row):
            width=1024 if key in {'TITLE','TITLE_COMPL','SOURCE','COMMENT','BREAKS'} else 128
            if len(value)>width or any(ord(c)<32 for c in value):raise ValueError('BIS field width/characters invalid')
        r=dict(zip(header,row));area=_text(r['REF_AREA'],8);code=_text(r['CURRENCY'],3)
        if re.fullmatch('[A-Z0-9]{2,8}',area) is None or re.fullmatch('[A-Z]{3}',code) is None:
            raise ValueError('BIS area/currency code invalid')
        month=_month(r['TIME_PERIOD'])
        g=grouped.setdefault((area,code),{'obs':{},'states':{},'duplicates':0,'flags':set(),'regime':set()})
        if month in g['states']:
            g['duplicates']+=1;g['states'][month]='duplicate';g['obs'].pop(month,None)
            continue
        state='valid'
        if (r['FREQ'],r['COLLECTION'],r['OBS_STATUS'],r['UNIT_MULT'])!=('M','A','A','0'):state='status'
        try:
            # Parse original token; no scaling, imputation or area aggregation.
            value=Decimal(r['OBS_VALUE']); number=float(value)
            if not value.is_finite() or value<=0 or not math.isfinite(number) or number<=0:state='invalid'
        except (InvalidOperation,ValueError,OverflowError):state='invalid';number=None
        g['states'][month]=state
        if state=='valid':g['obs'][month]=number
        attrs={k:r[k] for k in ('TITLE','TITLE_COMPL','SOURCE','COMMENT','BREAKS') if r.get(k)}
        for k,v in attrs.items():
            if k=='BREAKS' or re.search(r'currency (?:change|break)|redenominat|discontinu|regime',v,re.I):
                g['flags'].add('currency_regime_review_required');g['regime'].add((k,v))
        area_codes.setdefault(area,set()).add(code)
    records,values=[],{}
    for (area,code),g in sorted(grouped.items()):
        flags=set(g['flags'])
        if len(area_codes[area])>1:flags.add('reference_area_currency_change')
        evidence={'required_attributes':{'FREQ':'M','COLLECTION':'A','OBS_STATUS':'A','UNIT_MULT':'0'},
                  'regime_evidence':[{'attribute':k,'value':v} for k,v in sorted(g['regime'])],
                  'area_currency_codes':sorted(area_codes[area]),'csv_header':header}
        from .observed_admission import canonical
        regime_id=hashlib.sha256(canonical({k:evidence[k] for k in ('regime_evidence','area_currency_codes')})).hexdigest()
        sid=f'BIS:bis-xru-monthly-average-v1:{area}:{code}:A:R{regime_id}'
        rec=_base(source,profile,sid,code+' / '+area,'fx','exchange_rate',code+'/USD',code,
                  {'REF_AREA':area,'CURRENCY':code,'FREQ':'M','COLLECTION':'A'},evidence)
        rec['category']='reserve_basket' if code=='XDR' else 'currency'
        rec['regime_id']=regime_id
        rec['regime_decision']='no flagged break in parsed metadata; historical authenticity not inferred'
        reason='duplicate_source_key' if g['duplicates'] else None
        if flags:
            disposition=dispositions.get(sid)
            evidence_sha=hashlib.sha256(canonical(evidence)).hexdigest()
            # A disposition binds the exact parsed evidence, not an APPROVED flag.
            if not isinstance(disposition,dict) or set(disposition)!= {'evidence_sha256','disposition_sha256','decision'} or (
                    disposition['evidence_sha256']!=evidence_sha or disposition['decision']!='retain_single_source_regime'
                    or re.fullmatch('[0-9a-f]{64}',disposition['disposition_sha256']) is None):
                reason=reason or 'currency_regime_disposition_required'
                rec['regime_decision']='unavailable_pending_explicit_regime_disposition'
            else:
                rec['regime_disposition']=disposition
                rec['regime_decision']='explicit_root_disposition_single_source_no_splice'
        if code=='USD':reason=reason or 'identity_currency_target_excluded'
        values[sid]=_finish(rec,g['obs'],g['states'],flags=flags,reason=reason,duplicates=g['duplicates'])
        records.append(rec)
    by_code={}
    for rec in records:by_code.setdefault(rec['currency'],[]).append(rec)
    for code,lines in by_code.items():
        candidates=sorted((r for r in lines if r['forecast_eligible']),key=lambda r:r['locator']['REF_AREA'])
        if code=='EUR':candidates=[r for r in candidates if r['locator']['REF_AREA']=='DE']
        chosen=candidates[0]['series_id'] if candidates else None
        for rec in lines:
            rec['alternative_areas']=sorted(r['locator']['REF_AREA'] for r in lines)
            rec['representative_series_id']=chosen
            rec['representative_policy']='EUR=DE; otherwise lexicographic eligible area; no splice or averaging'
            if rec['forecast_eligible'] and rec['series_id']!=chosen:
                rec.update(forecast_eligible=False,eligibility_reason='alternative_area_not_representative' if chosen else 'canonical_area_unavailable')
    unknown=set(dispositions)-set(values)
    if unknown:raise ValueError('foreign regime disposition')
    return records,values


def build_catalog(wb_raw,bis_raw,*,sources,profile,regime_dispositions=None):
    """Parse only supplied bounded bytes; reconstruct all metadata and eligibility.

    ROOT's host independently verifies canonical catalog equality before fitting.
    Source-specific limits are checked here too; this function is not admission.
    """
    if set(sources)!= {'wb','bis'}:raise ValueError('both original source descriptors required')
    if not isinstance(wb_raw,bytes) or not isinstance(bis_raw,bytes):raise ValueError('source bytes required')
    if len(wb_raw)>MAX_WB_BYTES or len(bis_raw)>MAX_BIS_BYTES:raise ValueError('source byte bound')
    for key,raw in [('wb',wb_raw),('bis',bis_raw)]:
        if sources[key].get('provider')!=('World Bank' if key=='wb' else 'BIS'):
            raise ValueError('catalog source/provider adapter mismatch')
        _source(sources[key],raw,profile)
    wr,wv=_wb(wb_raw,sources['wb'],profile);br,bv=_bis(bis_raw,sources['bis'],profile,regime_dispositions or {})
    records=tuple(sorted(wr+br,key=lambda r:r['series_id']))
    return CatalogSnapshot(records,MappingProxyType({k:MappingProxyType(v) for k,v in {**wv,**bv}.items()}),canonical_catalog(records))
