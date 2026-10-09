"""Owned XLSX fixtures only; no market original or statistical model fitting."""
import copy
import io
import json
import os
import re
from pathlib import Path
import zipfile

import pytest
from openpyxl import Workbook

from procure.evaluation import observed_admission as admission
from procure.evaluation import observed_runner as runner


def owned_fixture(tmp_path, *, cell=None, archive_edit=None, label_edit=None):
    inputs = tmp_path / "inputs"
    inputs.mkdir(mode=0o700)
    source = inputs / "owned_fake.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Monthly Prices"
    sheet.cell(1, 1, "OWNED SYNTHETIC TEST FIXTURE")
    # Independent fixture contract from the bound current source manifest,
    # rather than copying the reader's constants into its own affirmative test.
    for col, label, unit in ((65, "Copper", "($/mt)"), (63, "Aluminum", "($/mt)")):
        sheet.cell(5, col, label); sheet.cell(6, col, unit)
    for i in range(801):
        index = 1960 * 12 + i
        sheet.cell(i + 7, 1, f"{index // 12:04d}M{index % 12 + 1:02d}")
        sheet.cell(i + 7, 65, float(100 + i))
        sheet.cell(i + 7, 63, float(200 + i))
    if cell is not None:
        sheet.cell(687, 65, cell)  # first selected month2016-09
    if label_edit:
        label_edit(sheet)
    book.save(source); book.close()
    if type(cell) is int and cell > 2**53:
        # openpyxl's writer rounds large ints. Generate the intended native
        # integer token explicitly; this does not certify real XML decimals.
        def exact_integer(members):
            return [(n, re.sub(rb'(<c r="BM687"[^>]*><v>)[^<]+',
                               lambda m: m[1] + str(cell).encode(), d)
                     if n == "xl/worksheets/sheet1.xml" else d) for n,d in members]
        archive_edit = exact_integer
    if archive_edit:
        raw = source.read_bytes()
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            members = [(m.filename, z.read(m)) for m in z.infolist()]
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            for name, data in archive_edit(members):
                z.writestr(name, data)
        source.write_bytes(out.getvalue())
    raw = source.read_bytes()  # owned generator, not an observed source
    receipt = {"schema_version": admission.SCHEMA, "fixture_id": "owned-fixture",
               "is_synthetic": True, "observed_authenticity": False, "observed_execution_authorized": False,
               "source_sha256": admission.digest(raw), "source_bytes": len(raw),
               "as_of_utc": "2026-10-08T00:00:00.000000Z"}
    receipt_raw = admission.canonical(receipt)
    protocol = {**runner.PROTOCOL_FIELDS, "status": "INFRASTRUCTURE_ONLY_OBSERVED_EXECUTION_HOLD",
                "observed_execution_authorized": False, "source_sha256": receipt["source_sha256"],
                "protocol_sha256": "b" * 64, "snapshot_sha256": "a" * 64}
    manifest = {"is_synthetic": True, "observed_authenticity": False, "observed_execution_authorized": False,
                "sha256": receipt["source_sha256"], "bytes": len(raw), "source_id": "owned_fixture"}
    pp, mp = inputs / "protocol.json", inputs / "manifest.json"
    pp.write_bytes(admission.canonical(protocol)); mp.write_bytes(admission.canonical(manifest))
    binding = admission.register_owned_fixture(source, receipt_raw, controls=(pp, mp))
    run = admission.register_owned_fixture_run(binding, pp, mp, tmp_path / "run")
    return source, receipt_raw, binding, run, pp, mp, tmp_path / "run"


def test_registered_public_reader_hashes_and_parses_same_bytes(tmp_path, monkeypatch):
    source, raw, binding, *_ = owned_fixture(tmp_path)
    import openpyxl
    real = openpyxl.load_workbook
    seen = []
    def traced(stream, **kwargs):
        assert isinstance(stream, io.BytesIO)
        seen.append(admission.digest(stream.getvalue()))
        return real(stream, **kwargs)
    monkeypatch.setattr(openpyxl, "load_workbook", traced)
    with admission.open_observed_native_snapshot(source, raw, expected_admission=binding) as read:
        assert list(read.series["copper"])[0] == "2016-09"
        assert len(read.series["copper"]) == len(read.series["aluminum"]) == 120
        assert read.series["copper"]["2016-09"] == 780.0
        assert read.receipt["full_parent_snapshot_exposed"] is True
        assert read.receipt["physical_test_actuals_untouched"] is False
        assert read.receipt["is_synthetic"] is True
        assert read.receipt["observed_authenticity"] is False
        basis = read.receipt["stored_value_basis"]["copper"]["2016-09"]
        assert basis == {"hex": (780.0).hex(), "numerator": "780", "denominator": "1"}
    assert seen == [json.loads(raw)["source_sha256"]]
    assert read.lifetime.closed


@pytest.mark.parametrize("fake", [None, object(), admission.RootAdmissionBinding("owned-fixture")])
def test_unregistered_authority_refuses_before_any_source_access(tmp_path, monkeypatch, fake):
    def forbidden(*args, **kwargs):
        raise AssertionError("body or metadata access before admission")
    monkeypatch.setattr(os, "lstat", forbidden)
    monkeypatch.setattr(os, "open", forbidden)
    with pytest.raises(ValueError, match="registered.*HOLD"):
        admission.open_observed_native_snapshot(tmp_path / "absent.xlsx", b"{}", expected_admission=fake)


def test_copied_binding_and_altered_receipt_are_not_authority(tmp_path):
    source, raw, binding, *_ = owned_fixture(tmp_path)
    with pytest.raises(ValueError, match="registered"):
        admission.open_observed_native_snapshot(source, raw, expected_admission=copy.copy(binding))
    with pytest.raises(ValueError, match="foreign"):
        admission.open_observed_native_snapshot(source, raw + b" ", expected_admission=binding)


@pytest.mark.parametrize("value", [True, "text", "=1+2", "#VALUE!", 0, -1, 1e-310, 9007199254740993])
def test_required_native_cells_refuse_coercion(tmp_path, value):
    source, raw, binding, *_ = owned_fixture(tmp_path, cell=value)
    with pytest.raises(ValueError, match="native|integer"):
        admission.open_observed_native_snapshot(source, raw, expected_admission=binding)


@pytest.mark.parametrize("edit", [
    lambda s: s.cell(687, 65).__setattr__("value", None),
    lambda s: s.cell(7, 1).__setattr__("value", "1960M02"),
    lambda s: s.cell(5, 65).__setattr__("value", "wrong"),
    lambda s: s.cell(807, 1).__setattr__("value", None),
    lambda s: s.cell(7, 200, "dimension bound"),
])
def test_required_cells_labels_headers_and_dimensions_are_complete(tmp_path, edit):
    source, raw, binding, *_ = owned_fixture(tmp_path, label_edit=edit)
    with pytest.raises(ValueError):
        admission.open_observed_native_snapshot(source, raw, expected_admission=binding)


@pytest.mark.parametrize("edit", [
    lambda m: m + [(m[0][0], m[0][1])],
    lambda m: m + [("../rogue.xml", b"<x/>")],
    lambda m: [(n, b'<!DOCTYPE x [<!ENTITY y "z">]><x/>' if n == "xl/workbook.xml" else d) for n,d in m],
    lambda m: m + [("xl/_rels/evil.rels", b'<Relationships><Relationship TargetMode="External" Target="https://invalid/"/></Relationships>')],
    lambda m: m + [("huge.bin", b"x" * (admission.MAX_MEMBER + 1))],
])
def test_zip_boundaries_reject_unsafe_archives(tmp_path, edit):
    source, raw, binding, *_ = owned_fixture(tmp_path, archive_edit=edit)
    with pytest.raises(ValueError, match="ZIP|XML|external"):
        admission.open_observed_native_snapshot(source, raw, expected_admission=binding)


def test_source_replacement_hardlink_and_control_change_refuse(tmp_path):
    source, raw, binding, _, pp, *_ = owned_fixture(tmp_path)
    original = source.read_bytes()
    other = source.with_name("replacement.xlsx")
    other.write_bytes(original); os.replace(other, source)
    with pytest.raises(ValueError, match="identity"):
        admission.open_observed_native_snapshot(source, raw, expected_admission=binding)
    # Re-registration is a new trusted fixture, not mutation of the old grant.
    binding = admission.register_owned_fixture(source, raw, controls=(pp,))
    os.link(source, source.with_name("hardlink"))
    with pytest.raises(ValueError, match="identity"):
        admission.open_observed_native_snapshot(source, raw, expected_admission=binding)


def test_parse_time_source_or_control_replacement_prevents_result(tmp_path, monkeypatch):
    source, raw, binding, _, pp, *_ = owned_fixture(tmp_path)
    real = admission._parse_native
    def changed(data):
        parsed = real(data)
        pp.write_bytes(pp.read_bytes() + b" ")
        return parsed
    monkeypatch.setattr(admission, "_parse_native", changed)
    with pytest.raises(ValueError, match="identity|close"):
        admission.open_observed_native_snapshot(source, raw, expected_admission=binding)


def test_first_fstat_failure_still_closes_owned_fd_once(tmp_path, monkeypatch):
    source, raw, binding, *_ = owned_fixture(tmp_path)
    real_fstat, real_close = os.fstat, os.close
    seen, closed = [], []
    def broken(fd):
        seen.append(fd)
        raise OSError("owned injected first fstat failure")
    def close(fd):
        closed.append(fd); real_close(fd)
    monkeypatch.setattr(os, "fstat", broken)
    monkeypatch.setattr(os, "close", close)
    with pytest.raises(OSError):
        admission.open_observed_native_snapshot(source, raw, expected_admission=binding)
    assert closed == seen and len(closed) == 1
    monkeypatch.setattr(os, "fstat", real_fstat)


def test_reader_close_uncertainty_never_retries_fd(tmp_path, monkeypatch):
    source, raw, binding, *_ = owned_fixture(tmp_path)
    read = admission.open_observed_native_snapshot(source, raw, expected_admission=binding)
    real_close = os.close
    closed = []
    def injected(fd):
        closed.append(fd); real_close(fd)
        if len(closed) == 1:
            raise OSError("injected close uncertainty")
    monkeypatch.setattr(os, "close", injected)
    with pytest.raises(ValueError, match="close/postclose"):
        read.close()
    assert len(closed) == len(set(closed)) == len(read.lifetime.fds)


def test_registration_never_accepts_observed_flags(tmp_path):
    source, raw, *_ = owned_fixture(tmp_path)
    grant = json.loads(raw)
    grant.update(is_synthetic=False, observed_authenticity=True, observed_execution_authorized=True)
    with pytest.raises(ValueError, match="only synthetic"):
        admission.register_owned_fixture(source, admission.canonical(grant))


def test_actual_xml_rows_cannot_hide_behind_small_dimension_metadata(tmp_path):
    def hidden(members):
        return [(name, data.replace(b"</sheetData>",
                    b'<row r="1100"><c r="A1100" t="inlineStr"><is><t>hidden</t></is></c></row></sheetData>')
                 if name == "xl/worksheets/sheet1.xml" else data) for name, data in members]
    source, raw, binding, *_ = owned_fixture(tmp_path, archive_edit=hidden)
    with pytest.raises(ValueError, match="actual.*row|expansion"):
        admission.open_observed_native_snapshot(source, raw, expected_admission=binding)


def test_named_observed_bytes_cannot_be_relabelled_fixture_before_access(tmp_path, monkeypatch):
    source, raw, *_ = owned_fixture(tmp_path)
    grant = json.loads(raw)
    grant["source_sha256"] = "ea1c350827878ea3bbe30e3cda16a13fd3bd5b409b8458940dc94a36b5a33154"
    def forbidden(*args, **kwargs):
        raise AssertionError("observed original accessed")
    monkeypatch.setattr(os, "open", forbidden)
    with pytest.raises(ValueError, match="observed original"):
        admission.register_owned_fixture(source, admission.canonical(grant))


def test_source_hash_mismatch_refuses_before_parser(tmp_path, monkeypatch):
    source, raw, *_ = owned_fixture(tmp_path)
    grant = json.loads(raw); grant["source_sha256"] = "0" * 64
    raw = admission.canonical(grant)
    binding = admission.register_owned_fixture(source, raw)
    def forbidden(*args, **kwargs):
        raise AssertionError("hash mismatch reached parser")
    monkeypatch.setattr(admission, "_parse_native", forbidden)
    with pytest.raises(ValueError, match="hash mismatch"):
        admission.open_observed_native_snapshot(source, raw, expected_admission=binding)


def test_private_root_and_code_identity_changes_refuse_before_body(tmp_path, monkeypatch):
    source, raw, binding, *_ = owned_fixture(tmp_path)
    source.parent.chmod(0o755)
    with pytest.raises(ValueError, match="identity"):
        admission.open_observed_native_snapshot(source, raw, expected_admission=binding)
    source.parent.chmod(0o700)  # owned fixture cleanup, never source marker deletion
    binding = admission.register_owned_fixture(source, raw)
    real_stat = os.lstat
    def changed(path, *args, **kwargs):
        info = real_stat(path, *args, **kwargs)
        if Path(path) == Path(admission.__file__):
            from types import SimpleNamespace
            return SimpleNamespace(st_dev=info.st_dev, st_ino=info.st_ino, st_mode=info.st_mode, st_uid=info.st_uid,
                                   st_nlink=info.st_nlink, st_size=info.st_size, st_mtime_ns=info.st_mtime_ns + 1,
                                   st_ctime_ns=info.st_ctime_ns)
        return info
    monkeypatch.setattr(os, "lstat", changed)
    with pytest.raises(ValueError, match="identity"):
        admission.open_observed_native_snapshot(source, raw, expected_admission=binding)


def test_utf16_dtd_is_refused(tmp_path):
    def edit(members):
        return [(n, '<?xml version="1.0" encoding="utf-16"?><!DOCTYPE x [<!ENTITY y "z">]><x/>'.encode("utf-16")
                 if n == "xl/workbook.xml" else data) for n,data in members]
    source, raw, binding, *_ = owned_fixture(tmp_path, archive_edit=edit)
    with pytest.raises(ValueError, match="DTD/entity"):
        admission.open_observed_native_snapshot(source, raw, expected_admission=binding)
