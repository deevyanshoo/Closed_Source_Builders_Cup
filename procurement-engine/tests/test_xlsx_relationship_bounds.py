"""Workbook relationships must not route unbounded sheets around ZIP preflight."""
from __future__ import annotations

import io
import warnings
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest
from openpyxl import Workbook
from openpyxl.styles import Color, Font

from procure.documents import DocumentReadError, read_document
from procure.documents import spreadsheet


_SHEET = "xl/worksheets/sheet1.xml"
_ALT_SHEET = "xl/alt/sheet1.xml"
_RELS = "xl/_rels/workbook.xml.rels"
_TYPES = "[Content_Types].xml"
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_TYPE_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


def _workbook() -> bytes:
    wb = Workbook()
    wb.active["A1"] = "a"
    wb.active["B2"] = "b"
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _repack(data: bytes, *, target: str | None = None, external: bool = False,
            duplicate_id: bool = False, malformed_rels: bool = False,
            move_sheet: bool = False, duplicate_sheet: bool = False,
            unsupported_type: bool = False) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, "w") as dest:
        for info in source.infolist():
            content = source.read(info)
            name = info.filename
            if name == _RELS:
                if malformed_rels:
                    content = b"<Relationships><Relationship"
                else:
                    root = ElementTree.fromstring(content)
                    rel = next(child for child in root if child.get("Type", "").endswith("/worksheet"))
                    if target is not None:
                        rel.set("Target", target)
                    if external:
                        rel.set("TargetMode", "External")
                    if unsupported_type:
                        rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chartsheet")
                    if duplicate_id:
                        duplicate = ElementTree.SubElement(root, f"{{{_REL_NS}}}Relationship", rel.attrib)
                        duplicate.set("Target", "/xl/styles.xml")
                    content = ElementTree.tostring(root)
            elif name == _TYPES and move_sheet:
                root = ElementTree.fromstring(content)
                sheet = next(child for child in root if child.get("PartName") == "/" + _SHEET)
                sheet.set("PartName", "/" + _ALT_SHEET)
                content = ElementTree.tostring(root)
            elif name == _SHEET and move_sheet:
                name = _ALT_SHEET
            dest.writestr(name, content, compress_type=zipfile.ZIP_DEFLATED)
            if name == _ALT_SHEET and duplicate_sheet:
                with warnings.catch_warnings():
                    warnings.filterwarnings("ignore", message="Duplicate name:", category=UserWarning)
                    dest.writestr(name, content, compress_type=zipfile.ZIP_DEFLATED)
    return output.getvalue()


@pytest.mark.parametrize("target", ["/xl/alt/sheet1.xml", "alt/sheet1.xml"])
def test_alternate_worksheet_path_is_checked_before_both_workbook_loads(monkeypatch, target):
    data = _repack(_workbook(), target=target, move_sheet=True)
    original = spreadsheet.load_workbook
    calls = []

    def recording_load(*args, **kwargs):
        calls.append(kwargs.get("data_only"))
        return original(*args, **kwargs)

    monkeypatch.setattr(spreadsheet, "MAX_SHEET_CELLS", 1)
    monkeypatch.setattr(spreadsheet, "load_workbook", recording_load)
    with pytest.raises(DocumentReadError) as err:
        read_document(data, filename="alternate.xlsx")
    assert err.value.code == "too_large"
    assert calls == []


@pytest.mark.parametrize("target", ["/xl/alt/sheet1.xml", "alt/sheet1.xml"])
def test_valid_alternate_worksheet_path_keeps_locators(target):
    data = _repack(_workbook(), target=target, move_sheet=True)
    doc = read_document(data, filename="alternate.xlsx")
    assert [(segment.locator, segment.text) for segment in doc.segments] == [
        ("Sheet!A1", "a"), ("Sheet!B2", "b")]


@pytest.mark.parametrize("change", [
    {"malformed_rels": True},
    {"target": "https://example.com/sheet.xml", "external": True},
    {"target": "/xl/alt/missing.xml"},
    {"target": "../outside.xml"},
    {"unsupported_type": True},
    {"duplicate_id": True},
    {"target": "/xl/alt/sheet1.xml", "move_sheet": True, "duplicate_sheet": True},
])
def test_invalid_worksheet_relationship_fails_before_workbook_load(monkeypatch, change):
    data = _repack(_workbook(), **change)
    monkeypatch.setattr(spreadsheet, "load_workbook", lambda *args, **kwargs: pytest.fail("workbook loaded"))
    with pytest.raises(DocumentReadError) as err:
        read_document(data, filename="invalid.xlsx")
    assert err.value.code == "unreadable"


def test_theme_zero_is_white_and_theme_one_is_dark():
    wb = Workbook()
    for coordinate, theme in (("A1", 0), ("A2", 1)):
        cell = wb.active[coordinate]
        cell.value = "Visible quote"
        cell.font = Font(color=Color(theme=theme))
    output = io.BytesIO()
    wb.save(output)

    doc = read_document(output.getvalue(), filename="theme.xlsx")
    assert "low_visibility" in doc.segment("Sheet!A1").flags
    assert "low_visibility" not in doc.segment("Sheet!A2").flags


def test_q18_visible_prices_are_not_marked_low_visibility():
    fixture = Path(__file__).parent / "fixtures/offers/heldout/SYN-Q18_southern_cross_two_prices.xlsx"
    doc = read_document(fixture.read_bytes(), filename=fixture.name)
    assert "low_visibility" not in doc.segment("Offers!D9").flags
    assert "low_visibility" not in doc.segment("Terms!A5").flags
